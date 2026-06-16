from __future__ import annotations

import hmac
import html
import hashlib
import io
import json
import logging
import os
import re
import time
import traceback
import uuid
from difflib import SequenceMatcher
from datetime import datetime, timezone
from typing import Any, Dict, List, Literal, Optional, Tuple, cast

from fastapi import APIRouter, File, HTTPException, Request, Response, UploadFile
from fastapi.responses import JSONResponse
from starlette.responses import Response
from pydantic import BaseModel, Field, ValidationError, field_validator

from backend.config.anchor_network_config import (
    ALLOWED_AGREEMENT_ANCHOR_NETWORKS,
    anchor_cadence_summary,
    daily_equivalent_block_count_for_network,
)
from backend.config.runtime_environment import (
    clamp_recipient_token_ttl_seconds,
    recipient_access_token_required,
    recipient_token_ttl_max_seconds,
    recipient_token_ttl_min_seconds,
)
from backend.handlers.verifier_api_handler import get_batch
from backend.proof.agreement_receipt import create_agreement_receipt_response
from backend.security.agreement_read_scope import (
    assert_agreement_full_draft_read_allowed,
    assert_agreement_recipient_write_allowed,
    recipient_access_token_from_request,
    validate_recipient_access_token_for_agreement,
)
from backend.security.recipient_access_token import (
    RECIPIENT_LINK_INVALID_OR_EXPIRED,
    mint_recipient_access_token,
    verify_recipient_access_token,
)
from backend.agreements.placeholder_template_safety import (
    primary_agreement_plain_field_and_value,
    strip_html_agreement_scan_text,
    validate_user_visible_agreement_text,
)
from backend.agreements.premium_dev_context_leak import (
    premium_document_text_has_dev_context_leak,
    sanitize_premium_intake_for_retry,
    serialize_context_clean,
)
from backend.agreements.premium_generation_intelligence import build_premium_generation_intelligence_brief
from backend.agreements.premium_agreement_validation import (
    AgreementValidationResult,
    validatePremiumAgreementDraft,
)
from backend.agreements.premium_agreement_finalization import (
    PremiumFinalizationResult,
    finalize_premium_agreement_if_needed,
)
from backend.agreements.premium_full_draft_quality_gate import (
    build_free_reference_blob,
    build_premium_full_draft_repair_user_payload,
    evaluate_premium_full_draft_quality,
    premium_full_draft_repair_system_prompt,
)
from backend.agreements.premium_simple_consulting_size_guard import enrich_user_payload_for_simple_consulting
from backend.agreements.paid_pro_server_timing import (
    PaidProServerTiming,
    maybe_attach_server_timing_header,
    paid_pro_perf_trace_requested,
)
from backend.cors_policy import attach_cors_from_request, log_premium_full_draft_cors_proof
from backend.agreements.pro_redline_diff import compute_pro_redline_block_diff
from backend.agreements.premium_refine_narrow import (
    classify_narrow_amendment_prompt,
    try_apply_narrow_amendment,
)
from backend.agreements.revision_surgical import (
    MINIMAL_REVISION_RETRY_SUFFIX,
    instruction_requests_material_rewrite,
    is_overbroad_structured_revision,
)
from backend.agreements.premium_intent_schema import (
    build_premium_intent_skeleton,
    evaluate_premium_intent_schema,
    resolve_premium_intent_key,
)
from backend.services.agreement_signing_lock_store import (
    assert_draft_exists,
    clear_signing_lock,
    read_signing_lock,
    write_signing_lock,
)
from backend.utils.timeline_store import TimelineStore
from backend.llm_router import ExternalAIBlockedError, OPENAI_API_KEY, call_legal_llm, resolve_llm_model_for_access_class
from backend.security.privilege_policy import first_privilege_airlock_block_diagnostic
from backend.llm_usage_guard import (
    build_llm_trace_context,
    client_fingerprint,
    peek_recipient_remaining,
    recipient_prompt_allowed,
    recipient_try_acquire_llm_slot,
    usage_response_header,
    validate_instruction_size,
    validate_negotiate_payload_size,
    validate_negotiate_text,
)
from backend.config.agreement_signing_token import (
    SigningTokenSecretMissingInProductionError,
    detected_signing_token_env_var,
    operator_signing_token_secret_configured,
    resolve_signing_token_secret_raw,
    review_link_mint_enabled,
)
from backend.config.feed_anchor_policy import settlement_anchor_network_hint
from backend.proof_status.store import ProofLayerStore
from backend.services import document_service
from backend.services.agreement_draft_store import list_draft_agreement_ids_newest_first, load_draft
from backend.services.agreement_pdf_story_capability import (
    RECIPIENT_PREVIEW_PDF_STORY_RENDER_MODES,
    assess_agreement_pdf_story_capability,
)
from backend.services.agreement_vs01_pdf_seed import agreement_rendered_html_to_pdf_bytes
from backend.services.claw_feed_service import record_public_feed_event_if_applicable
from backend.services.claw_feed_store import get_claw_feed_store
from backend.treasury.treasury_usage_hooks import record_usage_ledger_event
from backend.utils.enforce import org_id_from_subject, resolve_subject_from_request
from backend.usage_economics.constants import WATERMARK_LABEL
from backend.usage_economics.policy import (
    assert_can_complete_agreement,
    assert_can_create_draft,
    review_first_paid_pro_persist_bypass,
    assert_free_incomplete_draft_not_expired,
    assert_registered_owner_matches,
    economics_overlay_for_agreement,
    record_agreement_finalized,
    record_ai_call,
    record_draft_created,
    require_claw_org_id_header,
    usage_summary_for_subject,
    workspace_lists_agreement_for_subject,
)
from backend.utils.canon_json import canon_json_bytes, canon_sha256_hex, sha256_hex
from backend.lawdog_dashboard.draft_persistence import save_draft_and_sync_dashboard_metadata
from backend.lawdog_dashboard.workspace_index import (
    fallback_summary_from_supabase_row,
    merge_workspace_index_agreement_ids,
    supabase_rows_by_id_for_subject,
)

log = logging.getLogger(__name__)


def _save_draft_sync(
    draft: Dict[str, Any],
    request: Optional[Request] = None,
    *,
    subject_ref: Optional[str] = None,
) -> None:
    """Persist draft JSON and sync Phase A Supabase dashboard metadata when org is known."""
    subj = (subject_ref or "").strip() or None
    if not subj and request is not None:
        subj = resolve_subject_from_request(request)
    save_draft_and_sync_dashboard_metadata(draft, subject_ref=subj)


def _openai_key_diagnostics() -> Dict[str, Any]:
    """
    For logs only: key presence, length, last 4 chars of key — never log the full secret.
    Re-reads env at call time (avoids stale module-level snapshot after test monkeypatches).
    """
    k = (os.environ.get("OPENAI_API_KEY") or "").strip()
    if not k:
        return {"openai_key": "missing"}
    return {"openai_key": "present", "openai_key_len": len(k), "openai_key_suffix": k[-4:] if len(k) >= 4 else "?"}


def _classify_premium_llm_failure(exc: BaseException) -> str:
    """Single-line reason for operator logs; never include user PII from prompts."""
    if isinstance(exc, ExternalAIBlockedError):
        return f"airlock:{getattr(exc, 'block_reason', None) or 'blocked'}"
    if isinstance(exc, RuntimeError) and "OPENAI_API_KEY" in str(exc):
        return "openai_key_not_configured"
    if isinstance(exc, (json.JSONDecodeError, ValueError)):
        return "llm_output_parse_failed"
    eid = type(exc).__name__
    if "Timeout" in eid or "ReadTimeout" in eid or "connect" in eid.lower():
        return f"network_or_timeout:{eid}"
    if "OpenAI" in eid or eid in ("APIError", "APIStatusError", "AuthenticationError", "RateLimitError", "InternalServerError"):
        return f"openai_client:{eid}"
    if "BadRequest" in eid or "NotFound" in eid:  # invalid model, etc.
        return f"openai_client:{eid}"
    return eid


router = APIRouter(prefix="/api/agreements", tags=["agreements-v2"])

# Premium paid revision: if normalized draft text is still >= this similar to the prior version, auto-retry LLM.
PREMIUM_REVISION_SIMILARITY_CEILING = float(
    os.environ.get("CLAW_PREMIUM_REVISION_SIMILARITY_CEILING", "0.87")
)
PREMIUM_REVISION_MAX_ATTEMPTS = max(1, int(os.environ.get("CLAW_PREMIUM_REVISION_MAX_ATTEMPTS", "3")))


def _timeline_db_path_agreements() -> str:
    env_data = os.getenv("CLAW_DATA_DIR", "").strip()
    if env_data:
        data_dir = os.path.expanduser(env_data)
    else:
        prod = "/var/lib/claw"
        try:
            if os.path.isdir(prod) and os.access(prod, os.W_OK):
                data_dir = prod
            else:
                data_dir = os.path.expanduser("~/.claw")
        except Exception:
            data_dir = os.path.expanduser("~/.claw")
    return os.path.expanduser(
        os.getenv("CLAW_TIMELINE_DB_PATH", os.path.join(data_dir, "timeline.sqlite3"))
    )


def _agreements_timeline_store() -> TimelineStore:
    return TimelineStore(db_path=_timeline_db_path_agreements())


def _agreements_write_allowed() -> bool:
    return os.getenv("CLAW_NODE_MODE", "api").strip().lower() != "verifier"


def _owner_mutation_guards(request: Request, agreement_id: str, *, surface: str) -> None:
    require_claw_org_id_header(request)
    assert_registered_owner_matches(request, agreement_id)
    assert_free_incomplete_draft_not_expired(agreement_id, surface=surface)


def _signing_lock_active(agreement_id: str) -> bool:
    lock = read_signing_lock(agreement_id)
    return bool(lock and str(lock.get("locked_version_id") or "").strip())


def _assert_negotiation_not_locked(agreement_id: str) -> None:
    if _signing_lock_active(agreement_id):
        raise HTTPException(status_code=400, detail="negotiation_locked")


def _recipient_link_mint_key_ok(request: Request) -> bool:
    required = os.getenv("CLAW_RECIPIENT_LINK_MINT_KEY", "").strip()
    if not required:
        return True
    got = (request.headers.get("X-Claw-Recipient-Link-Mint-Key") or "").strip()
    if len(got) != len(required):
        return False
    return hmac.compare_digest(got.encode("utf-8"), required.encode("utf-8"))


_agreements_log = logging.getLogger("claw.agreements")


def _agreement_id_short(agreement_id: str) -> str:
    aid = (agreement_id or "").strip()
    return aid[:8] if len(aid) >= 8 else aid or "unknown"


def _signing_token_secret_bytes(*, agreement_id: str | None = None) -> bytes:
    try:
        return resolve_signing_token_secret_raw().encode("utf-8")
    except SigningTokenSecretMissingInProductionError as e:
        _agreements_log.warning(
            "[review-first-env-token-secret-missing] agreementIdShort=%s claw_environment=%s",
            _agreement_id_short(agreement_id or ""),
            os.getenv("CLAW_ENVIRONMENT", "local").strip().lower(),
        )
        raise HTTPException(
            status_code=422,
            detail={
                "code": "signing_token_secret_not_configured",
                "message": str(e),
            },
        ) from e


def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def _recipient_access_fail(
    code: str,
    *,
    message: Optional[str] = None,
    status_code: int = 403,
) -> HTTPException:
    return HTTPException(
        status_code=status_code,
        detail={
            "code": code,
            "message": message or RECIPIENT_LINK_INVALID_OR_EXPIRED,
        },
    )


def _draft_dict_fully_executed(draft_body: Dict[str, Any]) -> bool:
    """True when the agreement recorded completion — no further signing ceremony actions."""
    audit = draft_body.get("audit_log") or []
    for e in audit:
        if str(_audit_event_dict(e).get("event_type") or "") == "signed":
            return True
    return False


def _agreement_draft_fully_executed(draft: AgreementDraft) -> bool:
    return _draft_dict_fully_executed(draft.model_dump())


# --- Jurisdiction + party display (high-trust copy; keep deterministic) ---
_US_JURISDICTION_CANON: Dict[str, str] = {
    "al": "Alabama",
    "ak": "Alaska",
    "az": "Arizona",
    "ar": "Arkansas",
    "ca": "California",
    "co": "Colorado",
    "ct": "Connecticut",
    "de": "Delaware",
    "dc": "District of Columbia",
    "fl": "Florida",
    "ga": "Georgia",
    "hi": "Hawaii",
    "id": "Idaho",
    "il": "Illinois",
    "in": "Indiana",
    "ia": "Iowa",
    "ks": "Kansas",
    "ky": "Kentucky",
    "la": "Louisiana",
    "me": "Maine",
    "md": "Maryland",
    "ma": "Massachusetts",
    "mi": "Michigan",
    "mn": "Minnesota",
    "ms": "Mississippi",
    "mo": "Missouri",
    "mt": "Montana",
    "ne": "Nebraska",
    "nv": "Nevada",
    "nh": "New Hampshire",
    "nj": "New Jersey",
    "nm": "New Mexico",
    "ny": "New York",
    "nc": "North Carolina",
    "nd": "North Dakota",
    "oh": "Ohio",
    "ok": "Oklahoma",
    "or": "Oregon",
    "pa": "Pennsylvania",
    "ri": "Rhode Island",
    "sc": "South Carolina",
    "sd": "South Dakota",
    "tn": "Tennessee",
    "tx": "Texas",
    "ut": "Utah",
    "vt": "Vermont",
    "va": "Virginia",
    "wa": "Washington",
    "wv": "West Virginia",
    "wi": "Wisconsin",
    "wy": "Wyoming",
}


def _title_case_words(s: str) -> str:
    """Lightweight title case; small words lowercased when not first token."""
    s = (s or "").strip()
    if not s:
        return s
    small = {"of", "and", "the", "in", "on", "at", "for", "to", "a", "an"}
    parts = re.split(r"(\s+|-)", s)
    out: List[str] = []
    word_idx = 0
    for p in parts:
        if not p:
            continue
        if re.match(r"^\s+$", p) or p == "-":
            out.append(p)
            continue
        w = p
        low = w.lower()
        if word_idx > 0 and low in small:
            out.append(low)
        elif len(w) == 2 and w.isalpha() and w.isupper():
            out.append(w.upper())
        else:
            out.append(w[:1].upper() + w[1:].lower() if len(w) > 1 else w.upper())
        word_idx += 1
    return "".join(out)


def normalize_jurisdiction_display(raw: str) -> str:
    """Map common US inputs to canonical labels; otherwise title-case."""
    s = (raw or "").strip()
    if not s:
        return s
    key = re.sub(r"\s+", " ", s.lower().strip())
    if key in _US_JURISDICTION_CANON:
        return _US_JURISDICTION_CANON[key]
    full_map = {
        "new york": "New York",
        "new jersey": "New Jersey",
        "new mexico": "New Mexico",
        "new hampshire": "New Hampshire",
        "north carolina": "North Carolina",
        "north dakota": "North Dakota",
        "south carolina": "South Carolina",
        "south dakota": "South Dakota",
        "west virginia": "West Virginia",
        "rhode island": "Rhode Island",
        "district of columbia": "District of Columbia",
        "oklahoma": "Oklahoma",
    }
    if key in full_map:
        return full_map[key]
    return _title_case_words(s)


# Tokens models sometimes emit into party names/roles; strip from display + map roles to party_a/party_b by order.
_INTERNAL_PARTY_REF_RE = re.compile(
    r"\s*[\[(]?\s*(?:ORG|PARTY|CLIENT|COMPANY)_\d+\s*[\])]?\s*",
    re.IGNORECASE,
)


def _strip_internal_party_refs_from_name(name: str) -> str:
    s = (name or "").strip()
    if not s:
        return ""
    cleaned = _INTERNAL_PARTY_REF_RE.sub(" ", s)
    return re.sub(r"\s+", " ", cleaned).strip()


def _is_placeholder_party_role(role: str) -> bool:
    s = (role or "").strip()
    if not s:
        return False
    if bool(re.fullmatch(r"[\[(]?\s*(?:ORG|PARTY|CLIENT|COMPANY)_\d+\s*[\])]?", s, re.IGNORECASE)):
        return True
    return bool(re.fullmatch(r"[\[(]?\s*party_\d+\s*[\])]?", s, re.IGNORECASE))


def _fallback_role_for_party_index(idx: int) -> str:
    if idx == 0:
        return "party_a"
    if idx == 1:
        return "party_b"
    return "party"


def _sanitize_agreement_parties_in_order(parties: List[AgreementParty]) -> List[AgreementParty]:
    """Drop empty names after cleanup; remap internal placeholder roles to party_a / party_b by final order."""
    out: List[AgreementParty] = []
    for p in parties or []:
        name = _strip_internal_party_refs_from_name(str(p.name or ""))
        if not name:
            continue
        role_in = str(p.role or "party").strip() or "party"
        role_out = _fallback_role_for_party_index(len(out)) if _is_placeholder_party_role(role_in) else role_in
        pid = (p.id or "").strip() or None
        email = str(getattr(p, "email", None) or "").strip() or None
        phone = str(getattr(p, "phone", None) or "").strip() or None
        out.append(AgreementParty(name=name, role=role_out, id=pid, email=email, phone=phone))
    return out


def _draft_with_sanitized_parties(draft: AgreementDraft) -> AgreementDraft:
    sp = _sanitize_agreement_parties_in_order(list(draft.parties or []))
    return draft.model_copy(update={"parties": sp})


def _party_display_names_role(
    parties: List[AgreementParty],
) -> Tuple[Tuple[str, str], Tuple[str, str]]:
    """(name, role_label) pairs for the first two agreement parties for HTML render."""
    raw = list(parties or [])
    p0 = raw[0] if len(raw) > 0 else None
    p1 = raw[1] if len(raw) > 1 else None

    def nm(idx: int, p: Optional[AgreementParty]) -> str:
        if p is None:
            return "Party A" if idx == 0 else "Party B"
        n = _strip_internal_party_refs_from_name((p.name or "").strip())
        if n:
            return n
        return "Party A" if idx == 0 else "Party B"

    def role_lbl(idx: int, p: Optional[AgreementParty]) -> str:
        if p is None:
            return "Party"
        r = (p.role or "").strip()
        if _is_placeholder_party_role(r):
            return "Client" if idx == 0 else "Consultant"
        low = r.lower()
        if low in ("party_a",):
            return "Client"
        if low in ("party_b",):
            return "Consultant"
        if low in ("party", ""):
            return "Party"
        if not r:
            return "Party"
        return normalize_jurisdiction_display(r) if r.islower() else r

    return (nm(0, p0), role_lbl(0, p0)), (nm(1, p1), role_lbl(1, p1))


class AgreementParty(BaseModel):
    name: str
    role: str
    id: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None


class VersionSnapshot(BaseModel):
    version: int
    created_at: str
    note: Optional[str] = None


class AuditEvent(BaseModel):
    """Looser `event_type` so older drafts with extra audit kinds still validate."""
    event_type: str
    at: str
    field: Optional[str] = None
    value: Optional[Any] = None


class AgreementDraftCreate(BaseModel):
    title: str = ""
    jurisdiction: str = ""
    parties: List[AgreementParty] = Field(default_factory=list)
    purpose: str = ""
    payment_terms: str = ""
    duration: Optional[str] = None
    due_date: Optional[str] = None
    effective_date: Optional[str] = None
    feed_visibility: Literal["private", "link_only", "public"] = "private"
    feed_party_anonymize: bool = False
    feed_show_financial_summary: bool = False
    feed_anchor_network: Optional[str] = None
    # Stub: no processor — persisted for UI + future billing (amount, type, payer, condition).
    payment_request: Optional[Dict[str, Any]] = None
    payment_required: bool = False


class AgreementDraft(AgreementDraftCreate):
    id: str
    created_at: str
    updated_at: str
    versions: List[VersionSnapshot] = Field(default_factory=list)
    audit_log: List[AuditEvent] = Field(default_factory=list)
    review_sent_at: Optional[str] = None
    """Set when review invite Resend delivery orchestration completed (idempotent email guard)."""
    review_invite_emails_sent_at: Optional[str] = None
    """VS01 prepared signing packet for cross-browser recipient hydration (test346)."""
    vs01_signing_packet_v1: Optional[Dict[str, Any]] = None
    """Per-recipient invite delivery registry (JTIs, timestamps, resend counts)."""
    recipient_delivery_v1: Optional[Dict[str, Any]] = None
    workspace_archived_at: Optional[str] = None
    workspace_folder_id: Optional[str] = None
    workspace_tags: List[str] = Field(default_factory=list)
    # Premium / full-document plain text (optional; persisted for Pro surfaces + redline import).
    server_full_document_text: Optional[str] = None
    premium_server_full_document_text: Optional[str] = None
    premium_full_document_text: Optional[str] = None
    document_text: Optional[str] = None
    rendered_document_text: Optional[str] = None
    premium_render_source: Optional[str] = None
    """Pro review redline v1: pending import diff, reviewer suggestions, version ledger (JSON-only)."""
    pro_redline_v1: Optional[Dict[str, Any]] = None


def _merge_agreement_draft(base: AgreementDraft, **updates: Any) -> AgreementDraft:
    data = base.model_dump()
    for k, v in updates.items():
        data[k] = v
    return AgreementDraft.model_validate(data)


class AgreementFieldUpdateRequest(BaseModel):
    field: str
    value: Any


class AgreementParseRequest(BaseModel):
    intake_text: str
    # basic = default create-flow parse (cost tier); premium = post-pay completion re-parse.
    ai_model_class: Literal["basic", "premium"] = "basic"


class AgreementParseExtract(BaseModel):
    """Premium parse only: optional structured fields alongside the flat draft. Always empty if absent from model output."""

    material_asks: List[str] = Field(default_factory=list)
    agreement_family_hint: Optional[str] = None
    confidence: Optional[str] = None


class AgreementParseResponse(BaseModel):
    draft: AgreementDraftCreate
    # Optional observability (no user text). Set for premium always; for basic when CLAW_AGREEMENT_PARSE_CLIENT_DEBUG=1.
    parse_meta: Optional[Dict[str, Any]] = None
    # Premium-only structured extract; null for basic parse.
    extract: Optional[AgreementParseExtract] = None


class PremiumFullDraftContext(BaseModel):
    """Structured context from the client; all fields optional in practice (best-effort)."""

    title: str = ""
    jurisdiction: str = ""
    parties: List[AgreementParty] = Field(default_factory=list)
    purpose: str = ""
    payment_terms: str = ""
    duration: Optional[str] = None
    due_date: Optional[str] = None
    effective_date: Optional[str] = None
    termination_summary: Optional[str] = None
    additional_terms: Optional[str] = None
    agreement_family: str = ""
    material_asks: List[str] = Field(default_factory=list)
    payment: Optional[Dict[str, Any]] = None
    clause_pack_seed: Optional[str] = Field(
        default=None,
        max_length=12_000,
        description="Deterministic clause-coverage hints from the client; prioritize when drafting.",
    )
    deterministic_intent_id: Optional[str] = Field(
        default=None,
        max_length=128,
        description="Client intent bucket id (e.g. logo_brand, web_presence).",
    )
    intent_contract: Optional[Dict[str, Any]] = Field(
        default=None,
        description="Optional LawDog Pro intent contract: expected agreement type, title terms, must-cover, avoid-misclassify, user fact summary (guidance, not a template).",
    )


class PremiumFullDraftRequest(BaseModel):
    intake_text: str = Field(..., min_length=1)
    context: Optional[PremiumFullDraftContext] = None
    user_gap_answers: Optional[str] = Field(default=None, max_length=32_000)
    """Client session generation id — idempotency / stale-response correlation (optional)."""
    agreement_generation_id: Optional[str] = Field(default=None, max_length=128)
    """Short intake fingerprint from client (optional)."""
    intake_fingerprint: Optional[str] = Field(default=None, max_length=64)
    """Persisted agreement id when retrying after checkout (optional)."""
    agreement_id: Optional[str] = Field(default=None, max_length=128)
    """
    When true, the client is asking for a second pass because the first full draft was too close
    in substance to a free / stitched outline. Uses CLAW_LLM_MODEL_PREMIUM_REGEN (or premium default).
    """
    similarity_regeneration: bool = False


class ExtractedParty(BaseModel):
    name: str = ""
    role: str = ""


class ExtractedPartyRole(BaseModel):
    party_name: str = ""
    role: str = ""


class PaymentMilestone(BaseModel):
    label: str = ""
    amount: Optional[str] = None
    percentage: Optional[str] = None
    trigger: Optional[str] = None


class RecurringSupportTerms(BaseModel):
    amount: Optional[str] = None
    cadence: Optional[str] = None
    renewal: Optional[str] = None


class AgreementPaymentTerms(BaseModel):
    total_amount: Optional[str] = None
    currency: Optional[str] = None
    milestones: List[PaymentMilestone] = Field(default_factory=list)
    recurring_support: Optional[RecurringSupportTerms] = None


class AgreementOwnershipTerms(BaseModel):
    deliverable_ownership: Optional[str] = None
    retained_materials: Optional[str] = None


class AgreementTerminationTerms(BaseModel):
    convenience_termination: Optional[bool] = None
    breach_termination: Optional[bool] = None
    notice_period: Optional[str] = None


class AgreementConfidentialityTerms(BaseModel):
    included: bool = False
    survival: Optional[str] = None


class AgreementNoticesTerms(BaseModel):
    method: Optional[str] = None


class AgreementSupportTerms(BaseModel):
    included: Optional[bool] = None
    standard: Optional[str] = None


class AgreementThirdPartyDependencyTerms(BaseModel):
    included: Optional[bool] = None
    uptime_disclaimer: Optional[bool] = None


class AgreementExtractedTerms(BaseModel):
    parties: List[ExtractedParty] = Field(default_factory=list)
    party_roles: List[ExtractedPartyRole] = Field(default_factory=list)
    governing_law: Optional[str] = None
    payment_terms: Optional[AgreementPaymentTerms] = None
    ownership_terms: Optional[AgreementOwnershipTerms] = None
    termination_terms: Optional[AgreementTerminationTerms] = None
    confidentiality: Optional[AgreementConfidentialityTerms] = None
    notices: Optional[AgreementNoticesTerms] = None
    support_terms: Optional[AgreementSupportTerms] = None
    third_party_dependency_terms: Optional[AgreementThirdPartyDependencyTerms] = None
    electronic_signatures: Optional[bool] = None


class AgreementAmbiguity(BaseModel):
    id: str = ""
    topic: str = ""
    description: str = ""
    severity: Literal["low", "medium", "high"] = "medium"
    source: Optional[str] = None


class AgreementConflict(BaseModel):
    id: str = ""
    topic: str = ""
    description: str = ""
    conflicting_values: List[str] = Field(default_factory=list)
    severity: Literal["low", "medium", "high"] = "medium"


class MissingMaterialTerm(BaseModel):
    id: str = ""
    topic: str = ""
    reason: str = ""
    severity: Literal["low", "medium", "high"] = "medium"


class RecommendedQuestion(BaseModel):
    id: str = ""
    topic: str = ""
    question: str = ""
    reason: str = ""
    priority: Literal["low", "medium", "high"] = "medium"


class AgreementQualityFlag(BaseModel):
    id: str = ""
    topic: str = ""
    description: str = ""
    severity: Literal["low", "medium", "high"] = "medium"


class AgreementIntelligence(BaseModel):
    extracted_terms: AgreementExtractedTerms = Field(default_factory=AgreementExtractedTerms)
    ambiguities: List[AgreementAmbiguity] = Field(default_factory=list)
    conflicts: List[AgreementConflict] = Field(default_factory=list)
    missing_material_terms: List[MissingMaterialTerm] = Field(default_factory=list)
    recommended_questions: List[RecommendedQuestion] = Field(default_factory=list)
    quality_flags: List[AgreementQualityFlag] = Field(default_factory=list)


class PremiumFullDraftResponse(BaseModel):
    title: str = ""
    agreement_family: str = ""
    document_text: str = ""
    authoritative_draft: str = ""
    agreement_intelligence: AgreementIntelligence = Field(default_factory=AgreementIntelligence)
    agreement_validation: Optional[AgreementValidationResult] = None
    server_full_document_text: str = ""
    server_repair_document_text: str = ""
    key_terms_found: List[str] = Field(default_factory=list)
    missing_material_info: List[str] = Field(default_factory=list)
    generation_outcome: Literal["ok", "needs_details", "degraded"] = "ok"
    schema_validation_reasons: List[str] = Field(default_factory=list)
    """When generation used a non-model fallback (degraded), machine-safe reason for UI + logs."""
    server_generation_failure_code: str = ""
    """Operator-safe short text; no secrets."""
    server_generation_failure_message: str = ""
    """False when degraded/empty output must not be treated as a successful Pro generation."""
    generation_ok: bool = True
    """True when the client may retry premium-full-draft without a new free draft."""
    retryable: bool = False


class PremiumFinalizationClarificationAnswer(BaseModel):
    question_id: Optional[str] = None
    question: str = Field(..., min_length=1, max_length=4000)
    answer: str = Field(..., min_length=1, max_length=12000)

    @field_validator("question_id", "question", "answer", mode="before")
    @classmethod
    def _strip_optional_text(cls, v: Any) -> Any:
        if isinstance(v, str):
            return v.strip()
        return v


class PremiumFinalizationRequest(BaseModel):
    original_intake: str = Field(..., min_length=1, max_length=120_000)
    first_draft: str = Field(..., min_length=1, max_length=200_000)
    agreement_intelligence: Optional[AgreementIntelligence] = None
    agreement_validation: Optional[AgreementValidationResult] = None
    clarification_answers: List[PremiumFinalizationClarificationAnswer] = Field(default_factory=list, max_length=50)
    force_finalize: bool = False

    @field_validator("original_intake", "first_draft", mode="before")
    @classmethod
    def _strip_required_text(cls, v: Any) -> Any:
        if isinstance(v, str):
            return v.strip()
        return v


def _classify_premium_full_draft_failure(exc: BaseException) -> tuple[str, str]:
    """
    Return (failure_code, safe_log_line) for premium full draft.
    Distinguishes auth, timeout, parse, airlock, empty output, unknown.
    """
    et = type(exc).__name__
    em = str(exc).lower()
    if isinstance(exc, RuntimeError) and "openai_api_key" in em:
        return "missing_openai_key", "missing_openai_key:RuntimeError"
    if isinstance(exc, ExternalAIBlockedError):
        br = getattr(exc, "block_reason", None) or "airlock"
        codes = getattr(exc, "policy_reason_codes", ()) or ()
        codes_s = ",".join(str(c) for c in codes) if codes else ""
        detail = f"airlock:{br}" + (f" policy_codes={codes_s}" if codes_s else "")
        return "airlock_blocked", detail
    try:
        from openai import APIConnectionError, APIStatusError, APITimeoutError, AuthenticationError, RateLimitError

        if isinstance(exc, AuthenticationError):
            return "openai_auth", f"openai_auth:{et}"
        if isinstance(exc, APIStatusError):
            sc = getattr(exc, "status_code", None)
            if sc == 401:
                return "openai_auth", "openai_http_401"
            if sc == 404:
                return "openai_model_not_found", "openai_http_404_model_or_path"
            if sc == 429:
                return "openai_rate_limit", "openai_http_429"
            if sc is not None and int(sc) >= 500:
                return "openai_server", f"openai_http_{sc}"
        if isinstance(exc, (APITimeoutError, TimeoutError)):
            return "openai_timeout", f"timeout:{et}"
        if isinstance(exc, APIConnectionError):
            return "openai_connection", f"connection:{et}"
        if isinstance(exc, RateLimitError):
            return "openai_rate_limit", f"rate_limit:{et}"
    except ImportError:
        pass
    if isinstance(exc, json.JSONDecodeError):
        return "json_parse", f"json_decode:{et}"
    if isinstance(exc, ValueError):
        if "invalid_json" in em or "invalid_json_object" in em:
            return "json_parse", f"value_invalid_json:{et}"
        if "empty_document" in em or em == "empty_document_text":
            return "empty_output", "empty_document_text"
        if "premium_dev_context" in em or "dev_context" in em:
            return "dev_context_leak", "dev_context_leak"
        if "too_large" in em:
            return "payload_limits", et
    if "timeout" in em or "timed out" in em or "read timed out" in em:
        return "openai_timeout", f"timeout_string:{et}"
    return "unknown", f"{et}"


_SUPPRESS_DEGRADED_FAKE_DOCUMENT_CODES: frozenset[str] = frozenset(
    {
        "airlock_blocked",
        "dev_context_leak",
    }
)


def _build_premium_full_draft_fallback_document(
    intake_s: str,
    ctx_dict: Optional[Dict[str, Any]],
    failure_code: str,
) -> str:
    """Structured preview from intake + context when the Pro model path is unavailable (no repeated filler clauses)."""
    title = ""
    if ctx_dict:
        title = str(ctx_dict.get("title") or "").strip()
    if not title:
        title = "Agreement"
    blob = build_free_reference_blob(intake_s, ctx_dict).strip()
    if len(blob) < 400:
        blob = f"{(intake_s or '').strip()}\n\n{blob}".strip()
    # Single neutral completion note — never emit many copies of the same generic “operative terms” line.
    completion = (
        "Complete operative commercial terms (scope, fees, liability, termination, and signatures) "
        "from the summary above; refine after internal or external review before signing."
    )
    return (
        f"# {title}\n\n"
        f"*We saved your Pro upgrade. The automated full pass was not available for this run ({failure_code}). "
        f"Below is a structured summary from your notes — you can edit freely. Use **Retry Pro draft** later for another full pass, or keep refining below.*\n\n"
        f"## Summary from your intake\n\n{blob}\n\n"
        f"## Commercial framework\n\n{completion}\n"
    )


def _premium_full_draft_degraded_response(
    *,
    intake_s: str,
    ctx_dict: Optional[Dict[str, Any]],
    failure_code: str,
    failure_message: str,
    primary_full: str = "",
    repair_body: str = "",
) -> PremiumFullDraftResponse:
    suppress_body = failure_code in _SUPPRESS_DEGRADED_FAKE_DOCUMENT_CODES
    doc = (
        ""
        if suppress_body
        else _build_premium_full_draft_fallback_document(intake_s, ctx_dict, failure_code)
    )
    fam = ""
    if ctx_dict:
        fam = str(ctx_dict.get("agreement_family") or "").strip()
    log.error(
        "premium_full_draft event=degraded_response failure_code=%s failure_message=%s doc_len=%s suppress_body=%s",
        failure_code,
        failure_message[:200],
        len(doc),
        int(suppress_body),
    )
    if suppress_body:
        log.warning(
            "[CLAW] premium generation blocked category=%s stage=model_path suppress_fallback_document=1",
            failure_code,
        )
    else:
        log.info("[CLAW] premium degraded fallback_document failure_code=%s", failure_code)
    srv_full = "" if suppress_body else (primary_full or "")
    srv_repair = "" if suppress_body else (repair_body or "")
    empty_intelligence = AgreementIntelligence()
    agreement_validation = _validate_and_log_premium_agreement_draft(
        authoritative_draft=doc,
        agreement_intelligence=empty_intelligence,
        original_intake=intake_s,
        stage=f"degraded:{failure_code}",
    )
    return PremiumFullDraftResponse(
        title=str((ctx_dict or {}).get("title") or "").strip() or "Agreement",
        agreement_family=fam,
        document_text=doc,
        authoritative_draft=doc,
        agreement_intelligence=empty_intelligence,
        agreement_validation=agreement_validation,
        server_full_document_text=srv_full,
        server_repair_document_text=srv_repair,
        key_terms_found=[],
        missing_material_info=[f"pro_model_unavailable:{failure_code}"],
        generation_outcome="degraded",
        schema_validation_reasons=(
            [f"fallback_suppressed:{failure_code}"] if suppress_body else [f"fallback:{failure_code}"]
        ),
        server_generation_failure_code=failure_code,
        server_generation_failure_message=failure_message,
        generation_ok=bool(doc.strip()),
        retryable=suppress_body,
    )


def _degraded_user_message_for_code(code: str) -> str:
    """Safe, non-technical copy for API clients (no secrets). Calm, trust-preserving; optional retry is secondary."""
    m: Dict[str, str] = {
        "missing_openai_key": "Your LawDog Pro agreement is ready for review. Your Pro access is saved; you can keep editing the text below. When the AI service is restored, **Retry Pro draft** is available if you want another automated pass.",
        "openai_model_not_found": "The configured AI model is not available for this API. Ask your admin to set CLAW_LLM_MODEL_PREMIUM to a valid model id, then use **Retry Pro draft**.",
        "openai_auth": "Your agreement is ready. You can refine any wording below, or use **Retry Pro draft** in a few minutes if you want a fresh pass.",
        "openai_rate_limit": "Your agreement is ready. You can refine any wording below, or **Retry Pro draft** shortly if you want another pass.",
        "openai_timeout": "Your agreement is ready. You can refine any wording below, or use **Retry Pro draft** in a few minutes if you want a fresh pass.",
        "openai_connection": "Your agreement is ready. You can refine any wording below, or **Retry Pro draft** after checking your network.",
        "openai_server": "Your agreement is ready. You can refine any wording below, or use **Retry Pro draft** in a few minutes if you want another pass.",
        "json_parse": "Your agreement is ready. You can refine any wording below, or try **Retry Pro draft** to regenerate the full pass.",
        "airlock_blocked": (
            "LawDog Pro could not complete the full drafting pass. Your upgrade is saved. "
            "Please retry, or continue editing the starter draft."
        ),
        "empty_output": "Your agreement is ready. You can refine any wording below, or use **Retry Pro draft** for another pass.",
        "dev_context_leak": "Your agreement is ready. You can refine any wording below, or use **Retry Pro draft** for a fresh pass.",
        "payload_limits": "Your agreement is ready. Try shortening the intake and using **Retry Pro draft**, or keep editing the text below.",
    }
    return m.get(
        code,
        "Your agreement is ready. You can refine any wording below. Your Pro purchase is intact — you can also use **Retry Pro draft** in a few minutes for another full pass if you like.",
    )


def _premium_full_draft_sanitize_wire_nested(obj: Any) -> Any:
    """Ensure all strings are UTF-8-encodable for JSON over HTTP/2 (lone surrogates → replacement)."""
    if isinstance(obj, str):
        try:
            obj.encode("utf-8")
            return obj
        except UnicodeEncodeError:
            return obj.encode("utf-8", errors="replace").decode("utf-8")
    if isinstance(obj, dict):
        return {str(k): _premium_full_draft_sanitize_wire_nested(v) for k, v in obj.items()}
    if isinstance(obj, list):
        return [_premium_full_draft_sanitize_wire_nested(v) for v in obj]
    return obj


def _premium_full_draft_model_to_wire_dict(model: PremiumFullDraftResponse) -> Dict[str, Any]:
    dumped = model.model_dump(mode="json")
    return _premium_full_draft_sanitize_wire_nested(dumped)


def _premium_full_draft_finalize_http_response(
    model: PremiumFullDraftResponse,
    *,
    intake_len: int,
    session_hint: str,
    server_timing: Optional[PaidProServerTiming] = None,
    request: Optional[Request] = None,
) -> Response:
    """
    Single JSON serialization to bytes — avoids streaming/partial frames and catches wire-unsafe text
    before any response bytes are committed.
    """
    serialize_started = time.perf_counter()
    try:
        wire = _premium_full_draft_model_to_wire_dict(model)
        raw = json.dumps(wire, ensure_ascii=False, separators=(",", ":"), allow_nan=False)
    except (TypeError, ValueError, UnicodeEncodeError) as exc:
        log.exception(
            "[premium-full-draft] event=failure stage=response_serialize status=503 code=premium_full_draft_response_serialization_failed "
            "exc_type=%s intake_len=%s session_hint=%s",
            type(exc).__name__,
            intake_len,
            session_hint,
        )
        err = {
            "detail": {
                "code": "premium_full_draft_response_serialization_failed",
                "message": type(exc).__name__,
                "stage": "response_serialize",
                "route": "premium-full-draft",
            }
        }
        err_raw = json.dumps(err, ensure_ascii=False, separators=(",", ":"))
        err_response = Response(
            status_code=503,
            content=err_raw.encode("utf-8"),
            media_type="application/json; charset=utf-8",
        )
        err_response = attach_cors_from_request(request, err_response)
        if request is not None:
            log_premium_full_draft_cors_proof(
                request, err_response, note="finalize_serialize_error_503"
            )
        return err_response
    body_bytes = raw.encode("utf-8")
    doc_len = len(wire.get("document_text") or "") if isinstance(wire.get("document_text"), str) else 0
    gen_ok = wire.get("generation_ok")
    retryable = wire.get("retryable") is True
    gen_out = str(wire.get("generation_outcome") or "")
    fail_code = str(wire.get("server_generation_failure_code") or "")
    status_code = 200
    if gen_ok is False and retryable and doc_len == 0 and gen_out == "degraded":
        status_code = 503
        log.error(
            "[premium-full-draft-empty-output-forbidden] status=503 failure_code=%s "
            "intake_len=%s session_hint=%s document_text_len=0 retryable=1",
            fail_code,
            intake_len,
            session_hint,
        )
    log.info(
        "[premium-full-draft] event=response_build status=%s intake_len=%s body_len=%s document_text_len=%s "
        "session_hint=%s generation_outcome=%s",
        status_code,
        intake_len,
        len(body_bytes),
        doc_len,
        session_hint,
        wire.get("generation_outcome"),
    )
    if server_timing is not None:
        server_timing.record(
            "backend_response_packaging",
            (time.perf_counter() - serialize_started) * 1000,
            bodyLen=len(body_bytes),
            documentTextLen=doc_len,
        )
        try:
            log.info(
                "[paid-pro-server-waterfall] %s",
                json.dumps(server_timing.to_wire(), ensure_ascii=False, default=str)[:12000],
            )
        except Exception:
            pass
    response = Response(
        status_code=status_code,
        content=body_bytes,
        media_type="application/json; charset=utf-8",
    )
    response = maybe_attach_server_timing_header(response, server_timing)
    response = attach_cors_from_request(request, response)
    if request is not None:
        log_premium_full_draft_cors_proof(
            request, response, note="finalize_http_response"
        )
    return response


class PremiumMissingFactsRequest(BaseModel):
    """Pre–full-draft: which high-value commercial facts are still open?"""

    intake_text: str = Field(..., min_length=1)
    context: Optional[PremiumFullDraftContext] = None


class PremiumMissingFactsResponse(BaseModel):
    questions: List[str] = Field(default_factory=list)


def _normalize_premium_missing_facts_result(raw: Dict[str, Any]) -> PremiumMissingFactsResponse:
    out: List[str] = []
    q = raw.get("questions")
    if isinstance(q, list):
        for x in q:
            s = str(x).strip() if x is not None else ""
            if s and len(out) < 5:
                s = s[:800] if len(s) > 800 else s
                out.append(s)
    return PremiumMissingFactsResponse(questions=out)


def _premium_missing_facts_system_prompt() -> str:
    return (
        "You help users finish a business agreement in CLAW before a final full draft is generated.\n"
        "You receive: the user's original intake, and (if present) structured `context` (free-path draft + extract: "
        "title, parties, purpose, payment, dates, material_asks, agreement_family, etc.).\n"
        "Task: list up to **5** high-value, **material** missing facts that would meaningfully change or sharpen the contract if known. "
        "Ask only about things a reasonable business person would need to lock before sending: e.g. exact payment amounts/cadence if completely absent, "
        "governing law/venue if TBD, key deliverable dates, exclusivity scope, IP ownership if contested, who approves ad spend, commission base, etc.\n"
        "Rules:\n"
        "- Do not ask for boilerplate the draft can already handle with neutral defaults.\n"
        "- No duplicate questions; no legal advice; not a law-firm voice.\n"
        "- If the materials are already specific enough to produce a strong first draft, return an **empty** questions array.\n"
        "Output ONLY valid JSON, no markdown, with the exact key: { \"questions\": string array } (0–5 items).\n"
    )


def _detect_premium_scenario_category(intake: str, agreement_family: str = "") -> Tuple[str, List[str]]:
    """
    Heuristic scenario bucket for premium full-draft routing.
    Category IDs must stay aligned with frontend `premiumScenarioCategory.ts`.
    """
    low = (intake or "").strip().lower()
    fam_low = (agreement_family or "").strip().lower()
    signals: List[str] = []

    def push(tok: str) -> None:
        if tok not in signals:
            signals.append(tok)

    if len(low) < 24:
        return "custom_mixed", ["short_intake"]

    if re.search(
        r"\b(influencer|ugc|creator|tiktok|instagram|youtube|podcast\s+sponsor|brand\s+deal|"
        r"whitelisting|paid\s+post|sponsorship)\b",
        low,
    ):
        push("creator_influencer")
        return "business_commercial", signals

    if re.search(r"\b(saas|subscription|software\s+as\s+a\s+service|api\s+access|platform\s+terms)\b", low):
        push("saas_platform")
        return "business_commercial", signals

    employment = bool(
        re.search(
            r"\b(employment|at-will|at\s+will|w-2|w2|salary|hourly\s+wage|employee\s+handbook|job\s+offer|position\s+title|"
            r"benefits\s+package|severance\s+package|non-?compete\s*\(employee|work\s+for\s+the\s+company\s+as\s+an\s+employee)\b",
            low,
        )
    ) or (bool(re.search(r"\b(employer|employee)\b", low)) and bool(re.search(r"\b(salary|wages?|pto|vacation\s+days|performance\s+review)\b", low)))
    if employment:
        push("employment")
        return "employment", signals

    loan = bool(
        re.search(
            r"\b(promissory\s+note|personal\s+loan|loan\s+to|lend\s+\$|lend\s+money|borrow\s+\$|iou|installment\s+loan|"
            r"principal\s+and\s+interest|\bapr\b|repayment\s+schedule|lender|borrower)\b",
            low,
        )
    ) or (bool(re.search(r"\bloan\b", low)) and bool(re.search(r"\b(repay|interest|principal|installment)\b", low)))
    if loan:
        push("loan")
        return "loan_payment", signals

    settlement = bool(
        re.search(
            r"\b(settlement\s+agreement|mutual\s+release|release\s+of\s+claims|dispute\s+is\s+settled|full\s+and\s+final\s+settlement|"
            r"confidential\s+settlement)\b",
            low,
        )
    ) or (
        bool(re.search(r"\brelease\b", low))
        and bool(re.search(r"\b(claims?|disputes?|liabilit)\b", low))
        and bool(re.search(r"\b(settle|resolved|dismiss)\b", low))
    )
    if settlement:
        push("settlement")
        return "settlement_dispute", signals

    prop = bool(
        re.search(
            r"\b(roommate|sublet|sub-lease|security\s+deposit|landlord|tenant|lessor|lessee|rental\s+agreement|lease\s+agreement|monthly\s+rent)\b",
            low,
        )
    ) or (bool(re.search(r"\blease\b", low)) and bool(re.search(r"\b(rent|premises|unit\s+at)\b", low))) or bool(
        re.search(r"\b(hoa|homeowners\s+association)\b", low)
    )
    if prop:
        push("property")
        return "property_roommate", signals

    family = bool(
        re.search(
            r"\b(spouse|divorce|custody|visitation|prenup|prenuptial|family\s+loan|gift\s+to|between\s+family|parent\s+and\s+child|"
            r"sibling|caregiving|elder\s+care)\b",
            low,
        )
    ) or (bool(re.search(r"\b(personal|family)\b", low)) and bool(re.search(r"\b(trust|care|support)\b", low)))
    if family and not re.search(r"\b(llc|inc\.?|corp|b2b|services\s+agreement|invoice)\b", low, re.I):
        push("family")
        return "family_personal", signals

    freelancer = bool(
        re.search(
            r"\b(1099|independent\s+contractor|freelance|consulting\s+agreement|statement\s+of\s+work|\bsow\b|retainer\s+fee|"
            r"hourly\s+rate\s+for\s+services|deliverables?\b)",
            low,
        )
    ) or ("contractor" in fam_low or "consulting" in fam_low or "services" in fam_low)
    if freelancer and not re.search(r"\b(full-?time\s+employee|w-2)\b", low):
        push("freelancer")
        return "freelancer_service", signals

    business = bool(
        re.search(
            r"\b(b2b|vendor|supplier|saas|enterprise|master\s+service|msa|commission\s+structure|referral\s+partner|agency\s+of\s+record|"
            r"llc\s+and\s+llc|corporation)\b",
            low,
        )
    ) or any(
        x in fam_low for x in ("generic_business", "services_agreement", "nda", "partnership")
    )
    if business:
        push("business")
        return "business_commercial", signals

    push("mixed")
    return "custom_mixed", signals


def _normalize_premium_full_draft_result(raw: Dict[str, Any]) -> PremiumFullDraftResponse:
    title = str(raw.get("title") or "").strip()
    fam = str(raw.get("agreement_family") or "").strip()
    authoritative = str(raw.get("authoritative_draft") or "").strip()
    doc = authoritative or str(raw.get("document_text") or "").strip()
    if len(doc) > 200_000:
        doc = doc[:200_000]
    intel_raw = raw.get("agreement_intelligence")
    intelligence_parse_ok = isinstance(intel_raw, dict)
    if intelligence_parse_ok:
        try:
            agreement_intelligence = AgreementIntelligence.model_validate(intel_raw)
        except Exception as exc:
            intelligence_parse_ok = False
            agreement_intelligence = AgreementIntelligence()
            log.warning(
                "[agreement-intelligence] event=parse_failure exc_type=%s keys=%s",
                type(exc).__name__,
                sorted(str(k) for k in intel_raw.keys())[:24] if isinstance(intel_raw, dict) else [],
            )
    else:
        agreement_intelligence = AgreementIntelligence()
        if intel_raw is not None:
            log.warning(
                "[agreement-intelligence] event=parse_failure exc_type=non_object raw_type=%s",
                type(intel_raw).__name__,
            )
    ktf: List[str] = []
    ma = raw.get("key_terms_found")
    if isinstance(ma, list):
        for x in ma:
            s = str(x).strip() if x is not None else ""
            if s and len(ktf) < 48:
                ktf.append(s)
    miss: List[str] = []
    mm = raw.get("missing_material_info")
    if isinstance(mm, list):
        for x in mm:
            s = str(x).strip() if x is not None else ""
            if s and len(miss) < 32:
                miss.append(s)
    return PremiumFullDraftResponse(
        title=title,
        agreement_family=fam,
        document_text=doc,
        authoritative_draft=doc,
        agreement_intelligence=agreement_intelligence,
        key_terms_found=ktf,
        missing_material_info=miss,
    )


def _log_agreement_intelligence_summary(
    intelligence: AgreementIntelligence,
    *,
    stage: str,
    elapsed_ms: Optional[float] = None,
) -> None:
    terms = intelligence.extracted_terms
    payment = terms.payment_terms
    log.info(
        "[agreement-intelligence] event=extracted stage=%s elapsed_ms=%s parties=%s governing_law=%s "
        "payment_total=%s milestones=%s ambiguities=%s conflicts=%s missing_terms=%s recommended_questions=%s quality_flags=%s",
        stage,
        round(elapsed_ms, 2) if elapsed_ms is not None else None,
        len(terms.parties),
        1 if terms.governing_law else 0,
        1 if payment and payment.total_amount else 0,
        len(payment.milestones) if payment else 0,
        len(intelligence.ambiguities),
        len(intelligence.conflicts),
        len(intelligence.missing_material_terms),
        len(intelligence.recommended_questions),
        len(intelligence.quality_flags),
    )
    if intelligence.ambiguities:
        log.info(
            "[agreement-intelligence] event=ambiguities stage=%s items=%s",
            stage,
            [
                {
                    "id": a.id,
                    "topic": a.topic,
                    "severity": a.severity,
                }
                for a in intelligence.ambiguities[:8]
            ],
        )
    if intelligence.conflicts:
        log.info(
            "[agreement-intelligence] event=conflicts stage=%s items=%s",
            stage,
            [
                {
                    "id": c.id,
                    "topic": c.topic,
                    "severity": c.severity,
                    "values": c.conflicting_values[:4],
                }
                for c in intelligence.conflicts[:8]
            ],
        )
    if intelligence.missing_material_terms:
        log.info(
            "[agreement-intelligence] event=missing_material_terms stage=%s items=%s",
            stage,
            [
                {
                    "id": m.id,
                    "topic": m.topic,
                    "severity": m.severity,
                }
                for m in intelligence.missing_material_terms[:8]
            ],
        )
    if intelligence.recommended_questions:
        log.info(
            "[agreement-intelligence] event=recommended_questions stage=%s items=%s",
            stage,
            [
                {
                    "id": q.id,
                    "topic": q.topic,
                    "priority": q.priority,
                }
                for q in intelligence.recommended_questions[:8]
            ],
        )


def _validate_and_log_premium_agreement_draft(
    *,
    authoritative_draft: str,
    agreement_intelligence: AgreementIntelligence,
    original_intake: str,
    stage: str,
) -> AgreementValidationResult:
    result = validatePremiumAgreementDraft(
        authoritativeDraft=authoritative_draft,
        agreementIntelligence=agreement_intelligence,
        originalIntake=original_intake,
    )
    log.info(
        "[agreement-validation] event=result stage=%s passed=%s failure_count=%s warning_count=%s failure_codes=%s",
        stage,
        int(result.passed),
        result.summary.failure_count,
        result.summary.warning_count,
        [f.code for f in result.failures[:16]],
    )
    return result


def _premium_full_draft_system_prompt() -> str:
    return (
        "You are a careful agreements drafter for real people and small businesses. Draft a complete, counterpart-ready agreement. "
        "Premium means **better fit, clearer structure, and smarter protections for the stated situation** — **not** longer generic boilerplate, "
        "stack-and-dump templates, or unrelated enterprise packs.\n"
        "The reader should feel: “this is a real agreement for my situation, not a rearranged outline.” The draft should feel complete, "
        "fair, and appropriate for the deal type—without padding for its own sake.\n"
        "**Generation intelligence brief (if present):** When `generation_intelligence_brief` is in the user JSON, treat "
        "`situation_line`, `tone_directive`, `must_address`, and `contradiction_notes` as binding routing for this run. "
        "Open with a recital that names **this** deal (who, relationship, commercial purpose) in plain language—not a context-free "
        "“parties desire to set forth” opener.\n"
        "**Contradictions:** If `contradiction_notes` is non-empty, choose **one** commercially reasonable interpretation, draft "
        "operative terms accordingly, and put unresolved forks in `missing_material_info` as short decisions—never encode both "
        "sides as binding law.\n"
        "**key_terms_found:** Use **specific** labels when facts allow (e.g. “$5k fixed fee, Net 30” not only “Payment terms”). "
        "8–16 items covering economics, scope, risk allocation, and lifecycle.\n"
        "**Signals `creator_influencer` / `saas_platform`:** Calibrate to creator usage/deliverables/FTC-style disclosure hooks or "
        "SaaS subscription/data/support/SLA patterns respectively—omit unrelated enterprise MSA packs.\n"
        "**Material LawDog Pro bar (not optional for this product):** The output must read as a **signed-ready** commercial agreement, not a reshuffled free outline. "
        "Use the **actual party names and roles** from the intake and `context.parties` (extract from raw text if needed; do not leave generic “Party A/B” when names are knowable). "
        "When payment, fees, or economics are stated, give a **clear numbered payment / compensation section** (e.g. 1. 2. 3. or 5.1 5.2) so terms are scannable. "
        "Include where fit for the fact pattern: **scope** (deliverables and out-of-scope), **IP and work product ownership**, **confidentiality** (or mutual NDA-style duties), "
        "**revisions or change requests** (e.g. cap on included revision rounds and how extras are approved), **termination** (convenience, cause, effect, return of materials), "
        "**governing law and venue** from the user’s **explicit** state or country only — if the user names a US state (e.g. **Oklahoma**), do **not** substitute another state (e.g. Delaware) unless the materials clearly say so. "
        "**Notices** (email and optional mail), **electronic signatures and counterparts**, and a complete **signature block** for each party with name, title, and date lines at the end.\n"
        "**Deterministic premium intent skeleton (if present in the user JSON):** When `deterministic_premium_intent_skeleton` is present, it is the **category-native spine** "
        "for this run (section topics, title rules, misroute warnings). Draft the full `document_text` so it clearly follows that skeleton in substance and headings before any "
        "optional generic provisions. The skeleton is **not** optional guidance when present — it fixes the deal type before free-form prose.\n"
        "**Deterministic product mapping (if present in `context`):** When `context` includes a non-empty `clause_pack_seed` and/or `deterministic_intent_id`, "
        "treat the provided `context.title` as the intended document / H1 title unless the raw intake *clearly* points to a different deal type, and use "
        "`clause_pack_seed` as a prioritized *coverage* checklist (operative clauses, not a literal paste — weave what fits; omit what does not).\n"
        "When `context.intent_contract` is present, use it for **agreement type fit, title expectations, must-cover operative topics, and wrong-category avoidance**; "
        "it is product routing guidance, not a form to paste or a requirement to add fictional facts.\n"
        "**Scenario routing (required):** The user JSON may include `scenario_category` and `scenario_category_signals` from deterministic routing. "
        "Calibrate the entire draft to that scenario. For `family_personal`, `property_roommate`, `employment`, `loan_payment`, and `settlement_dispute`, "
        "use a humane, plain-English tone suited to individuals. For `business_commercial`, `freelancer_service`, and `custom_mixed`, stay professional but still "
        "**omit** deep NDA packs, non-solicit, reverse-engineering bans, dense fee schedules, IP assignment, or arbitration/mediation stacks **unless** the intake "
        "or `user_gap_answers` explicitly calls for them.\n"
        "**Clause relevance:** Include **only** provisions a reasonable reader would tie to this fact pattern. Do not paste irrelevant vendor/enterprise boilerplate.\n"
        "**Length discipline:** Default to **6–12 major sections** (merge thin topics under clear headings). Add sections only when complexity truly requires it.\n"
        "**Final silent check:** Would a normal user feel *this was made for my situation*? If it reads generic or padded, revise before emitting JSON.\n"
        "**Universal drafting methodology (behave as a universal drafter, not a form filler or outline generator):**\n"
        "1) **Posture / stance:** At the start, fix an internal drafting posture. Default: **balanced commercial**—practical, even-handed, "
        "suitable for arms-length B2B or business counterparties. Adjust only when the fact pattern or standard market norms for that "
        "deal type clearly warrant asymmetry (e.g., heavy confidentiality, sponsor/influencer dynamics, or fiduciary-sounding administration); "
        "if you lean asymmetric, do so in a defensible, commercially explicable way.\n"
        "2) **Missing facts (do not collapse):** Never shrink to bullets, a skeleton, or a ‘starter’ feel because some amounts or dates are unknown. "
        "Use **neutral, law-style defaults** where it is safe (schedules, payment mechanics, net-day placeholders, TBD governing law, reasonable cure) "
        "and keep the narrative flowing as a real agreement, not a questionnaire.\n"
        "3) **Clause selection (fact pattern, not bloat):** Provisions below are a menu of **candidates**—weave in what fits; omit what does not. "
        "Avoid irrelevant templates that signal ‘generic M&A’ or unrelated deal types. Match tone and depth to the industry and risk the user actually described.\n"
        "4) **Headings and structure (deal-shaped, not slot-shaped):** Organize the document with **headings and article titles that name this deal** "
        "(e.g. deliverables, commissions, ad spend, data, or trustee duties) rather than five anonymous buckets. The numbered checklist in this prompt is a "
        "coverage map, not a required title sheet—reorder, merge, or split for clarity, but do not turn it into a bare ‘Scope / Payment / Term’ shell.\n"
        "5) **User asks (traceability):** Every **material** term the user wanted must appear in **fully drafted operative text** (not only recitals or a vague ‘noted’ line). "
        "If the intake is messy, you still disambiguate and carry each ask into a specific clause or unmissable section.\n"
        "6) **Completeness vs outline:** The `document_text` must be a **complete first-pass agreement** in full sentences and subsections, not a numbered outline, table of "
        "contents, “TBD in Schedule A” stub, or preview teaser. A reader can sign a version like this in principle after filling remaining specifics.\n"
        "7) **Opening and ordering:** After parties/recitals, lead with the issues that matter most for the chosen **agreement family** and risk (e.g. NDA/ confidentiality "
        "first; referral/commission first; long-running services: scope and change control early; data-heavy: privacy/usage and security; governance-heavy: roles and decision rights). "
        "Order sections for reader sense, not fixed boilerplate order.\n"
        "8) **Internal consistency:** Re-read your own draft: **roles** (who is client, provider, payor) must align across every section; **payment, approvals, and termination** "
        "must not conflict (e.g. no cure period that contradicts immediate termination the user required; no fee structure that changes mid-paragraph). Fix contradictions before output.\n"
        "9) **Specificity register:** Be **concrete and mirror the user’s nouns, numbers, and timelines** when the intake is specific. When the facts are only abstract, use "
        "professional, neutral, commercially common language and placeholders—**do not invent** granular facts, regulated outcomes, or party-specific commitments.\n"
        "10) **Internal completeness audit (before you output JSON):** Silently check: (a) structural completeness for this family; (b) all material user asks are present; "
        "(c) no cross-section conflicts; (d) the body reads as a real agreement, not a checklist. Revise the draft in your head until it passes, then return only the JSON.\n"
        "Rules:\n"
        "- If the user JSON includes `user_gap_answers`, treat that string as the user’s direct answers to pre-finalization questions; "
        "incorporate them as material facts in the agreement where not contradictory to the rest of the materials.\n"
        "- Be practical, readable, and well structured. Do not over-lawyer. No legal advice disclaimers in the body text.\n"
        "- Do not fabricate specific dollar amounts, dates, or party legal names that are not in the materials. "
        "If economics are unclear, use neutral placeholder phrasing (e.g., fees, invoicing, net days) without inventing numbers.\n"
        "- Preserve every material user ask from the intake in operative language somewhere in the agreement. "
        "If the user asked for ownership, confidentiality, reporting, spend approval, term/termination, FTC or industry rules, make explicit sections for those asks.\n"
        "- Add restrictive covenants (non-compete, exclusivity, non-circumvent) only if requested in the materials or "
        "clearly central to the stated deal. Otherwise omit or use light, mutual commercial boilerplate if appropriate.\n"
        "- **HOUSE STYLE — SECTION NUMBERING (document structure only):** Do not create subsection numbering unless at least "
        "two sibling subsections exist within the same section. Avoid orphan subsections. A section containing only one provision "
        "should render as a main heading followed by body paragraph text—not a lone N.1 label (e.g. use "
        "\"7. Governing Law\\n\\nThis Agreement shall be governed by…\" rather than \"7. Governing Law\\n\\n7.1 This Agreement…\"). "
        "Use subsection numbering only when multiple sibling subsections are present (e.g. 7.1, 7.2, 7.3). "
        "This instruction applies to document structure only; do not change substantive legal content.\n"
        "- Use clear numbering (1., 1.1, (a)…) and professional headings consistent with the house style above.\n"
        "- **Contextual standard commercial terms:** Use the intake and any structured `agreement_family` / deal labels in the JSON `context` "
        "to decide which standard clauses to weave in. Do not include every item below in every document—**select** based on "
        "relevance (e.g., SaaS, agency, referral, influencer, contractor, NDA, partnership, services retainer, trustee/fiduciary-style duties).\n"
        "- **Partnership/co-owner/family-business routing:** If the intake signals cofounders/partners/family operation/joint ownership "
        "(e.g., coffee cart, profit split, shared business control), prioritize: ownership allocations, capital contributions, "
        "expense controls, profit distribution cadence and definitions, authority controls, work expectations, exit/buyout mechanics, "
        "deadlock resolution, and signatures. De-prioritize generic agency/marketing/enterprise NDA boilerplate unless directly relevant.\n"
        "  • **Authority and capacity:** When parties are companies or the deal implies signing on behalf of an organization, a short line that each signer/representative "
        "has authority to bind the party is normal; skip or slim down for two individuals only if clearly a consumer-style simple deal.\n"
        "  • **Cooperation in good faith:** A concise mutual ‘reasonable cooperation’ obligation where ongoing performance, approvals, or shared execution matters "
        "(services, campaigns, data access, handoffs). Omit for bare payment-only or one-page NDAs if it would bloat needlessly.\n"
        "  • **Force majeure:** Include for ongoing performance, deliverables, or time-sensitive services; often shortened or omitted for one-off commission/referral settlement style deals.\n"
        "  • **Breach; notice; cure period:** For material defaults, a commercially standard pattern: written notice, a **reasonable cure period** (e.g. 10–30 days, stated as a range or placeholder) "
        "for curable breaches before termination. Align cure window with the deal (longer for enterprise; shorter for time-critical pilots only if the facts support it).\n"
        "  • **Independent contractor / not employment / no partnership:** Add when a person or small firm is performing services, referral, or creative work; distinguish from employment and joint venture "
        "unless the facts clearly require partnership language.\n"
        "  • **Taxes and withholding:** Clarify that each party bears its own taxes unless a specific withholding rule appears in the materials; for contractor/referral/commission structures, state plainly "
        "who issues forms (e.g. 1099-style) **without** inventing tax advice—use neutral 'as required by law' phrasing when specifics are unknown.\n"
        "  • **Confidentiality: customary exceptions** (e.g. public domain, independently developed, rightfully received from a third party, required by law with notice where practicable) when you include a confidentiality section—keep tight, not a treatise.\n"
        "  • **Injunctive relief (and similar):** For confidentiality, data misuse, or IP, note that the parties understand **money damages may be inadequate** and that **equitable relief** (e.g. injunction) may be "
        "sought, subject to the forum—keep commercially standard and mutual where appropriate; omit if the deal is purely about payment and no secrets/IP.\n"
        "  • **Survival:** State that provisions which by nature should survive (confidentiality, payment of accrued fees, IP, liability limits as stated, dispute resolution) survive termination. Tailor the list to what you actually included.\n"
        "  • **Order of precedence:** If the facts reference exhibits, SOW, schedules, or prior terms, a short ‘conflicts: this Agreement controls unless the exhibit expressly states otherwise’ (or the inverse if the user’s structure requires it). Omit if a single self-contained document.\n"
        "  • **Electronic / digital signatures and counterparts:** In miscellaneous or signatures, confirm that e-signatures and digital execution are **valid** and that counterparts together form one agreement.\n"
        "  • **Notices (digital):** In the notices section, allow delivery by email to designated addresses (and optional mail) with direction to update addresses; match business reality for remote parties.\n"
        "  • **Data / security (reasonable safeguards):** When the deal involves PII, customer data, ad platforms, or credentials, add brief mutual obligations to use **reasonable** administrative/technical **safeguards** and limit use to the agreement’s purpose—**not** full enterprise SOC2-style detail unless the intake demands it.\n"
        "Required coverage map (pick **only** items that fit this deal; merge into 6–12 sections where possible; omit labels that do not apply):\n"
        "1) Title\n"
        "2) Parties, recitals / background, purpose\n"
        "3) Definitions (only if they reduce repetition)\n"
        "4) Scope, responsibilities, deliverables, standards\n"
        "5) Compensation, fees, invoicing, payment terms, expenses/reimbursements\n"
        "6) Approvals, decision rights, spend or budget controls (if relevant)\n"
        "7) IP, work product, accounts, data, and background IP (as applicable)\n"
        "8) Confidentiality, privacy, and NDA elements as needed (include exceptions and remedies when confidentiality is material)\n"
        "9) Representations and warranties — light, mutual commercial standard, plus authority to enter and perform where context supports it\n"
        "10) Compliance: laws, regulatory / FTC / industry requirements when relevant to the user’s fact pattern\n"
        "11) Reporting, records, audit (if relevant)\n"
        "12) Subcontracting, assignment, and change of control (reasonable defaults)\n"
        "13) Exclusivity / non-solicit / non-circumvent / non-compete (only as requested or plainly relevant)\n"
        "14) Term, renewal, and effective date mechanics\n"
        "15) Suspension, termination, **material breach (notice and cure)**, **force majeure** (when ongoing duties exist), and transition or wind-down (as applicable)\n"
        "16) Indemnity, limitations of liability, and liability cap — commercially reasonable, mutual by default\n"
        "17) Dispute resolution, governing law, exclusive venue (use jurisdiction provided or neutral placeholder if TBD)\n"
        "18) **Notices** (including **email** as a permitted method where appropriate)\n"
        "19) Miscellaneous: **entire agreement; order of precedence (if needed);** amendments; **survival;** severability; waiver; **counterparts;** **electronic signatures** and valid digital execution\n"
        "20) Signature blocks for all parties and dates\n"
        "Output ONLY a single JSON object (no markdown, no code fences) with EXACT keys:\n"
        '{ "title": string, "agreement_family": string, "authoritative_draft": string, "agreement_intelligence": object, "key_terms_found": string array, "missing_material_info": string array }\n'
        "- `authoritative_draft` is the full agreement in plain text, not a summary. This must be a complete first-pass draft a user can review. "
        "For backward compatibility you may also include `document_text` with the exact same value, but `authoritative_draft` is required.\n"
        "- `key_terms_found`: 6–20 short labels for the major commercial points you actually included (including key standard protections you added).\n"
        "- `missing_material_info`: only material items still unknown after a fair read of the user materials, else [].\n"
        "- `agreement_family`: short human label of deal type, e.g. 'Marketing services retainer' or 'Referral commission'—use it to calibrate which standard clauses to emphasize.\n"
        "- `agreement_intelligence` is internal LawDog system intelligence. Extract only facts clearly supplied or reasonably inferable from the intake. "
        "Do not invent governing law, payment amounts, milestone schedules, signers, notice methods, ownership transfers, support obligations, or uptime promises. "
        "If uncertain, leave the extracted field null/empty and use `ambiguities` or `missing_material_terms`.\n"
        "- `agreement_intelligence` shape: { \"extracted_terms\": { \"parties\": [{\"name\": string, \"role\": string}], \"party_roles\": [{\"party_name\": string, \"role\": string}], "
        "\"governing_law\": string|null, \"payment_terms\": {\"total_amount\": string|null, \"currency\": string|null, \"milestones\": [{\"label\": string, \"amount\": string|null, \"percentage\": string|null, \"trigger\": string|null}], "
        "\"recurring_support\": {\"amount\": string|null, \"cadence\": string|null, \"renewal\": string|null}|null}|null, "
        "\"ownership_terms\": {\"deliverable_ownership\": string|null, \"retained_materials\": string|null}|null, "
        "\"termination_terms\": {\"convenience_termination\": boolean|null, \"breach_termination\": boolean|null, \"notice_period\": string|null}|null, "
        "\"confidentiality\": {\"included\": boolean, \"survival\": string|null}|null, \"notices\": {\"method\": string|null}|null, "
        "\"support_terms\": {\"included\": boolean|null, \"standard\": string|null}|null, "
        "\"third_party_dependency_terms\": {\"included\": boolean|null, \"uptime_disclaimer\": boolean|null}|null, \"electronic_signatures\": boolean|null }, "
        "\"ambiguities\": [{\"id\": string, \"topic\": string, \"description\": string, \"severity\": \"low\"|\"medium\"|\"high\", \"source\": string|null}], "
        "\"conflicts\": [{\"id\": string, \"topic\": string, \"description\": string, \"conflicting_values\": string[], \"severity\": \"low\"|\"medium\"|\"high\"}], "
        "\"missing_material_terms\": [{\"id\": string, \"topic\": string, \"reason\": string, \"severity\": \"low\"|\"medium\"|\"high\"}], "
        "\"recommended_questions\": [{\"id\": string, \"topic\": string, \"question\": string, \"reason\": string, \"priority\": \"low\"|\"medium\"|\"high\"}], "
        "\"quality_flags\": [{\"id\": string, \"topic\": string, \"description\": string, \"severity\": \"low\"|\"medium\"|\"high\"}] }.\n"
        "- Recommended questions must be sparse and high-signal. Duplicate or unnecessary clarification questions are harmful UX. "
        "Do NOT ask about governing law, payment structure, or ownership when already clearly supplied. Ask only if information is genuinely missing, materially ambiguous, conflicting, "
        "or would substantially improve specificity/enforceability (for example unclear venue, undefined acceptance criteria, ambiguous support renewal, conflicting milestone schedules).\n"
    )


class PremiumAgreementReviewRequest(BaseModel):
    """Lightweight post-pass: clarity + completeness nudges for the generated premium body."""

    intake_text: str = Field(..., min_length=1)
    document_text: str = Field(..., min_length=1)
    context: Optional[PremiumFullDraftContext] = None


class PremiumAgreementReviewResponse(BaseModel):
    strengths: List[str] = Field(default_factory=list)
    missing_or_weak_terms: List[str] = Field(default_factory=list)
    questions_for_user: List[str] = Field(default_factory=list)
    suggested_clause_upgrades: List[str] = Field(default_factory=list)
    priority_score: int = 0  # 0 – 100, higher = more to tighten before send


def _str_list_capped(raw: Any, cap: int) -> List[str]:
    out: List[str] = []
    if not isinstance(raw, list):
        return out
    for x in raw:
        s = str(x).strip() if x is not None else ""
        if s and len(out) < cap:
            s = s[:2000] if len(s) > 2000 else s
            out.append(s)
    return out


def _normalize_premium_agreement_review_result(raw: Dict[str, Any]) -> PremiumAgreementReviewResponse:
    sc = raw.get("priority_score")
    try:
        p = int(float(sc)) if sc is not None else 50
    except (TypeError, ValueError):
        p = 50
    if p < 0:
        p = 0
    if p > 100:
        p = 100
    return PremiumAgreementReviewResponse(
        strengths=_str_list_capped(raw.get("strengths"), 10),
        missing_or_weak_terms=_str_list_capped(raw.get("missing_or_weak_terms"), 10),
        questions_for_user=_str_list_capped(raw.get("questions_for_user"), 5),
        suggested_clause_upgrades=_str_list_capped(raw.get("suggested_clause_upgrades"), 8),
        priority_score=p,
    )


def _premium_agreement_review_system_prompt() -> str:
    return (
        "You are a crisp commercial editor helping someone finalize an agreement in CLAW (a product, not a law firm). "
        "You perform a **light review pass** on a draft the user is about to send. You do **not** rewrite the full agreement, "
        "re-state every clause, or use stiff legalese, firm branding, or caveated disclaimer tone.\n"
        "Objectives: improve perceived clarity and completeness, flag only **material** gaps or soft spots, and help the user "
        "self-correct. Optimize for a fast, calm desktop or mobile read—short lines, scannable, actionable.\n"
        "Rules:\n"
        "- `strengths`: 2–5 bullets max—what already works (business fit, key protections present, clear economics structure). "
        "Praise is specific, not generic.\n"
        "- `missing_or_weak_terms`: 0–6 items—only if something **material** is missing, under-specified, or internally inconsistent with the intake. "
        "No nitpicks. Skip filler like “could add a definitions section” unless the deal truly needs it.\n"
        "- `questions_for_user`: 0–4 high-value questions only—decisions the user (or their counterparty) should confirm before sending, "
        "e.g. jurisdiction, payment trigger, exclusivity scope, data retention. If nothing important is open, return [].\n"
        "- `suggested_clause_upgrades`: 0–5 short, concrete nudges—e.g. “Tie the cure period in Section X to the 30 days mentioned in the intake” "
        "or “Add a one-line cap reference if liability is open-ended for ad spend.” **Do not** paste a full replacement clause. "
        "No multi-paragraph rewrites.\n"
        "- `priority_score` integer 0–100: higher means more issues worth addressing before send; low 20s = light polish only; 80+ = meaningful gaps. "
        "If the draft is solid, keep score under 40.\n"
        "Output ONLY valid JSON, no markdown fences, with **exact** keys: "
        '{ "strengths": string array, "missing_or_weak_terms": string array, "questions_for_user": string array, '
        '"suggested_clause_upgrades": string array, "priority_score": number }.\n'
    )


PremiumRefineAction = Literal["update", "ask_missing", "ready"]


class PremiumRefineRequest(BaseModel):
    """Post-draft smart refinement: apply edits, surface gaps, or assess readiness (premium surface)."""

    current_document_text: str = Field(..., min_length=1)
    intake_text: str = Field(..., min_length=1)
    user_refinement_prompt: str = ""
    action: PremiumRefineAction = "update"
    surgical_preserve_retry: bool = Field(
        default=False,
        description="Second-pass: stronger preserve-full-document prompt + heading context (surgical QA).",
    )


class PremiumRefineResponse(BaseModel):
    updated_document_text: str = ""
    summary_changes: List[str] = Field(default_factory=list)
    readiness_score: int = Field(0, ge=0, le=100)
    suggested_next_step: Literal["edit", "review", "send"] = "review"


def _premium_refine_next_step_literal(v: Any) -> Literal["edit", "review", "send"]:
    s = str(v or "review").strip().lower()
    if s == "edit":
        return "edit"
    if s == "send":
        return "send"
    return "review"


def _normalize_premium_refine_result(
    raw: Dict[str, Any], *, action: str, current_doc: str
) -> PremiumRefineResponse:
    nxt = _premium_refine_next_step_literal(raw.get("suggested_next_step"))
    sc = raw.get("readiness_score")
    try:
        rs = int(float(sc)) if sc is not None else 0
    except (TypeError, ValueError):
        rs = 0
    if rs < 0:
        rs = 0
    if rs > 100:
        rs = 100
    if action == "ask_missing":
        items = _str_list_capped(raw.get("summary_changes"), 3)
        if not items and isinstance(raw.get("suggestions"), list):
            items = _str_list_capped(raw.get("suggestions"), 3)
        return PremiumRefineResponse(
            updated_document_text=current_doc,
            summary_changes=items[:3],
            readiness_score=rs,
            suggested_next_step=nxt,
        )
    changes = _str_list_capped(raw.get("summary_changes"), 12)
    new_doc = str(raw.get("updated_document_text") or "").strip()
    if not new_doc:
        new_doc = current_doc or ""
    return PremiumRefineResponse(
        updated_document_text=new_doc,
        summary_changes=changes,
        readiness_score=rs,
        suggested_next_step=nxt,
    )


PREMIUM_REFINE_UPDATE_FAIL_OPEN_USER_MESSAGE = (
    "We couldn't apply that update. Your current Pro agreement is unchanged. Try again."
)


def _premium_refine_ws_normalize_for_compare(text: str) -> str:
    """Whitespace-normalized fingerprint for detecting no-op refine outputs."""
    return " ".join((text or "").split())


def _premium_refine_instruction_is_substantive(user_prompt: str) -> bool:
    return len((user_prompt or "").strip()) >= 8


def _premium_refine_update_reject_unchanged_candidate(
    out: PremiumRefineResponse,
    *,
    current_doc: str,
    user_prompt: str,
    source: str,
) -> PremiumRefineResponse:
    """
    action=update: never treat an identical body as a successful refine when the user asked for a real change.
    """
    if not _premium_refine_instruction_is_substantive(user_prompt):
        return out
    cur = _premium_refine_ws_normalize_for_compare(current_doc)
    new = _premium_refine_ws_normalize_for_compare(out.updated_document_text)
    if cur != new:
        return out
    preserved = current_doc or ""
    _log_premium_refine_structured(
        "update_unchanged_candidate_rejected",
        {
            "source": source,
            "current_document_len": len(preserved),
            "out_len": len((out.updated_document_text or "").strip()),
            "instruction_len": len((user_prompt or "").strip()),
        },
    )
    return _premium_refine_update_fail_open_response(preserved)


def _premium_refine_update_fail_open_response(current_doc: str) -> PremiumRefineResponse:
    """Paid Pro action=update: never leave the client with 503 for LLM/parse outages — preserve document."""
    d = current_doc or ""
    return PremiumRefineResponse(
        updated_document_text=d,
        summary_changes=[PREMIUM_REFINE_UPDATE_FAIL_OPEN_USER_MESSAGE],
        readiness_score=35,
        suggested_next_step="review",
    )


def _log_premium_refine_structured(event: str, fields: Dict[str, Any]) -> None:
    """Temporary structured trace for Railway QA (premium-refine narrow + full refine)."""
    payload = {"event": event, **fields}
    try:
        blob = json.dumps(payload, ensure_ascii=False, default=str)
    except Exception:
        blob = json.dumps({"event": event, "serialization": "failed"}, default=str)
    if len(blob) > 12_000:
        blob = blob[:12_000] + "…"
    log.info("claw_premium_refine_trace %s", blob)


def _premium_refine_sample_major_headings(doc: str, cap: int = 40) -> List[str]:
    """Short list of likely section headings for surgical-retry context (best-effort, not exhaustive)."""
    out: List[str] = []
    for raw in (doc or "").splitlines():
        s = raw.strip()
        if not s:
            continue
        if s.startswith("#"):
            out.append(s[:200])
        elif re.match(r"^\d+(?:\.\d+)*\.\s+\S", s):
            out.append(s[:200])
        if len(out) >= cap:
            break
    return out


def _premium_refine_update_system_prompt(
    *, surgical_retry: bool = False, original_document_char_len: Optional[int] = None
) -> str:
    surgical_block = (
        "\n**Surgical / preserve-first update (required for action=update):**\n"
        "- Return the **COMPLETE** agreement in `updated_document_text`.\n"
        "- Do **not** summarize, abridge, or replace the whole agreement.\n"
        "- Do **not** omit sections, headings, numbering, parties, signature blocks, or schedules.\n"
        "- Make **only** the change implied by `user_refinement_prompt`.\n"
        "- If adding a clause, insert it into the most relevant existing section and preserve all other text unchanged.\n"
    )
    retry_block = ""
    if surgical_retry:
        len_hint = ""
        if isinstance(original_document_char_len, int) and original_document_char_len > 0:
            len_hint = f" The original agreement is approximately {original_document_char_len} characters — your output must be similar in length (full document). "
        retry_block = (
            "\n**Second attempt (prior output was too short):** Return the **full** agreement text. "
            "Your prior response was too short. Return the complete agreement with **only** the requested revision applied. "
            "Do not summarize, replace, or omit sections. Copy through the entire `current_document_text`, then edit in place. "
            "Do not shorten the document."
            + len_hint
            + "\n"
        )
    return (
        "You are a precise commercial contract editor in CLAW (a product, not a law firm). The user is refining "
        "their **full agreement text** after generation.\n"
        "Task: **apply** `user_refinement_prompt` to `current_document_text` — change, clarify, add protections, "
        "or tighten language they asked for. Keep party names, key numbers, and business intent aligned with the intake.\n"
        "Rules: Do not invent new economics or parties. Do not strip entire sections unless the user asked. "
        "When in doubt, make a minimal, targeted edit. Plain text only (no HTML).\n"
        + surgical_block
        + retry_block
        + "Output ONLY valid JSON, no markdown fences, with **exact** keys: "
        '{ "updated_document_text": string, "summary_changes": string array (1–6 short imperatives, what changed), '
        '"readiness_score": number 0–100 (higher = closer to send-ready after this pass), '
        '"suggested_next_step": "edit" | "review" | "send" }.\n'
        "Choose suggested_next_step: `edit` if more local edits likely; `review` if a human should scan once; `send` "
        "if the doc looks send-ready and the user’s request was small.\n"
    )


def _premium_refine_ask_missing_system_prompt() -> str:
    return (
        "You are a commercial deal assistant in CLAW. From `intake_text` and `current_document_text`, infer the "
        "**agreement type / deal shape**, then return **exactly 3** concise, high-value bullets: what is missing, "
        "under-specified, or worth double-checking **for this type** (not generic filler).\n"
        "Do not rewrite the agreement. Do not paste clauses.\n"
        "Output ONLY valid JSON, no markdown fences, with **exact** keys: "
        '{ "summary_changes": string array of exactly 3 strings, "readiness_score": number 0–100, '
        '"suggested_next_step": "edit" | "review" | "send" }.\n'
        "summary_changes = the three suggestions. Usually suggested_next_step is `edit` unless the document is very complete.\n"
    )


def _premium_refine_ready_system_prompt() -> str:
    return (
        "You are helping someone decide whether a generated agreement is ready to send. Read `intake_text` and "
        "`current_document_text`. Assess completeness vs typical expectations for the inferred deal type.\n"
        "Return brief, scannable `summary_changes` (2–5 items): strengths and/or last checks — not a full redraft. "
        "readiness_score 0–100: higher = ready to route to signers. suggested_next_step: `send` if strong, "
        "`review` if a careful read is best, `edit` if material gaps remain.\n"
        "The agreement text must NOT change in this mode — omit updated_document_text or set it to empty string.\n"
        "Output ONLY valid JSON, no markdown fences, with **exact** keys: "
        '{ "summary_changes": string array, "readiness_score": number, "suggested_next_step": "edit" | "review" | "send" '
        "}. Optional key `updated_document_text` may be empty; it will be ignored — do not return a full rewrite.\n"
    )


def _user_payload_premium_refine(
    body: PremiumRefineRequest, doc: str, trunc_note: str
) -> Dict[str, Any]:
    p: Dict[str, Any] = {
        "intake": (body.intake_text or "").strip(),
        "action": body.action,
        "current_document_text": doc + trunc_note,
    }
    if (body.user_refinement_prompt or "").strip():
        p["user_refinement_prompt"] = (body.user_refinement_prompt or "").strip()[:12_000]
    if body.action == "update" and bool(getattr(body, "surgical_preserve_retry", False)):
        p["refine_context"] = {
            "mode": "surgical_preserve_retry",
            "original_document_char_len": len(doc),
            "major_headings_sample": _premium_refine_sample_major_headings(doc),
        }
    return p


@router.post("/premium-refine", response_model=PremiumRefineResponse)
def premium_refine(request: Request, body: PremiumRefineRequest) -> PremiumRefineResponse:
    require_claw_org_id_header(request)
    ok_txt, msg_txt = validate_negotiate_text(body.intake_text, "owner")
    if not ok_txt:
        raise HTTPException(status_code=400, detail=msg_txt)
    request_ip = request.client.host if request.client else "unknown"
    if body.action == "update":
        u = (body.user_refinement_prompt or "").strip()
        if not u:
            raise HTTPException(
                status_code=400,
                detail="user_refinement_prompt is required for action=update",
            )
    doc = (body.current_document_text or "").strip()
    if not doc:
        raise HTTPException(status_code=400, detail="current_document_text is required")
    trunc_note = ""
    if len(doc) > 200_000:
        doc = doc[:200_000]
        trunc_note = " [Draft truncated to first ~200k characters.]"
    user_payload = _user_payload_premium_refine(body, doc, trunc_note)
    plen = len(json.dumps(user_payload, ensure_ascii=False))
    if plen > 270_000:
        raise HTTPException(status_code=400, detail="Input too large for premium refine")

    if body.action == "update":
        system = _premium_refine_update_system_prompt(
            surgical_retry=bool(body.surgical_preserve_retry),
            original_document_char_len=len(doc) if body.surgical_preserve_retry else None,
        )
        max_out = max(2000, int(os.environ.get("CLAW_PREMIUM_REFINE_UPDATE_MAX_TOKENS", "12000")))
    elif body.action == "ask_missing":
        system = _premium_refine_ask_missing_system_prompt()
        max_out = max(500, int(os.environ.get("CLAW_PREMIUM_REFINE_ASK_MAX_TOKENS", "2000")))
    else:
        system = _premium_refine_ready_system_prompt()
        max_out = max(500, int(os.environ.get("CLAW_PREMIUM_REFINE_READY_MAX_TOKENS", "2000")))

    llm_model = resolve_llm_model_for_access_class("premium")
    _raw = body.current_document_text or ""
    current_for_norm = _raw[:200_000] if len(_raw) > 200_000 else _raw

    pr_len = len((body.user_refinement_prompt or "").strip()) if body.action == "update" else 0
    log.info(
        "claw_premium route=premium_refine action=%s model=%s max_out=%d %s current_document_len=%d "
        "intake_text_len=%d user_refinement_prompt_len=%d",
        body.action,
        (llm_model or ""),
        max_out,
        _openai_key_diagnostics(),
        len(doc),
        len((body.intake_text or "").strip()),
        pr_len,
    )

    def _run_premium_refine_llm_path() -> PremiumRefineResponse:
        llm_text = call_legal_llm(
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": json.dumps(user_payload, ensure_ascii=False)},
            ],
            model=llm_model,
            max_tokens=max_out,
            temperature=0.2 if body.action == "update" else 0.15,
            airlock_profile="agreement_outbound",
        )
        log.info(
            "claw_premium route=premium_refine openai_response_chars=%d action=%s",
            len((llm_text or "").strip()),
            body.action,
        )
        parsed = _extract_json_object(llm_text)
        if body.action == "ask_missing" and (not isinstance(parsed.get("summary_changes"), list)) and isinstance(
            parsed.get("suggestions"), list
        ):
            parsed["summary_changes"] = parsed["suggestions"]
        out_llm = _normalize_premium_refine_result(parsed, action=body.action, current_doc=current_for_norm)
        if body.action == "ready" and not (out_llm.updated_document_text or "").strip():
            out_llm = out_llm.model_copy(update={"updated_document_text": current_for_norm})
        _safe_record_ai_call(request, request_ip)
        return out_llm

    try:
        if body.action == "update":
            u_narrow = (body.user_refinement_prompt or "").strip()
            instr_snip = u_narrow[:240] + ("…" if len(u_narrow) > 240 else "")
            narrow_kind_run: Optional[str] = None
            narrow_reason = "skipped"
            _log_premium_refine_structured(
                "update_enter",
                {
                    "action": body.action,
                    "narrow_kind": None,
                    "current_document_len": len(doc),
                    "instruction_snippet": instr_snip,
                },
            )
            try:
                narrow_kind_run = classify_narrow_amendment_prompt(u_narrow)
                if narrow_kind_run:
                    narrow_out = try_apply_narrow_amendment(
                        kind=narrow_kind_run,
                        current_document_text=doc,
                        user_refinement_prompt=u_narrow,
                        call_legal_llm_fn=call_legal_llm,
                        llm_model=llm_model,
                    )
                    if narrow_out is not None:
                        narrow_reason = "narrow_apply_ok"
                        parsed_narrow = {
                            "updated_document_text": narrow_out["updated_document_text"],
                            "summary_changes": narrow_out.get("summary_changes") or [],
                            "readiness_score": int(narrow_out.get("readiness_score") or 78),
                            "suggested_next_step": narrow_out.get("suggested_next_step") or "review",
                        }
                        out_narrow = _normalize_premium_refine_result(
                            parsed_narrow, action=body.action, current_doc=current_for_norm
                        )
                        log.info(
                            "claw_premium route=premium_refine narrow_amendment kind=%s out_len=%d current_len=%d",
                            narrow_kind_run,
                            len((out_narrow.updated_document_text or "").strip()),
                            len(doc),
                        )
                        _log_premium_refine_structured(
                            "narrow_success",
                            {
                                "action": body.action,
                                "narrow_kind": narrow_kind_run,
                                "narrow_success_reason": narrow_reason,
                                "current_document_len": len(doc),
                                "instruction_snippet": instr_snip,
                                "out_len": len((out_narrow.updated_document_text or "").strip()),
                            },
                        )
                        _safe_record_ai_call(request, request_ip)
                        return _premium_refine_update_reject_unchanged_candidate(
                            out_narrow,
                            current_doc=current_for_norm,
                            user_prompt=u_narrow,
                            source="narrow_amendment",
                        )
                    narrow_reason = "narrow_apply_returned_none"
                else:
                    narrow_reason = "not_classified_as_narrow"
            except Exception as narrow_exc:
                narrow_reason = f"narrow_exception:{type(narrow_exc).__name__}"
                log.warning(
                    "claw_premium route=premium_refine narrow_amendment_exception_fallback_to_llm err_type=%s repr=%r",
                    type(narrow_exc).__name__,
                    repr(narrow_exc),
                    exc_info=True,
                )
                _log_premium_refine_structured(
                    "narrow_exception",
                    {
                        "action": body.action,
                        "narrow_kind": narrow_kind_run,
                        "narrow_success_reason": narrow_reason,
                        "current_document_len": len(doc),
                        "instruction_snippet": instr_snip,
                        "exception_repr": repr(narrow_exc),
                    },
                )

            _log_premium_refine_structured(
                "full_refine_attempt",
                {
                    "action": body.action,
                    "narrow_kind": narrow_kind_run,
                    "narrow_success_reason": narrow_reason,
                    "current_document_len": len(doc),
                    "instruction_snippet": instr_snip,
                },
            )
            try:
                out_update = _run_premium_refine_llm_path()
                _log_premium_refine_structured(
                    "full_refine_success",
                    {
                        "action": body.action,
                        "narrow_kind": narrow_kind_run,
                        "narrow_success_reason": narrow_reason,
                        "fallback_full_refine_success": True,
                        "current_document_len": len(doc),
                        "instruction_snippet": instr_snip,
                        "out_len": len((out_update.updated_document_text or "").strip()),
                    },
                )
                return _premium_refine_update_reject_unchanged_candidate(
                    out_update,
                    current_doc=current_for_norm,
                    user_prompt=u_narrow,
                    source="full_llm_refine",
                )
            except Exception as full_exc:
                _log_premium_refine_structured(
                    "full_refine_failure_fail_open",
                    {
                        "action": body.action,
                        "narrow_kind": narrow_kind_run,
                        "narrow_success_reason": narrow_reason,
                        "fallback_full_refine_success": False,
                        "fallback_full_refine_failure_reason": type(full_exc).__name__,
                        "current_document_len": len(doc),
                        "instruction_snippet": instr_snip,
                        "exception_repr": repr(full_exc),
                    },
                )
                log.warning(
                    "claw_premium route=premium_refine update_fail_open err_type=%s repr=%r",
                    type(full_exc).__name__,
                    repr(full_exc),
                    exc_info=True,
                )
                return _premium_refine_update_fail_open_response(current_for_norm)

        out = _run_premium_refine_llm_path()
        return out
    except Exception as exc:
        kind = _classify_premium_llm_failure(exc)
        log.warning(
            "claw_premium route=premium_refine FAILED class=%s exc_type=%s action=%s model=%s %s err_snip=%s",
            kind,
            type(exc).__name__,
            body.action,
            (llm_model or ""),
            _openai_key_diagnostics(),
            (str(exc) or "")[:500].replace("\n", " "),
            exc_info=True,
        )
        raise HTTPException(
            status_code=503,
            detail={
                "code": "premium_refine_unavailable",
                "message": "Smart refinement is temporarily unavailable. Try again or edit the document directly.",
            },
        ) from exc


class PremiumFinalizeAuditContext(BaseModel):
    """Optional grounding for /premium-finalize-audit — not required."""

    agreement_family: str = ""
    material_asks: List[str] = Field(default_factory=list)
    user_gap_answers: str = ""
    party_labels: List[str] = Field(default_factory=list)
    parse_extract: Optional[Dict[str, Any]] = None
    premium_review: Optional[Dict[str, Any]] = None


class PremiumFinalizeAuditRequest(BaseModel):
    intake_text: str = Field(..., min_length=1)
    document_text: str = Field(..., min_length=1)
    context: Optional[PremiumFinalizeAuditContext] = None


class PremiumFinalizeAuditResponse(BaseModel):
    deal_specific_missing_terms: List[str] = Field(default_factory=list)
    placeholder_terms_found: List[str] = Field(default_factory=list)
    resolved_strengths: List[str] = Field(default_factory=list)
    best_next_step: Literal["edit", "review", "send"] = "review"
    confidence: Literal["low", "medium", "high"] = "medium"


def _safe_record_ai_call(request: Request, request_ip: str) -> None:
    """Usage accounting must not turn a successful LLM/narrow response into HTTP 503."""
    try:
        record_ai_call(subject_ref=resolve_subject_from_request(request), request_ip=request_ip or "unknown")
    except Exception as exc:
        log.warning(
            "claw_premium record_ai_call skipped err_type=%s",
            type(exc).__name__,
            exc_info=True,
        )


def _fallback_premium_agreement_review_response() -> PremiumAgreementReviewResponse:
    """Advisory-only: empty bullets so Pro send flow is never blocked by review LLM outages."""
    return PremiumAgreementReviewResponse(
        strengths=[],
        missing_or_weak_terms=[],
        questions_for_user=[],
        suggested_clause_upgrades=[],
        priority_score=35,
    )


def _fallback_premium_finalize_audit_response() -> PremiumFinalizeAuditResponse:
    """Advisory-only: neutral audit shell when finalize-audit LLM is unavailable."""
    return PremiumFinalizeAuditResponse(
        deal_specific_missing_terms=[],
        placeholder_terms_found=[],
        resolved_strengths=[],
        best_next_step="review",
        confidence="medium",
    )


class PremiumReviewRouteRequest(BaseModel):
    intake_text: str = Field(..., min_length=1)
    finalize_answers: str = ""
    agreement_text: str = Field(..., min_length=1)
    party_count: int = Field(2, ge=1, le=12)
    agreement_family: str = ""


class PremiumReviewRouteResponse(BaseModel):
    route: Literal["signature", "review", "fix"] = "review"
    confidence: Literal["low", "medium", "high"] = "medium"
    unresolved_items: List[str] = Field(default_factory=list)
    reasons: List[str] = Field(default_factory=list)
    send_readiness_score: int = Field(0, ge=0, le=100)
    recommended_cta: Literal["Send with confidence", "Start collaborative review", "Make these quick upgrades"] = (
        "Start collaborative review"
    )
    short_summary: str = ""


_PLACEHOLDER_SUBSTRINGS: Tuple[Tuple[str, str], ...] = (
    (r"(?i)\bto be agreed\b", "Unresolved: “to be agreed” still appears"),
    (r"(?i)\bto be defined\b", "Unresolved: “to be defined” still appears"),
    (r"(?i)\bto be determined\b", "Unresolved: “to be determined” still appears"),
    (r"(?i)\btbd\b", "Unresolved: “TBD” still appears"),
    (r"(?i)\bnot (?:yet )?specified\b", "Unresolved: “not specified” phrasing still appears"),
    (r"(?i)\bnot (?:yet )?finalized\b", "Unresolved: “not finalized” phrasing still appears"),
    (r"(?i)\brefine in review\b", "Unresolved: “refine in review” placeholder still appears"),
    (r"(?i)\bname not provided\b", "Unresolved: “name not provided” still appears"),
    (r"(?i)\bplaceholder party\b", "Placeholder party language still appears"),
    (r"(?i)\b\[insert\b", "Bracket/insert-style placeholder still appears"),
)


def _scan_document_placeholders_local(text: str) -> List[str]:
    if not (text or "").strip():
        return []
    out: List[str] = []
    for pat, label in _PLACEHOLDER_SUBSTRINGS:
        if re.search(pat, text) and label not in out:
            out.append(label)
        if len(out) >= 5:
            break
    return out[:5]


def _str_list_capped_5(raw: Any) -> List[str]:
    return _str_list_capped(raw, 5)[:5]


def _normalize_confidence(v: Any) -> str:
    s = str(v or "medium").strip().lower()
    if s in ("low", "medium", "high"):
        return s
    return "medium"


def _normalize_best_next(v: Any) -> str:
    s = str(v or "review").strip().lower()
    if s in ("edit", "review", "send"):
        return s
    return "review"


def _normalize_premium_finalize_audit_payload(raw: Dict[str, Any]) -> Dict[str, Any]:
    return {
        "deal_specific_missing_terms": _str_list_capped_5(raw.get("deal_specific_missing_terms")),
        "placeholder_terms_found": _str_list_capped_5(raw.get("placeholder_terms_found")),
        "resolved_strengths": _str_list_capped_5(raw.get("resolved_strengths")),
        "best_next_step": _normalize_best_next(raw.get("best_next_step")),
        "confidence": _normalize_confidence(raw.get("confidence")),
    }


def _merge_placeholder_terms_model_and_local(
    from_model: List[str], document_text: str
) -> List[str]:
    seen: set = set()
    out: List[str] = []
    for x in from_model:
        s = str(x).strip()
        if s and s not in seen:
            seen.add(s)
            out.append(s[:2000])
    for x in _scan_document_placeholders_local(document_text):
        if x not in seen:
            seen.add(x)
            out.append(x)
        if len(out) >= 5:
            break
    return out[:5]


def _premium_finalize_audit_system_prompt() -> str:
    return (
        "You are a **universal commercial agreement completion auditor** for CLAW (a product, not a law firm). "
        "Your job is to compare the user’s **intake** to the final **agreement text** and answer one question: "
        "**What is still unclear, weak, or only partially addressed for THIS specific deal?**\n"
        "Rules:\n"
        "- Prioritize **deal-specific** gaps: ownership of assets, profit/repayment waterfalls, spending authority, "
        "exit/buyout, bank controls, timing, performance obligations, and any subject the intake makes concrete but "
        "the agreement leaves vague — **over** generic “best practice” checklists.\n"
        "- Only mention standard topics (e.g. confidentiality, IP, indemnity) when the intake or the document’s "
        "substance **clearly** implicates them as open issues; do **not** list them by default.\n"
        "- Flag **placeholders and vague drafting** explicitly: e.g. to be agreed, TBD, not specified, name not provided, "
        "generic party labels where real names were expected, refine in review, bracket/insert text.\n"
        "- `deal_specific_missing_terms` (0–5): short, user-facing bullets; each a single clear gap. No long prose, "
        "no numbered clauses, no “consider consulting counsel.”\n"
        "- `placeholder_terms_found` (0–5): what vague or placeholder phrasing in the **document** is still "
        "problematic (or confirm none material).\n"
        "- `resolved_strengths` (0–5): what is already well wired between intake and doc (short).\n"
        "- `best_next_step`: `edit` if real drafting work left; `review` if a careful read/negotiation handoff is next; "
        "`send` only if the text is plausibly send-ready and gaps are minor/operational only.\n"
        "- `confidence` low|medium|high: how sure you are given messy inputs; low when the prompt is too thin to be sure.\n"
        "Output **only** valid JSON, no markdown, with **exact** keys: "
        '{ "deal_specific_missing_terms": string array, "placeholder_terms_found": string array, '
        '"resolved_strengths": string array, "best_next_step": "edit" | "review" | "send", "confidence": "low" | "medium" | "high" }.\n'
    )


def _premium_review_route_system_prompt() -> str:
    return (
        "You are an agreement completion reviewer for LawDog Pro (product workflow, not legal advice). "
        "Classify the best next action for this exact agreement into one route: signature, review, or fix.\n"
        "Inputs include intake, finalize answers, full agreement text, party count, and agreement family.\n"
        "Priority checks (highest first):\n"
        "- missing economics, payment triggers, or payout timing\n"
        "- missing ownership/control terms, buyout/exit mechanics, authority controls\n"
        "- placeholders/TBD/unclear party identities/dates\n"
        "- contradictory obligations or unclear responsibilities\n"
        "- family/friend/cofounder dynamics where fairness or consent likely needs review before signature\n"
        "Route rules:\n"
        "- `signature`: only if the agreement appears materially complete and clear with no meaningful unresolved placeholders.\n"
        "- `review`: default for reviewable drafts with understandable structure/economics, even if some placeholders or follow-ups remain.\n"
        "- `fix`: reserve for truly broken drafts only (missing core economics, contradictory terms, placeholder-heavy text, or missing essential party structure).\n"
        "Output JSON only (no markdown) with exact keys:\n"
        '{ "route": "signature"|"review"|"fix", "confidence": "low"|"medium"|"high", '
        '"unresolved_items": string array, "reasons": string array, "send_readiness_score": number, '
        '"recommended_cta": "Send for signature"|"Send for review"|"Fix a few items first", '
        '"short_summary": string }.\n'
        "Keep unresolved_items/reasons concise and user-facing.\n"
    )


def _normalize_review_route(v: Any) -> str:
    s = str(v or "review").strip().lower()
    if s in ("signature", "review", "fix"):
        return s
    return "review"


def _route_to_cta(route: str) -> str:
    if route == "signature":
        return "Send with confidence"
    if route == "fix":
        return "Make these quick upgrades"
    return "Start collaborative review"


def _normalize_premium_review_route_result(raw: Dict[str, Any]) -> PremiumReviewRouteResponse:
    route = _normalize_review_route(raw.get("route"))
    conf = _normalize_confidence(raw.get("confidence"))
    sc = raw.get("send_readiness_score")
    try:
        score = int(float(sc)) if sc is not None else 0
    except (TypeError, ValueError):
        score = 0
    if score < 0:
        score = 0
    if score > 100:
        score = 100
    unresolved = _str_list_capped(raw.get("unresolved_items"), 5)
    reasons = _str_list_capped(raw.get("reasons"), 5)
    summary = str(raw.get("short_summary") or "").strip()[:600]
    if not summary:
        if route == "signature":
            summary = "Agreement looks complete enough to send for signatures."
        elif route == "fix":
            summary = "A few unresolved items should be fixed before sending."
        else:
            summary = "This deal likely benefits from review/redline before signatures."
    return PremiumReviewRouteResponse(
        route=cast(Literal["signature", "review", "fix"], route),
        confidence=cast(Literal["low", "medium", "high"], conf),
        unresolved_items=unresolved,
        reasons=reasons,
        send_readiness_score=score,
        recommended_cta=cast(
            Literal["Send with confidence", "Start collaborative review", "Make these quick upgrades"],
            _route_to_cta(route),
        ),
        short_summary=summary,
    )


def _is_generic_route_unresolved_item(item: str) -> bool:
    s = (item or "").strip().lower()
    if not s:
        return True
    generic = (
        "governing law",
        "venue",
        "boilerplate",
        "miscellaneous",
        "remaining review placeholders",
        "generic",
    )
    return any(k in s for k in generic)


def _derive_deal_specific_route_items(payload: Dict[str, Any]) -> List[str]:
    intake = str(payload.get("intake_text") or "").lower()
    finalize = str(payload.get("finalize_answers") or "").lower()
    agreement = str(payload.get("agreement_text") or "").lower()
    family = str(payload.get("agreement_family") or "").lower()
    text = f"{intake}\n{finalize}\n{agreement}\n{family}"

    is_referral = any(k in text for k in ("referral", "commission", "no bypass", "lead"))
    is_family_coowner = any(k in text for k in ("cousin", "family", "co-owner", "coffee cart", "shared ownership"))
    is_marketing = any(k in text for k in ("campaign", "ad spend", "meta", "tiktok", "google ads"))

    if is_referral:
        return [
            "Confirm the net-revenue definition used to calculate commission.",
            "Confirm payout timing, ledger statement format, and chargeback/refund offsets.",
            "Confirm no-bypass scope and post-termination survival period.",
        ]
    if is_family_coowner:
        return [
            "Confirm ownership allocation for equipment, profits, and voting/control rights.",
            "Confirm buyout trigger events and valuation method (including dispute fallback).",
            "Confirm spending-approval limits plus bank/cash-control workflow.",
        ]
    if is_marketing:
        return [
            "Confirm spend-approval thresholds and emergency override workflow.",
            "Confirm account/data ownership boundaries (platform access, pixels, creatives).",
            "Confirm performance-remedy triggers and termination rights.",
        ]
    return [
        "Confirm remaining commercial economics and approval mechanics.",
        "Confirm party obligations and escalation path for disputes.",
        "Confirm send-ready signatures, dates, and operational handoff terms.",
    ]


def _polish_premium_review_route_output(
    out: PremiumReviewRouteResponse, payload: Dict[str, Any]
) -> PremiumReviewRouteResponse:
    intake = str(payload.get("intake_text") or "").lower()
    finalize = str(payload.get("finalize_answers") or "").lower()
    family = str(payload.get("agreement_family") or "").lower()
    context = f"{intake}\n{finalize}\n{family}"
    is_referral = any(k in context for k in ("referral", "commission", "no bypass"))
    is_family_coowner = any(k in context for k in ("cousin", "family", "co-owner", "coffee cart", "shared ownership"))

    unresolved = [x for x in out.unresolved_items if str(x).strip()]
    non_generic = [x for x in unresolved if not _is_generic_route_unresolved_item(x)]
    derived = _derive_deal_specific_route_items(payload)
    if out.route == "review":
        # REVIEW should feel intentional and premium, not second-best.
        if is_family_coowner:
            out.short_summary = "Best next step: send for review so both co-owners align on control and exit terms before signature."
            out.reasons = (
                [r for r in out.reasons if str(r).strip()]
                + ["Shared ownership dynamics are strongest when both sides review commercially sensitive terms first."]
            )[:5]
        elif is_referral:
            out.short_summary = "Best next step: send for review so commission economics and no-bypass boundaries are confirmed before signature."
            out.reasons = (
                [r for r in out.reasons if str(r).strip()]
                + ["Counterparty review helps lock commission mechanics and post-termination protections."]
            )[:5]
        else:
            out.short_summary = "Best next step: send for review to confirm remaining commercial details before signature."
        out.recommended_cta = "Send for review"
        if out.send_readiness_score < 65:
            out.send_readiness_score = 65
        # Replace generic unresolved bullets with deal-specific commercial bullets.
        curated = (non_generic + derived)[:3]
        out.unresolved_items = curated
    return out


def _looks_placeholder_heavy(agreement_text: str, placeholder_hits: List[str]) -> bool:
    text = (agreement_text or "").lower()
    hit_count = len([x for x in placeholder_hits if str(x).strip()])
    token_hits = 0
    for k in ("tbd", "to be agreed", "to be selected in review", "name not provided", "party_a", "party_b"):
        if k in text:
            token_hits += 1
    return hit_count >= 4 or token_hits >= 3


def _has_core_economics(agreement_text: str) -> bool:
    text = (agreement_text or "").lower()
    if re.search(r"\$\s?\d", text):
        return True
    if re.search(r"\b\d{1,3}\s?%\b", text):
        return True
    econ_terms = (
        "commission",
        "payment",
        "payout",
        "invoice",
        "fee",
        "rent",
        "revenue",
        "profit",
        "net revenue",
        "operating reserve",
    )
    return any(t in text for t in econ_terms)


def _has_party_structure(agreement_text: str) -> bool:
    text = (agreement_text or "").lower()
    return any(t in text for t in ("party a", "party b", "party_a", "party_b", "between"))


def _retune_review_route_thresholds(out: PremiumReviewRouteResponse, payload: Dict[str, Any]) -> PremiumReviewRouteResponse:
    agreement = str(payload.get("agreement_text") or "")
    placeholder_hits = [str(x) for x in (payload.get("placeholder_hits") or []) if str(x).strip()]
    unresolved = [str(x).lower() for x in out.unresolved_items]
    party_count = int(payload.get("party_count") or 2)

    placeholder_heavy = _looks_placeholder_heavy(agreement, placeholder_hits)
    has_econ = _has_core_economics(agreement)
    has_parties = _has_party_structure(agreement)
    contradiction_flag = any("contradict" in x for x in unresolved)
    missing_core_flag = any(
        k in x
        for x in unresolved
        for k in ("cannot understand economics", "missing core economics", "missing essential party structure")
    )

    # Keep signature strict: if unresolved placeholders/material holes exist, downgrade to review.
    if out.route == "signature" and (placeholder_heavy or len(out.unresolved_items) >= 2):
        out.route = "review"
        out.recommended_cta = "Send for review"
        out.send_readiness_score = min(out.send_readiness_score, 82)
        if "Review before signature is recommended." not in out.reasons:
            out.reasons = (out.reasons + ["Review before signature is recommended."])[:5]
        if not out.short_summary:
            out.short_summary = "This agreement is close, but should be reviewed before signatures."

    # FIX only for truly broken drafts; otherwise prefer REVIEW.
    if out.route == "fix":
        truly_broken = contradiction_flag or missing_core_flag or placeholder_heavy or (not has_econ) or (
            party_count >= 2 and not has_parties
        )
        if not truly_broken:
            out.route = "review"
            out.recommended_cta = "Send for review"
            out.send_readiness_score = max(60, min(out.send_readiness_score, 78))
            if not out.short_summary:
                out.short_summary = "Agreement is reviewable but should be sent for review before signatures."
            if "Review is appropriate while remaining items are finalized." not in out.reasons:
                out.reasons = (out.reasons + ["Review is appropriate while remaining items are finalized."])[:5]
    return out


def _unresolved_risk_rank(item: str) -> int:
    s = (item or "").lower()
    if any(k in s for k in ("economics", "payment", "payout", "commission", "revenue", "profit", "fee")):
        return 0
    if any(k in s for k in ("ownership", "control", "authority", "buyout", "exit", "voting")):
        return 1
    if any(k in s for k in ("placeholder", "tbd", "party", "name", "date", "governing law", "venue")):
        return 2
    if any(k in s for k in ("workflow", "format", "notice", "logistics", "process")):
        return 3
    return 4


def _apply_route_psychology_copy(out: PremiumReviewRouteResponse) -> PremiumReviewRouteResponse:
    unresolved_sorted = sorted([x for x in out.unresolved_items if str(x).strip()], key=_unresolved_risk_rank)[:5]
    out.unresolved_items = unresolved_sorted
    if out.route == "review":
        out.recommended_cta = "Start collaborative review"
        out.short_summary = "Best next step: collaborate on a focused review now to protect deal momentum."
    elif out.route == "signature":
        out.recommended_cta = "Send with confidence"
        out.short_summary = "Ready to send: this agreement looks commercially complete and signature-ready."
    else:
        out.recommended_cta = "Make these quick upgrades"
        out.short_summary = "Almost there: make these quick upgrades now, then send with confidence."
    return out


def _fallback_premium_review_route(payload: Dict[str, Any]) -> PremiumReviewRouteResponse:
    intake = str(payload.get("intake_text") or "").lower()
    finalize_answers = str(payload.get("finalize_answers") or "").lower()
    agreement = str(payload.get("agreement_text") or "").lower()
    family = str(payload.get("agreement_family") or "").lower()
    party_count = int(payload.get("party_count") or 2)
    placeholder_hits = [str(x) for x in (payload.get("placeholder_hits") or []) if str(x).strip()]

    unresolved: List[str] = []
    reasons: List[str] = []
    route: Literal["signature", "review", "fix"] = "review"
    score = 72
    conf: Literal["low", "medium", "high"] = "medium"

    severe_placeholders = len(placeholder_hits) >= 4
    if severe_placeholders:
        unresolved.append("Clear unresolved placeholders before sending.")
        route = "fix"
        score = 42
        conf = "low"

    fairness_signal = any(k in f"{intake} {finalize_answers}" for k in ["cousin", "family", "co-owner", "partner"])
    ownership_signal = any(k in f"{agreement} {finalize_answers}" for k in ["ownership", "buyout", "fair market value"])
    missing_dates_signal = any(k in agreement for k in ["to be selected in review", "tbd governing law", "to be agreed"])

    if route != "fix":
        if party_count >= 2 and (fairness_signal or "partnership" in family or ownership_signal):
            route = "review"
            score = min(score, 76)
            reasons.append("Multi-party ownership/control terms are likely to benefit from bilateral review before signature.")
        else:
            route = "signature"
            score = 86
            conf = "low"
            reasons.append("No major unresolved gaps were detected by fallback routing checks.")

    if missing_dates_signal:
        unresolved.append("Confirm governing law/venue and any remaining review placeholders.")
        if route == "signature":
            route = "review"
            score = 74
            conf = "medium"

    if not reasons:
        reasons.append("Fallback route generated because live route model was temporarily unavailable.")

    raw = {
        "route": route,
        "confidence": conf,
        "unresolved_items": unresolved[:5],
        "reasons": reasons[:5],
        "send_readiness_score": score,
        "recommended_cta": _route_to_cta(route),
        "short_summary": (
            "This deal is best sent for review before signatures."
            if route == "review"
            else "A few unresolved items should be fixed before sending."
            if route == "fix"
            else "Agreement appears complete enough to send for signatures."
        ),
    }
    out = _retune_review_route_thresholds(_normalize_premium_review_route_result(raw), payload)
    out = _polish_premium_review_route_output(out, payload)
    return _apply_route_psychology_copy(out)


@router.post("/premium-finalize-audit", response_model=PremiumFinalizeAuditResponse)
def premium_finalize_audit(request: Request, body: PremiumFinalizeAuditRequest) -> PremiumFinalizeAuditResponse:
    require_claw_org_id_header(request)
    ok_txt, msg_txt = validate_negotiate_text(body.intake_text, "owner")
    if not ok_txt:
        raise HTTPException(status_code=400, detail=msg_txt)
    request_ip = request.client.host if request.client else "unknown"
    doc = (body.document_text or "").strip()
    if not doc:
        raise HTTPException(status_code=400, detail="document_text is required")
    trunc = ""
    if len(doc) > 200_000:
        doc = doc[:200_000]
        trunc = " [Document truncated.]"
    ctx_obj: Dict[str, Any] = {
        "intake": (body.intake_text or "").strip()[:120_000],
        "agreement": doc + trunc,
    }
    if body.context is not None:
        c = body.context
        pctx: Dict[str, Any] = {
            "agreement_family": (c.agreement_family or "").strip(),
            "material_asks": [str(x).strip() for x in (c.material_asks or []) if str(x).strip()][:16],
            "user_gap_answers": (c.user_gap_answers or "")[:20_000],
            "party_labels": [str(x).strip() for x in (c.party_labels or []) if str(x).strip()][:20],
        }
        if c.parse_extract is not None:
            pctx["parse_extract"] = c.parse_extract
        if c.premium_review is not None:
            pctx["premium_review"] = c.premium_review
        ctx_obj["grounding"] = pctx
    if len(json.dumps(ctx_obj, ensure_ascii=False)) > 280_000:
        raise HTTPException(status_code=400, detail="Input too large for premium finalize audit")
    llm_model = resolve_llm_model_for_access_class("premium")
    max_out = max(600, int(os.environ.get("CLAW_PREMIUM_FINALIZE_AUDIT_MAX_TOKENS", "3000")))
    log.info(
        "claw_premium route=premium_finalize_audit model=%s max_out=%d %s document_len=%d intake_len=%d",
        (llm_model or ""),
        max_out,
        _openai_key_diagnostics(),
        len(doc),
        len((body.intake_text or "").strip()),
    )
    try:
        llm_text = call_legal_llm(
            messages=[
                {"role": "system", "content": _premium_finalize_audit_system_prompt()},
                {"role": "user", "content": json.dumps(ctx_obj, ensure_ascii=False)},
            ],
            model=llm_model,
            max_tokens=max_out,
            temperature=0.1,
            airlock_profile="agreement_outbound",
        )
        log.info("claw_premium route=premium_finalize_audit openai_response_chars=%d", len((llm_text or "").strip()))
        parsed = _extract_json_object(llm_text)
        n = _normalize_premium_finalize_audit_payload(parsed)
        n["placeholder_terms_found"] = _merge_placeholder_terms_model_and_local(
            [str(x) for x in n.get("placeholder_terms_found", []) if str(x).strip()],
            (body.document_text or "")[:200_000],
        )
        n["deal_specific_missing_terms"] = n["deal_specific_missing_terms"][:5]
        n["resolved_strengths"] = n["resolved_strengths"][:5]
        if not isinstance(n.get("best_next_step"), str) or n["best_next_step"] not in ("edit", "review", "send"):
            n["best_next_step"] = "review"
        if n.get("confidence") not in ("low", "medium", "high"):
            n["confidence"] = "medium"
        out = PremiumFinalizeAuditResponse(
            deal_specific_missing_terms=[str(x) for x in n.get("deal_specific_missing_terms", []) if str(x).strip()][
                :5
            ],
            placeholder_terms_found=[str(x) for x in n.get("placeholder_terms_found", []) if str(x).strip()][:5],
            resolved_strengths=[str(x) for x in n.get("resolved_strengths", []) if str(x).strip()][:5],
            best_next_step=cast(Literal["edit", "review", "send"], n["best_next_step"]),
            confidence=cast(Literal["low", "medium", "high"], n["confidence"]),
        )
        _safe_record_ai_call(request, request_ip)
        return out
    except Exception as exc:
        kind = _classify_premium_llm_failure(exc)
        log.warning(
            "claw_premium route=premium_finalize_audit FAILED class=%s exc_type=%s model=%s %s err_snip=%s (fail_open_200)",
            kind,
            type(exc).__name__,
            (llm_model or ""),
            _openai_key_diagnostics(),
            (str(exc) or "")[:500].replace("\n", " "),
            exc_info=True,
        )
        return _fallback_premium_finalize_audit_response()


@router.post("/premium-review-route", response_model=PremiumReviewRouteResponse)
def premium_review_route(request: Request, body: PremiumReviewRouteRequest) -> PremiumReviewRouteResponse:
    require_claw_org_id_header(request)
    ok_txt, msg_txt = validate_negotiate_text(body.intake_text, "owner")
    if not ok_txt:
        raise HTTPException(status_code=400, detail=msg_txt)
    request_ip = request.client.host if request.client else "unknown"
    agreement = (body.agreement_text or "").strip()
    if not agreement:
        raise HTTPException(status_code=400, detail="agreement_text is required")
    if len(agreement) > 220_000:
        agreement = agreement[:220_000]
    payload: Dict[str, Any] = {
        "intake_text": (body.intake_text or "").strip(),
        "finalize_answers": (body.finalize_answers or "").strip()[:20_000],
        "agreement_text": agreement,
        "party_count": int(body.party_count),
        "agreement_family": (body.agreement_family or "").strip(),
        "placeholder_hits": _scan_document_placeholders_local(agreement),
    }
    if len(json.dumps(payload, ensure_ascii=False)) > 300_000:
        raise HTTPException(status_code=400, detail="Input too large for premium review route")
    llm_model = resolve_llm_model_for_access_class("premium")
    max_out = max(600, int(os.environ.get("CLAW_PREMIUM_REVIEW_ROUTE_MAX_TOKENS", "2200")))
    try:
        llm_text = call_legal_llm(
            messages=[
                {"role": "system", "content": _premium_review_route_system_prompt()},
                {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
            ],
            model=llm_model,
            max_tokens=max_out,
            temperature=0.1,
            airlock_profile="agreement_outbound",
        )
        parsed = _extract_json_object(llm_text)
        out = _retune_review_route_thresholds(_normalize_premium_review_route_result(parsed), payload)
        out = _polish_premium_review_route_output(out, payload)
        out = _apply_route_psychology_copy(out)
        _safe_record_ai_call(request, request_ip)
        return out
    except Exception as exc:
        log.warning("premium_review_route failed err=%s", type(exc).__name__, exc_info=True)
        return _fallback_premium_review_route(payload)


class AgreementRenderResponse(BaseModel):
    id: str
    rendered_html: str


RecipientPreviewPdfExportKind = Literal["original", "proposed", "redline"]


class RecipientPreviewPdfExportBody(BaseModel):
    """Client supplies the same scrubbed HTML shown in the preview (original / proposed) or inline-styled redline HTML."""

    export_kind: RecipientPreviewPdfExportKind = "original"
    html: str = Field(..., min_length=1, max_length=1_200_000)

    @field_validator("html", mode="before")
    @classmethod
    def _trim_html(cls, v: object) -> str:
        return ("" if v is None else str(v)).strip()


_RECIPIENT_PREVIEW_PDF_FILENAMES: Dict[str, str] = {
    "original": "lawdog-original-draft.pdf",
    "proposed": "lawdog-proposed-draft.pdf",
    "redline": "lawdog-redline-preview.pdf",
}

_RECIPIENT_PDF_EXPORT_UNAVAILABLE_USER = (
    "PDF export is temporarily unavailable. Please use Copy or Download text for now."
)


SessionType = Literal["owner", "recipient"]


class AgreementReviseRequest(BaseModel):
    instruction: str
    session_type: SessionType = "owner"
    # When False, run the same revise pipeline but do not persist (workspace preview / compare).
    persist: bool = True
    # basic = STARTER_PREVIEW revise (minimal edits). premium = PREMIUM_SURGICAL by default, or
    # PREMIUM_MATERIAL_UPGRADE when the instruction explicitly asks for a broad polish (similarity-based auto-retry).
    ai_model_class: Literal["basic", "premium"] = "basic"


class AgreementRefineRequest(BaseModel):
    """Minimal body for create-flow refine bar (delegates to POST /revise)."""

    instruction: str


class AgreementCommitRevisionRequest(BaseModel):
    """Persist a revised draft from a prior preview (POST /revise with persist=False)."""

    instruction: str = ""
    draft: AgreementDraftCreate


class RecipientProposalRequest(BaseModel):
    """Recipient preview outcome: stored as audit-only pending until owner applies (canonical draft unchanged)."""

    instruction: str = ""
    draft: AgreementDraftCreate
    rendered_html: str = ""
    proposer_id: str = ""
    proposer_display_name: str = ""


class RecipientProposalResolveBody(BaseModel):
    proposal_id: str = ""


class RecipientProposalFinalizeBody(BaseModel):
    proposal_id: str = ""


STAGED_RECIPIENT_PROPOSALS_KEY = "staged_recipient_proposals"


class RecipientApproveBody(BaseModel):
    message: str = ""
    participant_id: str = ""
    participant_display_name: str = ""


class SigningCeremonyStartBody(BaseModel):
    participant_id: str = ""


class SigningCeremonyCompleteBody(BaseModel):
    participant_id: str = ""
    typed_name: str = ""
    locked_version_id: str = ""


class SigningInviteTargetBody(BaseModel):
    email: str = ""
    display_name: str = ""
    signing_url: str = ""
    signer_role_id: str = ""
    is_owner: bool = False


class SigningLinksSentBody(BaseModel):
    packet_revision: Optional[str] = None
    document_id: Optional[str] = None
    portable_packet: Optional[Dict[str, Any]] = None
    targets: List[SigningInviteTargetBody] = Field(default_factory=list)


class Vs01SignerCompleteBody(BaseModel):
    signer_role_id: str = ""
    participant_id: str = ""
    document_id: str = ""
    display_name: str = ""
    signed_at: Optional[str] = None
    signed_date_iso: str = ""
    signed_date_display: str = ""
    portable_packet: Optional[Dict[str, Any]] = None


class ReviewRecipientEmailCorrectBody(BaseModel):
    participant_id: str = ""
    new_email: str = ""
    resend_invite: bool = True


class SigningRecipientEmailCorrectBody(BaseModel):
    participant_id: str = ""
    new_email: str = ""
    signer_role_id: Optional[str] = None
    signing_url: Optional[str] = None
    resend_invite: bool = True


class RecipientInviteResendBody(BaseModel):
    phase: str = ""
    participant_id: str = ""
    signing_url: Optional[str] = None
    signer_role_id: Optional[str] = None


NegotiationPosture = Literal[
    "cooperative",
    "firm",
    "protective",
    "fast_close",
    "founder_friendly",
    "investor_friendly",
]

# Prepended to suggestion-generation (same LLM path; shapes output only).
NEGOTIATION_POSTURE_GUIDANCE: Dict[str, str] = {
    "cooperative": (
        "Cooperative: preserve goodwill; seek compromise; use soft, balanced language; "
        "prefer mutual concessions."
    ),
    "firm": (
        "Firm: clear boundaries; direct language; emphasize original intent and commercial reasonableness; "
        "limit unnecessary concessions."
    ),
    "protective": (
        "Protective: minimize risk; preserve termination, payment protection, liability limits, and confidentiality; "
        "push back on scope creep and vague obligations."
    ),
    "fast_close": (
        "Fast-close: prioritize speed to signature; minimize friction; favor simple, practical compromises; "
        "avoid extra complexity."
    ),
    "founder_friendly": (
        "Founder-friendly: favor operator flexibility; protect cash, control, IP, and exit flexibility; "
        "resist overcommitment."
    ),
    "investor_friendly": (
        "Investor-friendly: favor clarity, control rights, reporting, downside protection, auditability, "
        "and stronger enforcement hooks; more formal tone."
    ),
}


class NegotiateAssistRequest(BaseModel):
    """Assist-only: summaries and strategy options. Mutations still use POST /revise."""

    mode: Literal["summary", "options", "both"] = "both"
    recipient_instruction: str = ""
    prior_snapshot: Optional[Dict[str, Any]] = None
    current_snapshot: Optional[Dict[str, Any]] = None
    negotiation_posture: NegotiationPosture = "cooperative"
    session_type: SessionType = "owner"
    ai_model_class: Literal["basic", "premium"] = "basic"


NegotiationRiskTier = Literal["low_risk", "economic_impact", "manual_legal_review"]

_RISK_TIER_ORDER: Dict[str, int] = {
    "low_risk": 0,
    "economic_impact": 1,
    "manual_legal_review": 2,
}

_RISK_TIER_LABEL: Dict[str, str] = {
    "low_risk": "Low risk",
    "economic_impact": "Economic impact",
    "manual_legal_review": "Manual legal review",
}

_RISK_TIER_EXPLANATION: Dict[str, str] = {
    "low_risk": "This appears administrative or clarifying and may be low-risk to accept.",
    "economic_impact": "This changes the deal economics and should be reviewed carefully before accepting.",
    "manual_legal_review": (
        "This may affect legal rights or enforceability and likely needs manual legal review."
    ),
}

_RISK_TIER_HELPER: Dict[str, str] = {
    "low_risk": "You may be able to accept this quickly if it matches your intent.",
    "economic_impact": "Review pricing, scope, timing, and tradeoffs before responding.",
    "manual_legal_review": "Consider legal review before accepting this change.",
}

# Heuristic keyword buckets (informational triage only — not legal advice).
_RISK_LEGAL_KWS = (
    "indemnity",
    "indemnif",
    "liability",
    "limitation of liability",
    "cap on liability",
    "venue",
    "jurisdiction",
    "governing law",
    "choice of law",
    "arbitration",
    "ip ",
    "intellectual property",
    "confidential",
    "exclusiv",
    "non-compete",
    "non compete",
    "noncompete",
    "injunction",
    "compliance",
    "regulatory",
    "representations and warranties",
    "warranty",
)

_RISK_ECON_KWS = (
    "payment",
    "fee",
    "fees",
    "compensation",
    "milestone",
    "refund",
    "royalty",
    "revenue share",
    "rev share",
    "equity",
    "ownership",
    "price",
    "pricing",
    "penalt",
    "liquidated damages",
    "late fee",
    "invoice",
    "retainer",
    "deposit",
    "discount",
    "reimburs",
)

# “Low risk” hints — only persuasive if no legal/econ signals (or very weak).
_RISK_LOW_KWS = (
    "typo",
    "wording",
    "formatting",
    "clarif",
    "defined term",
    "notice address",
    "contact info",
    "administrative",
    "caption",
    "header",
    "numbering",
)


def _negotiate_risk_heuristic(instruction: str) -> Dict[str, Any]:
    """Rules-first risk triage. Always returns a complete risk_assessment dict."""
    text = (instruction or "").strip().lower()
    if not text:
        return {
            "tier": "low_risk",
            "label": _RISK_TIER_LABEL["low_risk"],
            "explanation": "No change text was provided; triage is informational only.",
            "rationale": "Unable to classify an empty request.",
            "helper_text": "Review the full revision in context before responding.",
            "confidence": "low",
        }

    legal_hits = [k for k in _RISK_LEGAL_KWS if k in text]
    econ_hits = [k for k in _RISK_ECON_KWS if k in text]
    low_hits = [k for k in _RISK_LOW_KWS if k in text]

    if legal_hits:
        tier = "manual_legal_review"
        rationale = (
            f"Mentions themes often tied to rights or enforceability (e.g. {', '.join(legal_hits[:3])})."
            if len(legal_hits) <= 3
            else "Touches topics often associated with legal rights, limits, or enforceability."
        )
    elif econ_hits:
        tier = "economic_impact"
        rationale = (
            f"Touches commercial or payment-related themes (e.g. {', '.join(econ_hits[:3])})."
            if len(econ_hits) <= 3
            else "Appears to engage payment, pricing, or commercial exposure."
        )
    elif low_hits and not econ_hits and not legal_hits:
        tier = "low_risk"
        rationale = (
            "Language suggests clarifying, formatting, or light administrative edits."
            if low_hits
            else "No strong economics or legal-structure keywords detected."
        )
    else:
        tier = "economic_impact"
        rationale = "No clear low-risk or legal-structure keywords; defaulting to a cautious commercial review."

    strength = len(legal_hits) + len(econ_hits) + len(low_hits)
    confidence: Literal["low", "medium", "high"]
    if strength >= 3 or (legal_hits and len(legal_hits) >= 2):
        confidence = "high"
    elif strength >= 1:
        confidence = "medium"
    else:
        confidence = "low"

    return {
        "tier": tier,
        "label": _RISK_TIER_LABEL[tier],
        "explanation": _RISK_TIER_EXPLANATION[tier],
        "rationale": rationale[:280],
        "helper_text": _RISK_TIER_HELPER[tier],
        "confidence": confidence,
    }


def _normalize_risk_tier(raw: Any) -> Optional[str]:
    s = str(raw or "").strip().lower().replace("/", "_").replace(" ", "_")
    aliases = {
        "lowrisk": "low_risk",
        "low_risk": "low_risk",
        "economic": "economic_impact",
        "economic_impact": "economic_impact",
        "material_economics": "economic_impact",
        "legal": "manual_legal_review",
        "manual_legal_review": "manual_legal_review",
        "legal_review": "manual_legal_review",
    }
    return aliases.get(s) if s in aliases else (s if s in _RISK_TIER_ORDER else None)


def _merge_risk_assessment(heuristic: Dict[str, Any], llm_raw: Optional[Dict[str, Any]]) -> Dict[str, Any]:
    """Escalate to the more severe tier; blend rationale and optional LLM copy."""
    if not llm_raw or not isinstance(llm_raw, dict):
        return dict(heuristic)

    t_h = str(heuristic.get("tier") or "low_risk")
    t_llm = _normalize_risk_tier(llm_raw.get("tier")) or t_h
    if t_llm not in _RISK_TIER_ORDER:
        t_llm = t_h
    pick = t_h if _RISK_TIER_ORDER[t_h] >= _RISK_TIER_ORDER[t_llm] else t_llm

    rationale = ""
    if pick == t_llm:
        rationale = str(llm_raw.get("rationale") or "").strip()
    if not rationale and pick == t_h:
        rationale = str(heuristic.get("rationale") or "").strip()
    if not rationale:
        rationale = str(heuristic.get("rationale") or llm_raw.get("rationale") or "").strip()

    expl = str(llm_raw.get("explanation") or "").strip() if pick == t_llm else ""
    if not expl:
        expl = _RISK_TIER_EXPLANATION.get(pick, _RISK_TIER_EXPLANATION["low_risk"])
    hel = str(llm_raw.get("helper_text") or "").strip() if pick == t_llm else ""
    if not hel:
        hel = _RISK_TIER_HELPER.get(pick, _RISK_TIER_HELPER["low_risk"])

    conf = str(llm_raw.get("confidence") or heuristic.get("confidence") or "medium").lower()
    if conf not in ("low", "medium", "high"):
        conf = "medium"

    return {
        "tier": pick,
        "label": _RISK_TIER_LABEL.get(pick, pick),
        "explanation": expl[:400],
        "rationale": rationale[:280],
        "helper_text": hel[:400],
        "confidence": conf,
    }


def canonicalize_agreement(draft: Dict[str, Any]) -> str:
    return canon_json_bytes(draft).decode("utf-8")


def _extract_json_object(text: str) -> Dict[str, Any]:
    raw = (text or "").strip()
    if raw.startswith("```"):
        lines = raw.splitlines()
        if lines:
            lines = lines[1:]
        if lines and lines[-1].strip().startswith("```"):
            lines = lines[:-1]
        raw = "\n".join(lines).strip()
    if not raw.startswith("{") or not raw.endswith("}"):
        start = raw.find("{")
        end = raw.rfind("}")
        if start >= 0 and end > start:
            raw = raw[start : end + 1]
    parsed = json.loads(raw)
    if not isinstance(parsed, dict):
        raise ValueError("invalid_json_object")
    return parsed


def _normalize_parsed_draft(raw: Dict[str, Any]) -> AgreementDraftCreate:
    parties_in = raw.get("parties") if isinstance(raw.get("parties"), list) else []
    tmp: List[AgreementParty] = []
    for p in parties_in:
        if not isinstance(p, dict):
            continue
        name = str(p.get("name") or "").strip()
        role = str(p.get("role") or "party").strip() or "party"
        pid = str(p.get("id") or "").strip()
        email = str(p.get("email") or "").strip() or None
        phone = str(p.get("phone") or "").strip() or None
        tmp.append(AgreementParty(name=name, role=role, id=pid or None, email=email, phone=phone))
    parties = _ensure_agreement_parties_have_ids(tmp)
    due_date = str(raw.get("due_date") or "").strip() or None
    duration = str(raw.get("duration") or "").strip() or None
    if due_date and not duration:
        duration = f"until {due_date}"
    jurisdiction = str(raw.get("jurisdiction") or "").strip() or "TBD"
    effective_date = str(raw.get("effective_date") or "").strip() or None
    return AgreementDraftCreate(
        title=str(raw.get("title") or "").strip(),
        jurisdiction=jurisdiction,
        parties=parties,
        purpose=str(raw.get("purpose") or "").strip(),
        payment_terms=str(raw.get("payment_terms") or "").strip(),
        duration=duration,
        due_date=due_date,
        effective_date=effective_date,
    )


_PREMIUM_FAMILY_HINTS = frozenset(
    {
        "generic",
        "nda",
        "operating_agreement",
        "services",
        "partnership",
        "family_financial",
    }
)
_PREMIUM_CONFIDENCE = frozenset({"low", "medium", "high"})


def _parse_premium_intake_result(raw: Dict[str, Any]) -> Tuple[AgreementDraftCreate, AgreementParseExtract]:
    """
    Build validated draft from LLM JSON; pull optional extract keys without letting unknown keys break draft validation.
    """
    draft = _normalize_parsed_draft(raw)
    material_asks: List[str] = []
    ma = raw.get("material_asks")
    if isinstance(ma, list):
        for x in ma:
            s = str(x).strip() if x is not None else ""
            if s and len(material_asks) < 8:
                material_asks.append(s)
    hint_raw = raw.get("agreement_family_hint")
    hint_s: Optional[str] = None
    if hint_raw is not None:
        h = str(hint_raw).strip().lower()
        if h in _PREMIUM_FAMILY_HINTS:
            hint_s = h
    conf_raw = raw.get("confidence")
    conf_s: Optional[str] = None
    if conf_raw is not None:
        c = str(conf_raw).strip().lower()
        if c in _PREMIUM_CONFIDENCE:
            conf_s = c
    extract = AgreementParseExtract(
        material_asks=material_asks,
        agreement_family_hint=hint_s,
        confidence=conf_s,
    )
    return draft, extract


def _heuristic_parse_intake(intake_text: str) -> AgreementDraftCreate:
    t = (intake_text or "").strip()
    title = ""
    jurisdiction = "TBD"
    purpose = t
    payment_terms = ""
    duration = None
    due_date = None
    effective_date = None

    between_match = re.search(r"between\s+(.+?)\s+and\s+(.+?)(?:\.|,|$)", t, re.I)
    parties: List[AgreementParty] = []
    if between_match:
        p1 = between_match.group(1).strip(" ,.")
        p2 = between_match.group(2).strip(" ,.")
        if p1:
            parties.append(AgreementParty(name=p1, role="party_a"))
        if p2:
            parties.append(AgreementParty(name=p2, role="party_b"))
    for rx in [
        r"(?:agreement title is|title is|title:)\s*([^.,\n]+)",
        r"(?:called)\s+([^.,\n]+agreement)",
    ]:
        m = re.search(rx, t, re.I)
        if m:
            title = m.group(1).strip(" \"'")
            break
    if not title:
        title = "Agreement Draft"
    # Only accept jurisdiction when explicitly tied to law / formation (avoid "in Fitness Niche" style captures).
    j = re.search(
        r"\b(?:governed\s+by|governing\s+law|laws?\s+of|jurisdiction|venue\s+in)\s*[:\s]+"
        r"([A-Z][a-z]+(?:\s+[A-Z][a-z]+)?)\b",
        t,
        re.I,
    )
    if j:
        jurisdiction = j.group(1).strip(" ,.")
    else:
        m_state = re.search(
            r"\b(?:State\s+of|Commonwealth\s+of)\s+([A-Z][a-z]+(?:\s+[A-Z][a-z]+)?)\b",
            t,
        )
        if m_state:
            jurisdiction = f"State of {m_state.group(1).strip()}"
    money_terms = re.findall(r"\$[\d,]+(?:\s*(?:per|\/)\s*(?:week|month|year))?", t, re.I)
    money_k = re.findall(r"\$?\d+(?:\.\d+)?k(?:/(?:month|week|year|mo))?", t, re.I)
    merged_money = list(dict.fromkeys([*(money_terms or []), *(money_k or [])]))
    if merged_money:
        payment_terms = "; ".join(merged_money)
    due_match = re.search(r"(?:due|by)\s+([A-Za-z]+\s+\d{1,2},\s*\d{4})", t, re.I)
    if due_match:
        due_date = due_match.group(1).strip()
        duration = f"until {due_date}"
    return AgreementDraftCreate(
        title=title,
        jurisdiction=jurisdiction,
        parties=parties,
        purpose=purpose,
        payment_terms=payment_terms,
        duration=duration,
        due_date=due_date,
        effective_date=effective_date,
    )


def _watermark_active_for_agreement(agreement_id: str) -> bool:
    return bool(economics_overlay_for_agreement(agreement_id).get("watermark_required"))


def _collapse_duplicate_watermark_labels(text: str, label: str) -> str:
    """If the draft body already contains repeated watermark footers, keep a single copy."""
    if not text or not label:
        return text
    escaped = re.escape(label)
    pattern = re.compile(f"(?:{escaped})(?:\\s+{escaped})+")
    return pattern.sub(label, text)


def _strip_watermark_label_from_body(text: str, label: str) -> str:
    """
    Remove LawDog draft footer text from the operative body before HTML/PDF render.

    When economics watermark is on, the label is re-emitted once in a dedicated document-flow
    footer so PyMuPDF Story does not paint an absolute overlay across signature blocks.
    """
    if not text or not label:
        return text
    t = text.replace(label, "")
    t = re.sub(r"[ \t]+\n", "\n", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def _html_watermark_footer(wm_safe: str) -> str:
    """
    Non-absolute footer block: flows after the agreement body so HTML→PDF engines (e.g. Story)
    cannot composite it over mid-page signature / DEVELOPER content.
    """
    return (
        "<footer class='ldg-draft-footer' "
        "style='margin-top:36pt;padding-top:12pt;border-top:1px solid #cbd5e1;font-size:10pt;color:#64748b;"
        "text-align:center;line-height:1.35;break-inside:avoid;page-break-inside:avoid'>"
        f"{wm_safe}</footer>"
    )


def _purpose_looks_like_full_client_agreement_text(purpose: str) -> bool:
    """When the create-flow document editor persists the full preview into `purpose`, render it as the body (not nested in the short-form template)."""
    t = (purpose or "").strip()
    # LawDog Pro / paid full drafts: long operative bodies must not be squeezed into the short starter template.
    if len(t) >= 2400:
        return True
    if len(t) < 240:
        return False
    low = t.lower()
    if "this draft agreement preview is generated from your structured fields" in low:
        return True
    if "this draft llc operating agreement preview is generated" in low:
        return True
    if "by and between" in low and ("\n2. payment terms\n" in low or "\n2. payment terms\r\n" in low):
        return True
    return False


def _render_html(draft: AgreementDraft, *, watermark: bool = False) -> str:
    review_first_corpus, review_first_source = _review_first_final_corpus_from_draft(draft)
    if review_first_corpus:
        log.info(
            "[review-first-reviewer-corpus-selected] agreement_id_short=%s source=%s len=%s hash=%s",
            draft.id[:8],
            review_first_source,
            len(review_first_corpus),
            _review_first_corpus_hash(review_first_corpus),
        )
        purpose_for_body = (
            _strip_watermark_label_from_body(
                _collapse_duplicate_watermark_labels(review_first_corpus, WATERMARK_LABEL),
                WATERMARK_LABEL,
            )
            if watermark
            else review_first_corpus
        )
        body = html.escape(purpose_for_body)
        wm = html.escape(WATERMARK_LABEL)
        article = (
            "<article style='position:relative;max-width:720px;margin:0 auto'>"
            "<p style='text-align:center;color:#475569;font-size:12px;margin-bottom:12px'>"
            "Draft Agreement (non-binding template)</p>"
            "<pre style='white-space:pre-wrap;font-family:Georgia,serif;font-size:15px;line-height:1.65;"
            "color:#0f172a;margin:0;padding:0;border:0;background:transparent'>"
            f"{body}</pre>"
            "</article>"
        )
        return article if not watermark else f"{article}{_html_watermark_footer(wm)}"

    purpose_raw = (draft.purpose or "").strip()
    if _purpose_looks_like_full_client_agreement_text(purpose_raw):
        if watermark:
            purpose_for_body = _strip_watermark_label_from_body(
                _collapse_duplicate_watermark_labels(purpose_raw, WATERMARK_LABEL),
                WATERMARK_LABEL,
            )
        else:
            purpose_for_body = purpose_raw
        body = html.escape(purpose_for_body)
        wm = html.escape(WATERMARK_LABEL)
        article = (
            "<article style='position:relative;max-width:720px;margin:0 auto'>"
            "<p style='text-align:center;color:#475569;font-size:12px;margin-bottom:12px'>"
            "Draft Agreement (non-binding template)</p>"
            "<pre style='white-space:pre-wrap;font-family:Georgia,serif;font-size:15px;line-height:1.65;"
            "color:#0f172a;margin:0;padding:0;border:0;background:transparent'>"
            f"{body}</pre>"
            + (
                ""
                if re.search(r"\bIN WITNESS WHEREOF\b", purpose_for_body, flags=re.I)
                else (
                    "<p style='margin-top:18px;font-size:12px;color:#475569;text-align:center'>"
                    "Execution and signature placement are handled in the electronic signing step."
                    "</p>"
                )
            )
            + "</article>"
        )
        if not watermark:
            return article
        return f"{article}{_html_watermark_footer(wm)}"

    title = html.escape((draft.title or "").strip() or "Agreement")
    jurisdiction_raw = (draft.jurisdiction or "").strip() or "TBD"
    jurisdiction = html.escape(normalize_jurisdiction_display(jurisdiction_raw) or "TBD")
    effective_date = html.escape((draft.effective_date or "").strip() or "TBD")
    purpose = html.escape(purpose_raw or "TBD")
    payment_terms = html.escape((draft.payment_terms or "").strip() or "TBD")
    duration = html.escape((draft.duration or "").strip() or "TBD")
    due_date = html.escape((draft.due_date or "").strip() or "TBD")
    (party_a_name_raw, party_a_role_raw), (party_b_name_raw, party_b_role_raw) = _party_display_names_role(
        list(draft.parties or [])
    )
    party_a_name = html.escape(party_a_name_raw)
    party_b_name = html.escape(party_b_name_raw)
    party_a_role = html.escape(party_a_role_raw)
    party_b_role = html.escape(party_b_role_raw)
    wm = html.escape(WATERMARK_LABEL)
    article = (
        "<article style='position:relative'>"
        f"<h1 style='text-align:center;margin-bottom:6px'>{title}</h1>"
        "<p style='text-align:center;margin-top:0;color:#475569'>Draft Agreement (non-binding template)</p>"
        f"<p>This {title} (the \"Agreement\") is made effective as of {effective_date}, by and between "
        f"{party_a_name} ({party_a_role}) and {party_b_name} ({party_b_role}). The parties agree as follows:</p>"
        "<h2>1. Scope of Services</h2>"
        f"<p>{purpose}. The service provider will perform the services in a professional and workmanlike manner and "
        "will keep the client reasonably informed regarding project progress.</p>"
        "<h2>2. Compensation</h2>"
        f"<p>In consideration for the services, compensation is as follows: {payment_terms}.</p>"
        "<h2>3. Payment Terms</h2>"
        f"<p>Payments are due according to the agreed schedule. If applicable, final delivery is expected by {due_date}. "
        "Late payments may be subject to commercially reasonable collection procedures.</p>"
        "<h2>4. Term and Termination</h2>"
        f"<p>This Agreement begins on the effective date and remains in effect for {duration}, unless earlier terminated "
        "by either party on written notice for material breach or as otherwise agreed in writing.</p>"
        "<h2>5. Confidentiality</h2>"
        "<p>Each party shall keep confidential non-public business and technical information disclosed by the other party "
        "and shall use such information solely for performance under this Agreement.</p>"
        "<h2>6. Independent Contractor</h2>"
        "<p>The parties agree that the service provider is an independent contractor and not an employee, partner, or "
        "agent of the client, except as expressly authorized in writing.</p>"
        "<h2>7. Governing Law</h2>"
        f"<p>This Agreement is governed by the laws of {jurisdiction}, without regard to conflict of law principles.</p>"
        "<h2>8. Signatures</h2>"
        "<p>IN WITNESS WHEREOF, the parties have executed this Agreement as of the effective date.</p>"
        "<table style='width:100%;margin-top:16px;border-collapse:collapse'>"
        "<tr>"
        f"<td style='width:50%;padding-right:12px'><div style='border-bottom:1px solid #64748b;height:28px'></div><div style='font-size:12px;color:#475569'>"
        f"{party_a_name} Signature</div><div style='margin-top:8px;border-bottom:1px solid #cbd5e1;height:20px'></div><div style='font-size:12px;color:#475569'>Date</div></td>"
        f"<td style='width:50%;padding-left:12px'><div style='border-bottom:1px solid #64748b;height:28px'></div><div style='font-size:12px;color:#475569'>"
        f"{party_b_name} Signature</div><div style='margin-top:8px;border-bottom:1px solid #cbd5e1;height:20px'></div><div style='font-size:12px;color:#475569'>Date</div></td>"
        "</tr>"
        "</table>"
        "</article>"
    )
    if not watermark:
        return article
    return f"{article}{_html_watermark_footer(wm)}"


def _norm_revision_comparison_text(s: str) -> str:
    return " ".join((s or "").split()).lower()


def _revision_comparison_blob(d: AgreementDraft | AgreementDraftCreate) -> str:
    """Whitespace-normalized fingerprint of revision-relevant fields (for similarity checks)."""
    parties: List[Dict[str, str]] = []
    for p in d.parties or []:
        parties.append(
            {
                "name": _norm_revision_comparison_text(p.name),
                "role": _norm_revision_comparison_text(p.role),
            }
        )
    parties.sort(key=lambda x: (x["role"], x["name"]))
    bucket = {
        "due_date": _norm_revision_comparison_text(d.due_date or ""),
        "duration": _norm_revision_comparison_text(d.duration or ""),
        "effective_date": _norm_revision_comparison_text(d.effective_date or ""),
        "jurisdiction": _norm_revision_comparison_text(d.jurisdiction or ""),
        "parties": parties,
        "payment_terms": _norm_revision_comparison_text(d.payment_terms or ""),
        "purpose": _norm_revision_comparison_text(d.purpose or ""),
        "title": _norm_revision_comparison_text(d.title or ""),
    }
    return json.dumps(bucket, sort_keys=True, separators=(",", ":"), ensure_ascii=False)


def _revision_validation_corpus_text(revised: AgreementDraft | AgreementDraftCreate) -> str:
    """Lowercased prose used for lightweight post-revision checks (no full HTML render)."""
    parts = [
        revised.title or "",
        revised.jurisdiction or "",
        revised.purpose or "",
        revised.payment_terms or "",
    ]
    if getattr(revised, "duration", None):
        parts.append(str(revised.duration))
    if getattr(revised, "due_date", None):
        parts.append(str(revised.due_date))
    if getattr(revised, "effective_date", None):
        parts.append(str(revised.effective_date))
    return " ".join(parts).lower()


def _has_any(text: str, phrases: List[str]) -> bool:
    return any(phrase in text for phrase in phrases)


def _validate_revision_expectations(
    base: AgreementDraft,
    revised: AgreementDraftCreate,
    instruction: str,
) -> Dict[str, Any]:
    """
    Best-effort checks that explicit user asks appear in the coalesced draft prose.
    Non-blocking. Result is returned as JSON key ``revision_validation`` on revise/commit
    (equivalent to tagging the coalesced draft for the client; not stored on AgreementDraft).
    """
    issues: List[str] = []
    text = _revision_validation_corpus_text(revised)
    instr = (instruction or "").lower()

    # Cure / remedy — only when instruction clearly asks for it.
    instr_wants_cure = _has_any(
        instr,
        [
            "cure period",
            "cure within",
            "days to cure",
            "day to cure",
            "opportunity to cure",
            "time to cure",
            "right to cure",
            "days to remedy",
            "day to remedy",
            "remedy breach",
            "cure for breach",
            "cure a breach",
            "add cure",
            "include a cure",
            "cure before",
            "cure prior",
        ],
    ) or ("breach" in instr and "cure" in instr)
    if instr_wants_cure:
        cure_ok = _has_any(
            text,
            [
                "cure period",
                "cure within",
                "days to cure",
                "days to remedy",
                "opportunity to cure",
                "remedy breach within",
                "five business day",
                "5 business day",
                "five (5) business day",
                "5-day cure",
            ],
        )
        if not cure_ok:
            issues.append("missing_cure_period")

    # Non-disparagement — instruction mentions disparagement / non-disparagement.
    instr_wants_nd = _has_any(
        instr,
        [
            "non-disparagement",
            "non disparagement",
            "nondisparagement",
            "disparagement",
        ],
    )
    if instr_wants_nd:
        nd_ok = _has_any(
            text,
            [
                "non-disparagement",
                "non disparagement",
                "not disparage",
                "shall not disparage",
                "disparag",
            ],
        )
        if not nd_ok:
            issues.append("missing_non_disparagement")

    # ~45-day timeline — instruction asks for 45 / forty-five day delivery.
    instr_wants_45 = bool(re.search(r"\b45\s*days?\b", instr)) or _has_any(
        instr, ["forty-five days", "forty five days", "forty-five day", "forty five day"]
    )
    if instr_wants_45:
        timeline_ok = _has_any(
            text,
            [
                "45 days",
                "forty-five days",
                "forty five days",
                "within 45",
                "within forty-five",
                "within forty five",
            ],
        ) or bool(re.search(r"\b45\s*days?\b", text))
        if not timeline_ok:
            issues.append("timeline_not_updated")

    base_j = (base.jurisdiction or "").strip()
    rev_j = (revised.jurisdiction or "").strip()
    if base_j and not rev_j:
        issues.append("jurisdiction_dropped")

    base_pt = (base.payment_terms or "").strip()
    rev_pt = (revised.payment_terms or "").strip()
    if base_pt and not rev_pt:
        issues.append("payment_terms_dropped")

    return {"ok": len(issues) == 0, "issues": issues}


def agreement_revision_similarity(
    before: AgreementDraft | AgreementDraftCreate,
    after: AgreementDraft | AgreementDraftCreate,
) -> float:
    """Similarity ratio in [0, 1]; 1.0 means normalized revision blobs match."""
    sa = _revision_comparison_blob(before)
    sb = _revision_comparison_blob(after)
    if not sa and not sb:
        return 1.0
    return SequenceMatcher(None, sa, sb).ratio()


def _coalesce_revision_draft_with_base(
    base: AgreementDraft,
    revised: AgreementDraftCreate,
) -> AgreementDraftCreate:
    """
    After LLM revise, restore identity the model often drops: party emails/phones, empty scalar fields.
    Index-aligned party merge; extra parties from either side are kept when the other side lacks that slot.
    """
    b_parties = list(base.parties or [])
    r_parties = list(revised.parties or [])
    n = max(len(b_parties), len(r_parties))
    merged_parties: List[AgreementParty] = []
    for i in range(n):
        bp = b_parties[i] if i < len(b_parties) else None
        rp = r_parties[i] if i < len(r_parties) else None
        if bp is None and rp is None:
            continue
        if bp is None and rp is not None:
            merged_parties.append(rp)
            continue
        if rp is None and bp is not None:
            merged_parties.append(bp)
            continue
        assert bp is not None and rp is not None
        name = (rp.name or "").strip() or (bp.name or "").strip() or ("Party A" if i == 0 else "Party B")
        role = (rp.role or "").strip() or (bp.role or "").strip() or "party"
        pid = (rp.id or "").strip() or (bp.id or "").strip() or None
        em_r = (rp.email or "").strip()
        em_b = (bp.email or "").strip()
        email = em_r or em_b or None
        ph_r = (rp.phone or "").strip()
        ph_b = (bp.phone or "").strip()
        phone = ph_r or ph_b or None
        merged_parties.append(AgreementParty(name=name, role=role, id=pid, email=email, phone=phone))

    def _pick_text(rev: str, cur: str) -> str:
        r = (rev or "").strip()
        if r:
            return rev
        c = (cur or "").strip()
        return cur if c else rev

    def _pick_opt(rev: Optional[str], cur: Optional[str]) -> Optional[str]:
        if rev is not None and str(rev).strip():
            return rev
        return cur

    jurisdiction = _pick_text(revised.jurisdiction or "", base.jurisdiction or "")
    if not jurisdiction.strip():
        jurisdiction = "TBD"
    title = _pick_text(revised.title or "", base.title or "")
    purpose = _pick_text(revised.purpose or "", base.purpose or "")
    payment_terms = _pick_text(revised.payment_terms or "", base.payment_terms or "")
    duration = _pick_opt(revised.duration, base.duration)
    due_date = _pick_opt(revised.due_date, base.due_date)
    effective_date = _pick_opt(revised.effective_date, base.effective_date)

    # Preview/commit bodies often omit feed + payment_request; do not reset stored agreement state.
    payment_request = (
        revised.payment_request if revised.payment_request is not None else base.payment_request
    )
    payment_required = (
        revised.payment_required
        if revised.payment_request is not None
        else getattr(base, "payment_required", False)
    )

    return AgreementDraftCreate(
        title=title,
        jurisdiction=jurisdiction,
        parties=_ensure_agreement_parties_have_ids(merged_parties),
        purpose=purpose,
        payment_terms=payment_terms,
        duration=duration,
        due_date=due_date,
        effective_date=effective_date,
        feed_visibility=base.feed_visibility,
        feed_party_anonymize=base.feed_party_anonymize,
        feed_show_financial_summary=base.feed_show_financial_summary,
        feed_anchor_network=base.feed_anchor_network,
        payment_request=payment_request,
        payment_required=payment_required,
    )


def _revise_system_prompt_starter_preview() -> str:
    return (
        "You are a structured agreement editor for CLAW (STARTER_PREVIEW mode).\n"
        "Update the agreement state based on the user's edit instruction.\n"
        "Return ONLY strict JSON matching:\n"
        '{ "title":"", "jurisdiction":"", "parties":[{"name":"","role":"","email":null,"phone":null}], '
        '"purpose":"", "payment_terms":"", "duration":null, "due_date":null, "effective_date":null }\n'
        "Rules:\n"
        "- Preserve existing values unless user explicitly changes them.\n"
        "- Keep response concise and valid JSON.\n"
        "- Do not add commentary."
    )


def _revise_system_prompt_premium_surgical() -> str:
    return (
        "You are CLAW's premium agreement editor (PREMIUM_SURGICAL mode — negotiation-style redlines).\n"
        "Return ONLY strict JSON matching:\n"
        '{ "title":"", "jurisdiction":"", "parties":[{"name":"","role":"","email":null,"phone":null}], '
        '"purpose":"", "payment_terms":"", "duration":null, "due_date":null, "effective_date":null }\n'
        "MINIMAL CHANGE (default expectation):\n"
        "- Preserve original structure, clause numbering, headings, and sentence order inside each JSON prose field "
        "whenever possible.\n"
        "- Keep **unchanged language verbatim**; do not paraphrase sentences the user did not ask to change.\n"
        "- Prefer **insertions and short appended sentences** over replacing whole paragraphs or entire payment sections.\n"
        "- Modify only text **directly implicated** by the user's instruction; do not restate surrounding deal context.\n"
        "- Avoid injected commentary, “explainer” paragraphs, or commercial rewrites beyond the ask.\n"
        "Material facts: preserve party names and roles, jurisdiction, stated amounts and cadence, dates, durations, "
        "and existing obligations unless the instruction clearly changes them.\n"
        "Do not invent new parties, dollar amounts, dates, or jurisdictions absent from the current draft or instruction.\n"
        "Multi-part instructions: implement each item with the smallest edit that satisfies it.\n"
        "In the output JSON `parties` array, echo each party's `email` and `phone` from `current_draft.parties` when those "
        "fields exist (same order). Omitting them wipes client-visible contact data.\n"
        "Do not add commentary outside JSON."
    )


def _revise_system_prompt_premium_rewrite() -> str:
    return (
        "You are CLAW's premium agreement rewriter (PREMIUM_MATERIAL_UPGRADE mode — user asked for a broad polish).\n"
        "Paid users asked for a visibly upgraded, send-ready draft in this pass — not a near-copy of the prior wording.\n"
        "Return ONLY strict JSON matching:\n"
        '{ "title":"", "jurisdiction":"", "parties":[{"name":"","role":"","email":null,"phone":null}], '
        '"purpose":"", "payment_terms":"", "duration":null, "due_date":null, "effective_date":null }\n'
        "User intent is sacred; you may materially improve clarity, grouping, and tone in prose fields.\n"
        "Treat the revision instruction as authorization to materially upgrade the draft (not a minimal line patch).\n"
        "Preserve material facts: party names and roles, jurisdiction, pricing amounts and cadence, dates, durations, "
        "obligations the user already stated, and any bespoke requests in the instruction.\n"
        "Freely improve: vocabulary, sentence structure, clause order inside prose fields, readability, completeness, "
        "stronger commercial defaults, reduced ambiguity, and professional tone.\n"
        "Works for any agreement type (e.g. independent contractor, consulting, recurring services or retainer, NDA, "
        "landlord/tenant lease addendum): purpose and payment_terms should read like polished deal terms — clearer "
        "sections (short headings or bullets inside the JSON string where helpful), tighter obligations, less filler.\n"
        "title: specific and scannable; avoid lazy one-word titles when the deal supports more.\n"
        "Do not invent new parties, dollar amounts, dates, or jurisdictions absent from the current draft or instruction.\n"
        "Numbered / multi-part instructions: implement each clear item; do not drop unrelated commercial terms unless the user asked to remove them.\n"
        "In the output JSON `parties` array, echo each party's `email` and `phone` from `current_draft.parties` when those "
        "fields exist (same order). Omitting them wipes client-visible contact data.\n"
        "Do not add commentary outside JSON."
    )


def _premium_revision_retry_escalation(instruction: str, *, level: int) -> str:
    quoted = json.dumps((instruction or "").strip(), ensure_ascii=False)
    return (
        "\n\n---\nAUTO_REWRITE_ESCALATION (internal — do not echo this header in output): "
        "The last JSON was still too close to the prior draft text. "
        f"Escalation {level}/{PREMIUM_REVISION_MAX_ATTEMPTS}. "
        "Re-output the full JSON with substantially rewritten prose in title, purpose, and payment_terms: "
        "new flow and grouping in purpose, sharper obligations, fewer vague qualifiers, more professional send-ready tone. "
        "Preserve party names, roles, stated amounts, dates, jurisdiction, and commercial intent including: "
        f"{quoted}. "
        "Output must not read like a trivial reorder of the previous version."
    )


def _revise_instruction_fallback(current: AgreementDraft, instruction: str) -> AgreementDraftCreate:
    text = (instruction or "").strip()
    next_data = AgreementDraftCreate(
        title=current.title,
        jurisdiction=current.jurisdiction,
        parties=current.parties,
        purpose=current.purpose,
        payment_terms=current.payment_terms,
        duration=current.duration,
        due_date=current.due_date,
        effective_date=current.effective_date,
    )
    m = re.search(r"(?:change|set|update)\s+payment(?:\s+terms)?\s+to\s+(.+)$", text, re.I)
    if m:
        next_data.payment_terms = m.group(1).strip(" .")
    m = re.search(r"(?:change|set|make|update)\s+(?:term|duration)\s+(?:to|for)?\s+(.+)$", text, re.I)
    if m:
        next_data.duration = m.group(1).strip(" .")
    m = re.search(r"(?:change|set|update)\s+effective\s+date\s+(?:to)?\s+(.+)$", text, re.I)
    if m:
        next_data.effective_date = m.group(1).strip(" .")
    m = re.search(r"(?:change|set|update)\s+jurisdiction\s+(?:to)?\s+(.+)$", text, re.I)
    if m:
        next_data.jurisdiction = m.group(1).strip(" .")
    if re.search(r"\bconfidentiality\b", text, re.I) and "confidential" not in (next_data.purpose or "").lower():
        next_data.purpose = f"{(next_data.purpose or '').strip()}. Includes confidentiality obligations.".strip(" .")
    return next_data


def _append_unique_purpose_sentence(purpose: str, addition: str) -> str:
    clause = (addition or "").strip()
    if not clause:
        return (purpose or "").strip()
    p = (purpose or "").strip()
    head = clause[: min(72, len(clause))].lower()
    if head and head in p.lower():
        return p
    joiner = ". " if p and not p.endswith((".", "?", "!")) else " "
    return f"{p}{joiner}{clause}".strip()


def _recipient_deterministic_merge_instruction(base: AgreementDraft, instruction: str) -> AgreementDraftCreate:
    """
    Apply common reviewer phrases to draft fields without an LLM (used when revise output
    is still field-identical to the stored draft).
    """
    t = (instruction or "").strip()
    lo = t.lower()
    title = base.title or ""
    jurisdiction = base.jurisdiction or ""
    parties = list(base.parties or [])
    purpose = base.purpose or ""
    payment_terms = base.payment_terms or ""
    duration = base.duration
    due_date = base.due_date
    effective_date = base.effective_date

    fb = _revise_instruction_fallback(base, instruction)
    title = fb.title or title
    jurisdiction = fb.jurisdiction or jurisdiction
    parties = list(fb.parties or parties)
    purpose = fb.purpose or purpose
    payment_terms = fb.payment_terms or payment_terms
    duration = fb.duration if fb.duration is not None else duration
    due_date = fb.due_date if fb.due_date is not None else due_date
    effective_date = fb.effective_date if fb.effective_date is not None else effective_date

    # Net N / payment terms (handles "update the payment terms to Net 30", "use net 30", etc.)
    nm = re.search(r"(?:payment\s+terms?\s*(?:to|is|=|:)\s*)?net\s+(\d+)\b", lo, re.I)
    if nm:
        n = nm.group(1)
        if n:
            compact = re.sub(r"[\s_]", "", (payment_terms or "").lower())
            if f"net{n}" not in compact:
                payment_terms = (f"Net {n}. " + (payment_terms or "").strip()).strip()

    # Late payment fee / penalty language
    if re.search(r"late\s*(?:payment)?\s*(?:fee|penalt)|penalt\w*\s+for\s+late|interest\s+on\s+late", lo):
        purpose = _append_unique_purpose_sentence(
            purpose,
            "Late payment may incur fees, interest, or other remedies described in the payment terms.",
        )

    # Pause / suspend work when payment is late (captures "more than 15 days late")
    if re.search(r"\bpause\b.*\bwork\b|\bsuspend\b.*\bwork\b|\bmay\s+pause\b", lo) and (
        "payment" in lo or "invoice" in lo or "paid" in lo
    ):
        dm = re.search(r"more\s+than\s+(\d+)\s*days?\s+late|(\d+)\s*days?\s*(?:past\s*due|late)", lo)
        days = (dm.group(1) or dm.group(2)) if dm else "15"
        purpose = _append_unique_purpose_sentence(
            purpose,
            f"The developer may pause work if payment is more than {days} days late until amounts are brought current.",
        )

    # Delivery / deadline changes
    dvm = re.search(
        r"(?:delivery|deadline)\s*(?:date)?\s*(?:in|within|of|to|=|:)?\s*(\d+)\s*days?",
        lo,
        re.I,
    )
    if dvm:
        purpose = _append_unique_purpose_sentence(
            purpose,
            f"Delivery timeline: {dvm.group(1)} days from the effective date unless otherwise agreed in writing.",
        )

    # Revision rounds
    rvm = re.search(r"(\d+)\s*(?:revision|review)\s*(?:rounds?|cycles?)", lo, re.I)
    if rvm:
        purpose = _append_unique_purpose_sentence(
            purpose,
            f"Up to {rvm.group(1)} revision rounds are included unless the parties agree otherwise.",
        )

    # Governing law (explicit phrase only; stop before "and" / "add" / sentence end)
    if re.search(r"\bgoverning\s+law\b", lo):
        gvm = re.search(
            r"governing\s+law\s*(?:to|is|:)?\s*([A-Za-z][A-Za-z0-9\s]{1,40}?)(?:\s*[.,]|\s+and\b|\s+add\b|$)",
            t,
            re.I,
        )
        if gvm:
            jurisdiction = gvm.group(1).strip(" .")

    # Confidentiality / NDA
    if re.search(r"\bconfidentiality\b|\bnon[-\s]?disclosure\b|\bnda\b", lo, re.I):
        purpose = _append_unique_purpose_sentence(
            purpose,
            "The parties will treat confidential information under commercially reasonable confidentiality obligations.",
        )

    # Support / bug-fix period
    if re.search(r"bug[-\s]?fix|support\s*(?:and\s*maintenance\s*)?period|warranty\s*period|maintenance\s*window", lo, re.I):
        purpose = _append_unique_purpose_sentence(
            purpose,
            "A mutually agreed support and bug-fix period applies as further described in the statement of work.",
        )

    return AgreementDraftCreate(
        title=title,
        jurisdiction=jurisdiction,
        parties=parties,
        purpose=purpose,
        payment_terms=payment_terms,
        duration=duration,
        due_date=due_date,
        effective_date=effective_date,
    )


def _maybe_apply_recipient_deterministic_no_op_patch(
    base: AgreementDraft,
    instruction: str,
    revised: AgreementDraftCreate,
) -> AgreementDraftCreate:
    """
    Recipient-only: when the LLM/coalesced revision is field-identical to the base draft,
    merge deterministic reviewer instruction patterns so previews can still change fields.
    """
    if _revision_comparison_blob(base) != _revision_comparison_blob(revised):
        return revised
    patched = _recipient_deterministic_merge_instruction(base, instruction)
    if _revision_comparison_blob(base) != _revision_comparison_blob(patched):
        return patched
    return revised


def _revise_llm_once(
    current: AgreementDraft,
    instruction: str,
    *,
    trace_context: Optional[Dict[str, Any]],
    ai_model_class: Optional[str],
    system_prompt: str,
    max_tokens: int,
    temperature: float,
    payload_mode: str,
) -> Tuple[AgreementDraftCreate, bool]:
    """Returns (draft, llm_ok). On LLM failure uses deterministic fallback (llm_ok False)."""
    payload: Dict[str, Any] = {
        "mode": payload_mode,
        "instruction": instruction,
        "current_draft": {
            "title": current.title,
            "jurisdiction": current.jurisdiction,
            "parties": [p.model_dump() for p in current.parties],
            "purpose": current.purpose,
            "payment_terms": current.payment_terms,
            "duration": current.duration,
            "due_date": current.due_date,
            "effective_date": current.effective_date,
        },
    }
    try:
        llm_model = resolve_llm_model_for_access_class(ai_model_class)
        llm_text = call_legal_llm(
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
            ],
            model=llm_model,
            max_tokens=max_tokens,
            temperature=temperature,
            trace_context=trace_context,
            airlock_profile="agreement_outbound",
        )
        parsed = _extract_json_object(llm_text)
        return _normalize_parsed_draft(parsed), True
    except Exception:
        return _revise_instruction_fallback(current, instruction), False


def _revise_with_instruction(
    current: AgreementDraft,
    instruction: str,
    trace_context: Optional[Dict[str, Any]] = None,
    ai_model_class: Optional[str] = None,
) -> AgreementDraftCreate:
    tier = (ai_model_class or "").strip().lower()
    premium = tier == "premium"

    if not premium:
        out, _ok = _revise_llm_once(
            current,
            instruction,
            trace_context=trace_context,
            ai_model_class=ai_model_class,
            system_prompt=_revise_system_prompt_starter_preview(),
            max_tokens=350,
            temperature=0.0,
            payload_mode="STARTER_PREVIEW",
        )
        return out

    want_material = instruction_requests_material_rewrite(instruction)
    system_prompt = (
        _revise_system_prompt_premium_rewrite()
        if want_material
        else _revise_system_prompt_premium_surgical()
    )
    payload_mode = "PREMIUM_MATERIAL_UPGRADE" if want_material else "PREMIUM_SURGICAL"
    combined = (instruction or "").strip()
    last: Optional[AgreementDraftCreate] = None
    temps = (0.12, 0.28, 0.42)
    for attempt in range(PREMIUM_REVISION_MAX_ATTEMPTS):
        temperature = temps[min(attempt, len(temps) - 1)]
        revised, llm_ok = _revise_llm_once(
            current,
            combined,
            trace_context=trace_context,
            ai_model_class=ai_model_class,
            system_prompt=system_prompt,
            max_tokens=960,
            temperature=temperature,
            payload_mode=payload_mode,
        )
        last = revised
        if not llm_ok:
            return revised
        sim = agreement_revision_similarity(current, revised)
        if want_material:
            if sim <= PREMIUM_REVISION_SIMILARITY_CEILING:
                return revised
            if attempt < PREMIUM_REVISION_MAX_ATTEMPTS - 1:
                combined = (instruction or "").strip() + _premium_revision_retry_escalation(
                    instruction, level=attempt + 2
                )
            continue
        if not is_overbroad_structured_revision(current, revised, instruction):
            return revised
        if attempt < PREMIUM_REVISION_MAX_ATTEMPTS - 1:
            combined = (instruction or "").strip() + MINIMAL_REVISION_RETRY_SUFFIX
            continue
        break
    assert last is not None
    if not want_material and is_overbroad_structured_revision(current, last, instruction):
        det = _recipient_deterministic_merge_instruction(current, instruction)
        if _revision_comparison_blob(det) != _revision_comparison_blob(current):
            return det
    if not want_material and _revision_comparison_blob(last) == _revision_comparison_blob(current):
        det = _recipient_deterministic_merge_instruction(current, instruction)
        if _revision_comparison_blob(det) != _revision_comparison_blob(current):
            return det
    return last


def _load_or_404(agreement_id: str) -> AgreementDraft:
    try:
        raw = load_draft(agreement_id)
        return AgreementDraft.model_validate(raw)
    except KeyError:
        raise HTTPException(status_code=404, detail="agreement_not_found")


def _load_draft_dict_or_404(agreement_id: str) -> Dict[str, Any]:
    """Load raw draft JSON for read-only surfaces that must not 500 on legacy audit shapes."""
    try:
        raw = load_draft(agreement_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="agreement_not_found")
    except ValueError:
        raise HTTPException(status_code=404, detail="agreement_not_found")
    if not isinstance(raw, dict):
        raise HTTPException(status_code=404, detail="agreement_not_found")
    return raw


def _parse_intake_system_prompt_basic() -> str:
    """Standard create-flow parse (unchanged)."""
    return (
        "You are a structured agreement intake assistant for CLAW.\n"
        "Extract agreement details from the user's intake.\n"
        "Return ONLY strict JSON matching this schema:\n"
        '{ "title":"", "jurisdiction":"", "parties":[{"name":"","role":""}], '
        '"purpose":"", "payment_terms":"", "duration":null, "due_date":null, "effective_date":null }\n'
        "Rules:\n"
        "- If due_date exists and duration missing, set duration = \"until <due_date>\".\n"
        "- If effective_date missing, set null.\n"
        "- If jurisdiction is ambiguous, default to \"TBD\".\n"
        "- Use each party's real legal or business name when known; never use internal placeholder tokens "
        'like ORG_1, PARTY_1, or bracketed variants in names or roles — use party_a / party_b roles or plain words.\n'
        "- Do not add commentary. Do not wrap in markdown. Only JSON."
    )


def _parse_intake_system_prompt_premium() -> str:
    """Premium completion re-parse: same schema, higher judgment (basic prompt unchanged)."""
    return (
        "You are a sharp commercial intake writer for CLAW (PREMIUM path). Optimize for human wow: the reader should "
        "feel \"it actually read my situation\" — clear, specific, practical — not like a law-school outline.\n"
        "Return ONLY strict JSON matching this schema (same top-level fields as the standard path, plus optional fields):\n"
        '{ "title":"", "jurisdiction":"", "parties":[{"name":"","role":""}], "purpose":"", "payment_terms":"", "duration":null, "due_date":null, "effective_date":null, '
        '"material_asks": [], "agreement_family_hint": null, "confidence": null }\n'
        "Optional fields (premium; omit or set null/[] if not applicable):\n"
        '- "material_asks": string array, max 8 short bullets, each grounded in the user intake only; no invented economics or terms.\n'
        '- "agreement_family_hint": one of: "generic" | "nda" | "operating_agreement" | "services" | "partnership" | "family_financial", or null.\n'
        '- "confidence": one of "low" | "medium" | "high", or null.\n'
        "Shared rules (same as standard path):\n"
        "- If due_date exists and duration missing, set duration = \"until <due_date>\".\n"
        "- If effective_date missing, set null.\n"
        "- If jurisdiction is ambiguous, default to \"TBD\".\n"
        "- Use each party's real legal or business name when known; never use internal placeholder tokens "
        'like ORG_1, PARTY_1, or bracketed variants in names or roles — use party_a / party_b roles or plain words.\n'
        "Optional-field rules: material_asks must be concise; if nothing distinct from purpose/payment, use [].\n"
        "Premium-only guidance (JSON only; no markdown; no extra keys except the optional three above):\n"
        "- title: Must be instantly scannable and specific. Name the real transaction (product, role, or risk) using "
        "words from the intake when possible — e.g. \"Marketing Retainer — Q2 Campaign\", \"Beta API Access Agreement\", "
        "\"Contractor Agreement — Warehouse Staffing\". Banned: lone \"Agreement\", \"Contract\", \"Services Agreement\" "
        "without a distinguishing noun, or any title that could describe a thousand unrelated deals.\n"
        "- purpose: Write as situation-specific clauses in flowing prose (still one JSON string): mirror the user's "
        "goals, deliverables, constraints, and vocabulary — reuse their nouns (SKU, API, pilot, equity, venue, launch date) "
        "so it feels bespoke. Cover, when the intake supports it: what each side does, timelines or milestones, "
        "acceptance or success in plain terms, confidentiality / data handling if secrets or customer data appear, "
        "IP for work product when they are paying for creation or customization, non-solicit or exclusivity only if "
        "hiring or partnership language clearly points there, termination or notice if they mentioned ending the "
        "relationship, and liability or indemnity only if they hinted at risk (breach, infringement, injury). "
        "Prefer one concrete sentence per idea; no \"whereas\", no stacked semicolons, no Latin. If a sentence could apply "
        "to any business in the world without change, rewrite it with their details or delete it.\n"
        "- payment_terms: Tie money to their story — retainer vs milestone vs hourly vs commission, invoicing rhythm, "
        "net days, currency, expenses, late fees or caps only if mentioned. If numbers are missing but structure is clear, "
        "describe the structure in plain English without inventing amounts. One honest line if compensation is truly "
        "unsaid; never \"as discussed\", \"TBD\", or \"standard terms\" when any payment hint exists.\n"
        "- Voice: warm-neutral, confident, short sentences, active voice. Sound like a thoughtful colleague summarizing "
        "the deal back to them — not a generic policy memo.\n"
        "- Anti-filler: Do not use lines like \"parties will cooperate in good faith\", \"industry standard\", "
        "\"mutually beneficial\", \"best efforts\" without the user saying so; do not pad with definitions of "
        "\"Confidential Information\" unless they asked for an NDA-style definition.\n"
        "- Hard limits: Do not invent parties, dollar amounts, dates, or jurisdictions absent from the intake; do not "
        "add legal verbosity for show. Wow comes from precision and fit, not length.\n"
        "- Do not add commentary outside JSON. Do not wrap in markdown. Only JSON."
    )


@router.post("/parse", response_model=AgreementParseResponse)
def parse_agreement_intake(request: Request, body: AgreementParseRequest) -> AgreementParseResponse:
    require_claw_org_id_header(request)
    parse_debug_client = os.getenv("CLAW_AGREEMENT_PARSE_CLIENT_DEBUG", "").strip() == "1"
    if body.ai_model_class == "premium":
        system_prompt = _parse_intake_system_prompt_premium()
        parse_max_tokens = 1200
        prompt_label = "premium_v1"
    else:
        system_prompt = _parse_intake_system_prompt_basic()
        parse_max_tokens = 350
        prompt_label = "basic_v1"
    usage_holder: List[Dict[str, Any]] = []
    llm_model = resolve_llm_model_for_access_class(body.ai_model_class)
    tokens_param = "max_completion_tokens" if (llm_model or "").strip().lower().startswith("gpt-5") else "max_tokens"
    try:
        llm_text = call_legal_llm(
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": body.intake_text},
            ],
            model=llm_model,
            max_tokens=parse_max_tokens,
            temperature=0.0,
            usage_sink=usage_holder,
            airlock_profile="agreement_outbound",
        )
        parsed = _extract_json_object(llm_text)
        if body.ai_model_class == "premium":
            draft_out, extract_out = _parse_premium_intake_result(parsed)
        else:
            draft_out = _normalize_parsed_draft(parsed)
            extract_out = None
        ip = request.client.host if request.client else "unknown"
        record_ai_call(subject_ref=resolve_subject_from_request(request), request_ip=ip or "unknown")
        usage = usage_holder[0] if usage_holder else {}
        log.info(
            "agreement_parse ok ai_model_class=%s resolved_model=%s prompt_label=%s tokens=%s",
            body.ai_model_class,
            llm_model,
            prompt_label,
            usage,
        )
        log.info("[premium-api-ok] model=%s tokens_param=%s status=ok", llm_model, tokens_param)
        parse_meta_out: Optional[Dict[str, Any]] = None
        if body.ai_model_class == "premium" or parse_debug_client:
            intake_txt = body.intake_text or ""
            parse_meta_out = {
                "ai_model_class": body.ai_model_class,
                "resolved_model": llm_model,
                "prompt_label": prompt_label,
                "max_output_tokens": parse_max_tokens,
                "usage": usage,
                "premium_system_prompt_loaded": body.ai_model_class == "premium",
                "intake_text_char_len": len(intake_txt),
                "has_exact_word_upgrade_marker": "--- Complete Version: exact wording / notes to apply ---" in intake_txt,
            }
        return AgreementParseResponse(draft=draft_out, parse_meta=parse_meta_out, extract=extract_out)
    except Exception as exc:
        log.warning(
            "agreement_parse failed ai_model_class=%s resolved_model=%s fallback=%s err=%s",
            body.ai_model_class,
            llm_model,
            "none_for_premium" if body.ai_model_class == "premium" else "heuristic_basic",
            type(exc).__name__,
            exc_info=body.ai_model_class != "premium",
        )
        log.warning("[premium-api-fail] model=%s tokens_param=%s error=%s", llm_model, tokens_param, f"{type(exc).__name__}:{exc}")
        if body.ai_model_class == "premium":
            # Never silently downgrade premium completion to heuristic/basic-quality parse.
            log.error(
                "premium_parse_hard_fail_no_downgrade resolved_model=%s error_type=%s",
                llm_model,
                type(exc).__name__,
            )
            raise HTTPException(
                status_code=503,
                detail={
                    "code": "premium_parse_unavailable",
                    "message": "Premium agreement rewrite requires the premium model path; retry shortly.",
                    "resolved_model": llm_model,
                    "error_type": type(exc).__name__,
                },
            ) from exc
        fb_meta: Optional[Dict[str, Any]] = None
        if parse_debug_client:
            fb_meta = {
                "ai_model_class": "basic",
                "resolved_model": llm_model,
                "fallback": "heuristic_basic",
                "error_type": type(exc).__name__,
            }
        return AgreementParseResponse(draft=_heuristic_parse_intake(body.intake_text), parse_meta=fb_meta)


@router.post("/premium-missing-facts", response_model=PremiumMissingFactsResponse)
def premium_missing_facts(request: Request, body: PremiumMissingFactsRequest) -> PremiumMissingFactsResponse:
    """
    Pre–full-draft: up to 5 high-value open questions. Fail-open returns empty questions on model errors.
    """
    require_claw_org_id_header(request)
    ok_txt, msg_txt = validate_negotiate_text(body.intake_text, "owner")
    if not ok_txt:
        raise HTTPException(status_code=400, detail=msg_txt)
    request_ip = request.client.host if request.client else "unknown"
    user_payload: Dict[str, Any] = {"intake": (body.intake_text or "").strip()}
    if body.context is not None:
        user_payload["context"] = body.context.model_dump(exclude_none=True, mode="json")
    if len(json.dumps(user_payload, ensure_ascii=False)) > 200_000:
        raise HTTPException(status_code=400, detail="Input too large for premium missing-facts check")
    max_out = max(200, int(os.environ.get("CLAW_PREMIUM_MISSING_FACTS_MAX_TOKENS", "900")))
    llm_model = resolve_llm_model_for_access_class("premium")
    try:
        llm_text = call_legal_llm(
            messages=[
                {"role": "system", "content": _premium_missing_facts_system_prompt()},
                {"role": "user", "content": json.dumps(user_payload, ensure_ascii=False)},
            ],
            model=llm_model,
            max_tokens=max_out,
            temperature=0.1,
            airlock_profile="agreement_outbound",
            airlock_log_context="premium_missing_facts",
        )
        parsed = _extract_json_object(llm_text)
        out = _normalize_premium_missing_facts_result(parsed)
        record_ai_call(subject_ref=resolve_subject_from_request(request), request_ip=request_ip or "unknown")
        return out
    except Exception as exc:
        log.warning("premium_missing_facts fail_open err=%s", type(exc).__name__, exc_info=True)
        return PremiumMissingFactsResponse(questions=[])


def build_premium_full_draft_user_payload_for_airlock(
    body: PremiumFullDraftRequest,
) -> tuple[Dict[str, Any], Optional[Dict[str, Any]]]:
    """
    Build the ``user_payload`` dict serialized to the premium full-draft model user message
    (the same text evaluated by ``run_ai_airlock(..., policy_profile="agreement_outbound")``).

    Returns ``(user_payload, ctx_dict)``. Kept in sync with :func:`premium_full_draft` assembly.
    """
    intake_s = (body.intake_text or "").strip()
    ctx_dict: Optional[Dict[str, Any]] = (
        body.context.model_dump(exclude_none=True, mode="json") if body.context else None
    )
    intent_key = resolve_premium_intent_key(intake_s, ctx_dict)
    intent_skeleton = build_premium_intent_skeleton(intent_key, intake_s)
    user_payload: Dict[str, Any] = {"intake": intake_s}
    if body.context is not None:
        user_payload["context"] = body.context.model_dump(exclude_none=True, mode="json")
    fam_hint = (body.context.agreement_family or "").strip() if body.context else ""
    scen_cat, scen_sigs = _detect_premium_scenario_category(intake_s, fam_hint)
    user_payload["scenario_category"] = scen_cat
    user_payload["scenario_category_signals"] = scen_sigs[:12]
    user_payload["generation_intelligence_brief"] = build_premium_generation_intelligence_brief(
        intake_s,
        scenario_category=scen_cat,
        scenario_signals=scen_sigs,
    )
    if intent_key is not None:
        user_payload["deterministic_premium_intent_key"] = intent_key.value
    if intent_skeleton:
        user_payload["deterministic_premium_intent_skeleton"] = intent_skeleton
    uga = (body.user_gap_answers or "").strip()
    if uga:
        user_payload["user_gap_answers"] = uga
    sim_regen = bool(getattr(body, "similarity_regeneration", False))
    if sim_regen:
        user_payload["regeneration"] = "similarity_distinct"
        user_payload["regeneration_directive"] = (
            "This is a second pass: the first full draft was too close to a thin free outline. "
            "Rewrite the entire document_text to be clearly more detailed and non-overlapping in structure. "
            "Use real party names from the intake. If payment or fees are described, set them out in a **numbered** "
            "compensation section. Include as appropriate: scope, IP ownership, confidentiality, a limit on revision "
            "rounds, termination, governing law and venue from the user’s **stated** state (never replace Oklahoma with "
            "Delaware or swap states without intake support), notices, counterparts, e-sign, and full signature blocks."
        )
    user_payload = enrich_user_payload_for_simple_consulting(user_payload, intake_s, ctx_dict)
    return user_payload, ctx_dict


def _premium_finalization_clarification_payload(
    answers: List[PremiumFinalizationClarificationAnswer],
) -> List[Dict[str, str]]:
    out: List[Dict[str, str]] = []
    for answer in answers[:50]:
        out.append(
            {
                "question_id": (answer.question_id or "").strip(),
                "question": answer.question.strip(),
                "answer": answer.answer.strip(),
            }
        )
    return out


@router.post("/premium/finalize", response_model=PremiumFinalizationResult)
def premium_finalize(request: Request, body: PremiumFinalizationRequest) -> PremiumFinalizationResult:
    """
    Explicit Phase 4 Pro finalization/repair route.

    This route never runs from first-pass generation automatically. The client must
    call it deliberately after material clarification answers or a repair-needed state.
    """
    require_claw_org_id_header(request)
    ok_txt, msg_txt = validate_negotiate_text(body.original_intake, "owner")
    if not ok_txt:
        raise HTTPException(status_code=400, detail=msg_txt)
    first_draft = (body.first_draft or "").strip()
    if not first_draft:
        raise HTTPException(status_code=400, detail="first_draft is required")
    answers_json = json.dumps([a.model_dump(mode="json") for a in body.clarification_answers], ensure_ascii=False)
    payload_size = len(body.original_intake.encode("utf-8")) + len(first_draft.encode("utf-8")) + len(answers_json)
    if payload_size > 330_000:
        raise HTTPException(status_code=400, detail="Input too large for premium finalization")

    intelligence = body.agreement_intelligence or AgreementIntelligence()
    result = finalize_premium_agreement_if_needed(
        original_intake=body.original_intake,
        first_draft=first_draft,
        agreement_intelligence=intelligence,
        agreement_validation=body.agreement_validation,
        clarification_answers=_premium_finalization_clarification_payload(body.clarification_answers),
        force_finalize=body.force_finalize,
    )
    log.info(
        "[premium-finalize-route] finalized=%s reason=%s model_call_count=%s repair_attempted=%s repair_succeeded=%s",
        result.finalized,
        result.reason,
        result.model_call_count,
        result.repair_attempted,
        result.repair_succeeded,
    )
    return result


@router.post("/premium-full-draft", response_model=PremiumFullDraftResponse)
def premium_full_draft(request: Request, body: PremiumFullDraftRequest) -> Response:
    """
    One-shot premium model: full agreement document (not stitched field transforms).
    Returns JSON with `document_text` as the primary body for LawDog Pro read-only preview.
    """
    require_claw_org_id_header(request)
    ok_txt, msg_txt = validate_negotiate_text(body.intake_text, "owner")
    if not ok_txt:
        raise HTTPException(status_code=400, detail=msg_txt)
    request_ip = request.client.host if request.client else "unknown"
    max_out = max(2000, int(os.environ.get("CLAW_PREMIUM_FULL_DRAFT_MAX_TOKENS", "8000")))
    sim_regen = bool(getattr(body, "similarity_regeneration", False))
    llm_model = resolve_llm_model_for_access_class("premium_regen" if sim_regen else "premium")
    if sim_regen:
        log.info(
            "[CLAW] premium similarity retry model=%s",
            llm_model or "default",
        )
    client_gen = (getattr(body, "agreement_generation_id", None) or "").strip() or "n/a"
    client_fp = (getattr(body, "intake_fingerprint", None) or "").strip() or "n/a"
    server_timing: Optional[PaidProServerTiming] = None
    if paid_pro_perf_trace_requested(request):
        server_timing = PaidProServerTiming(
            trace_id=client_gen,
            session_generation_id=client_gen,
            intake_fingerprint=client_fp,
        )
        server_timing.mark_instant("backend_request_received")
    context_assembly_started = time.perf_counter()
    user_payload, ctx_dict = build_premium_full_draft_user_payload_for_airlock(body)
    if server_timing is not None:
        server_timing.record(
            "backend_context_assembly",
            (time.perf_counter() - context_assembly_started) * 1000,
            payloadJsonLen=len(json.dumps(user_payload, ensure_ascii=False)),
        )
    prompt_assembly_started = time.perf_counter()
    intake_s = (body.intake_text or "").strip()
    intent_key = resolve_premium_intent_key(intake_s, ctx_dict)
    intent_skeleton = build_premium_intent_skeleton(intent_key, intake_s)
    uga = (body.user_gap_answers or "").strip()
    airlock_wire_text = json.dumps(user_payload, ensure_ascii=False)
    scen_cat = cast(str, user_payload.get("scenario_category") or "custom_mixed")
    _raw_sigs = user_payload.get("scenario_category_signals")
    scen_sigs: List[str] = list(_raw_sigs) if isinstance(_raw_sigs, list) else []
    if uga:
        uga_hash = canon_sha256_hex(uga)
        log.info(
            "[gap-trace-backend] stage=premium_full_draft_prompt_assembly "
            "has_user_gap_answers=1 gap_len=%s gap_hash=%s prompt_label=user_payload.user_gap_answers",
            len(uga),
            uga_hash,
        )
    else:
        log.info(
            "[gap-trace-backend] stage=premium_full_draft_prompt_assembly "
            "has_user_gap_answers=0 gap_len=0 gap_hash= prompt_label=user_payload.user_gap_answers"
        )
    if len(json.dumps(user_payload, ensure_ascii=False)) > 240_000:
        raise HTTPException(status_code=400, detail="Input too large for premium full draft")
    session_hint = sha256_hex(intake_s.encode("utf-8"))[:16]
    client_agreement_id = (getattr(body, "agreement_id", None) or "").strip() or "n/a"
    ctx_title = (str((ctx_dict or {}).get("title") or "")[:120]) if ctx_dict else ""
    if server_timing is not None:
        server_timing.record(
            "backend_prompt_assembly",
            (time.perf_counter() - prompt_assembly_started) * 1000,
            intentKey=(intent_key.value if intent_key is not None else None),
            airlockWireLen=len(airlock_wire_text),
        )
    log.info(
        "[premium-full-draft] event=start agreement_id=%s session_hint=%s client_generation_id=%s intake_fingerprint=%s intake_len=%s context_title=%r sim_regen=%s payload_json_len=%s",
        client_agreement_id,
        session_hint,
        client_gen,
        client_fp,
        len(intake_s),
        ctx_title,
        int(sim_regen),
        len(json.dumps(user_payload, ensure_ascii=False)),
    )
    log.info(
        "claw_premium route=premium_full_draft model=%s sim_regen=%s max_out=%s %s intake_len=%s payload_json_len=%s",
        (llm_model or ""),
        int(sim_regen),
        max_out,
        _openai_key_diagnostics(),
        len(intake_s),
        len(json.dumps(user_payload, ensure_ascii=False)),
    )
    if not (OPENAI_API_KEY or "").strip():
        log.error("premium_full_draft event=config_error category=missing_env OPENAI_API_KEY_unset=1")
        dm = _premium_full_draft_degraded_response(
            intake_s=intake_s,
            ctx_dict=ctx_dict,
            failure_code="missing_openai_key",
            failure_message=_degraded_user_message_for_code("missing_openai_key"),
        )
        log.info(
            "[premium-full-draft] event=failure stage=config status=200 code=missing_openai_key session_hint=%s intake_len=%s",
            session_hint,
            len(intake_s),
        )
        return _premium_full_draft_finalize_http_response(
            dm,
            intake_len=len(intake_s),
            session_hint=session_hint,
            server_timing=server_timing,
            request=request,
        )
    try:
        if server_timing is not None:
            server_timing.mark_instant(
                "backend_llm_api_call_start",
                model=str(llm_model or ""),
                maxTokens=max_out,
            )
        llm_primary_started = time.perf_counter()
        llm_text = call_legal_llm(
            messages=[
                {"role": "system", "content": _premium_full_draft_system_prompt()},
                {"role": "user", "content": json.dumps(user_payload, ensure_ascii=False)},
            ],
            model=llm_model,
            max_tokens=max_out,
            temperature=0.2 if sim_regen else 0.15,
            airlock_profile="agreement_outbound",
            airlock_log_context="premium_full_draft:primary",
        )
        if server_timing is not None:
            server_timing.record(
                "backend_llm_primary",
                (time.perf_counter() - llm_primary_started) * 1000,
                responseChars=len((llm_text or "").strip()),
            )
        log.info(
            "claw_premium route=premium_full_draft openai_response_chars=%s model=%s",
            len((llm_text or "").strip()),
            (llm_model or ""),
        )
        intelligence_parse_start = time.perf_counter()
        parsed = _extract_json_object(llm_text)
        out_primary = _normalize_premium_full_draft_result(parsed)
        if server_timing is not None:
            server_timing.record(
                "backend_parse_normalize",
                (time.perf_counter() - intelligence_parse_start) * 1000,
                primaryDocLen=len((out_primary.document_text or "").strip()),
            )
        _log_agreement_intelligence_summary(
            out_primary.agreement_intelligence,
            stage="primary",
            elapsed_ms=(time.perf_counter() - intelligence_parse_start) * 1000,
        )
        log.info(
            "claw_premium route=premium_full_draft parse_ok=1 primary_doc_len=%s title_len=%s",
            len((out_primary.document_text or "").strip()),
            len((out_primary.title or "").strip()),
        )
        log.info(
            "[premium-full-draft] event=llm_success status=200 session_hint=%s document_text_len=%s title_len=%s",
            session_hint,
            len((out_primary.document_text or "").strip()),
            len((out_primary.title or "").strip()),
        )
        _safe_record_ai_call(request, request_ip)

        post_processing_started = time.perf_counter()
        free_blob = build_free_reference_blob(intake_s, ctx_dict)
        repair_used = False
        repair_body = ""
        out = out_primary

        _brief_for_grade = user_payload.get("generation_intelligence_brief") or {}
        _contradiction_notes = _brief_for_grade.get("contradiction_notes")
        if not isinstance(_contradiction_notes, list):
            _contradiction_notes = None

        def _grade_draft(graded: PremiumFullDraftResponse) -> Tuple[bool, List[str]]:
            ok_q, reasons_q = evaluate_premium_full_draft_quality(
                intake=intake_s,
                context=ctx_dict,
                draft_title=graded.title,
                draft_family=graded.agreement_family,
                draft_document_text=graded.document_text,
                scenario_category=scen_cat,
                contradiction_notes=_contradiction_notes,
            )
            ok_s, reasons_s = evaluate_premium_intent_schema(
                intent_key, graded.title, graded.document_text
            )
            merged = list(dict.fromkeys([*reasons_q, *reasons_s]))
            return (ok_q and ok_s), merged

        quality_grade_started = time.perf_counter()
        ok_all, reject_reasons = _grade_draft(out_primary)
        if server_timing is not None:
            server_timing.record(
                "backend_quality_grade",
                (time.perf_counter() - quality_grade_started) * 1000,
                path="primary",
                qualityOk=bool(ok_all),
                reasonCount=len(reject_reasons),
            )
        if not ok_all:
            log.info(
                "premium_full_draft_quality_event event=premium_full_draft_quality_or_schema_fail reasons=%s",
                ",".join(reject_reasons[:16]),
            )
            repair_payload = build_premium_full_draft_repair_user_payload(
                intake=intake_s,
                free_reference_blob=free_blob,
                rejected=out_primary.model_dump(),
                rejection_reasons=reject_reasons,
                scenario_category=scen_cat,
                scenario_signals=scen_sigs,
                context=ctx_dict,
                deterministic_premium_intent_skeleton=intent_skeleton,
                premium_intent_key=(intent_key.value if intent_key is not None else None),
            )
            if len(json.dumps(repair_payload, ensure_ascii=False)) > 260_000:
                raise ValueError("repair_payload_too_large")
            airlock_wire_text = json.dumps(repair_payload, ensure_ascii=False)
            llm_repair_started = time.perf_counter()
            llm_repair = call_legal_llm(
                messages=[
                    {"role": "system", "content": premium_full_draft_repair_system_prompt()},
                    {"role": "user", "content": json.dumps(repair_payload, ensure_ascii=False)},
                ],
                model=llm_model,
                max_tokens=max_out,
                temperature=0.22,
                airlock_profile="agreement_outbound",
                airlock_log_context="premium_full_draft:repair",
            )
            repair_llm_ms = (time.perf_counter() - llm_repair_started) * 1000
            if server_timing is not None:
                server_timing.record(
                    "backend_llm_repair",
                    repair_llm_ms,
                    responseChars=len((llm_repair or "").strip()),
                )
                server_timing.record(
                    "backend_llm_repair_or_regen",
                    repair_llm_ms,
                    path="repair",
                )
            intelligence_parse_start = time.perf_counter()
            parsed_repair = _extract_json_object(llm_repair)
            out = _normalize_premium_full_draft_result(parsed_repair)
            if server_timing is not None:
                server_timing.record(
                    "backend_parse_normalize",
                    (time.perf_counter() - intelligence_parse_start) * 1000,
                    primaryDocLen=len((out.document_text or "").strip()),
                    path="repair",
                )
            _log_agreement_intelligence_summary(
                out.agreement_intelligence,
                stage="repair",
                elapsed_ms=(time.perf_counter() - intelligence_parse_start) * 1000,
            )
            _safe_record_ai_call(request, request_ip)
            repair_used = True
            repair_body = (out.document_text or "").strip()
            log.info("premium_full_draft_quality_event event=premium_full_draft_repair_used")
        else:
            log.info("premium_full_draft_quality_event event=premium_full_draft_quality_pass")

        doc = (out.document_text or "").strip()
        has_leak, leak_hits = premium_document_text_has_dev_context_leak(doc)
        if has_leak:
            log.error(
                "dev_context_leak event=premium_full_draft stage=pre_return labels=%s",
                ",".join(leak_hits),
            )
            intake2 = sanitize_premium_intake_for_retry(intake_s)
            up_clean: Dict[str, Any] = {
                "intake": intake2,
                "scenario_category": scen_cat,
                "scenario_category_signals": scen_sigs[:12],
            }
            if body.context is not None and ctx_dict is not None:
                up_clean["context"] = serialize_context_clean(ctx_dict) or {}
            if uga:
                up_clean["user_gap_answers"] = sanitize_premium_intake_for_retry(uga)
            if len(json.dumps(up_clean, ensure_ascii=False)) > 240_000:
                raise ValueError("clean_premium_user_payload_too_large")
            airlock_wire_text = json.dumps(up_clean, ensure_ascii=False)
            llm_sanitized_started = time.perf_counter()
            llm_clean = call_legal_llm(
                messages=[
                    {"role": "system", "content": _premium_full_draft_system_prompt()},
                    {"role": "user", "content": json.dumps(up_clean, ensure_ascii=False)},
                ],
                model=llm_model,
                max_tokens=max_out,
                temperature=0.12,
                airlock_profile="agreement_outbound",
                airlock_log_context="premium_full_draft:sanitized_retry",
            )
            sanitized_llm_ms = (time.perf_counter() - llm_sanitized_started) * 1000
            if server_timing is not None:
                server_timing.record(
                    "backend_llm_sanitized_retry",
                    sanitized_llm_ms,
                    responseChars=len((llm_clean or "").strip()),
                )
                server_timing.record(
                    "backend_llm_repair_or_regen",
                    sanitized_llm_ms,
                    path="sanitized_retry",
                )
            intelligence_parse_start = time.perf_counter()
            parsed_c = _extract_json_object(llm_clean)
            out_clean = _normalize_premium_full_draft_result(parsed_c)
            if server_timing is not None:
                server_timing.record(
                    "backend_parse_normalize",
                    (time.perf_counter() - intelligence_parse_start) * 1000,
                    primaryDocLen=len((out_clean.document_text or "").strip()),
                    path="sanitized_retry",
                )
            _log_agreement_intelligence_summary(
                out_clean.agreement_intelligence,
                stage="sanitized_retry",
                elapsed_ms=(time.perf_counter() - intelligence_parse_start) * 1000,
            )
            _safe_record_ai_call(request, request_ip)
            doc_c = (out_clean.document_text or "").strip()
            has_leak2, leak_h2 = premium_document_text_has_dev_context_leak(doc_c)
            if has_leak2 or not doc_c:
                log.error(
                    "dev_context_leak event=premium_full_draft_sanitized_retry_failed labels=%s",
                    ",".join(leak_h2),
                )
                raise ValueError("premium_dev_context_leak_in_output")
            out = out_clean
            out_primary = out_clean
            doc = doc_c
            repair_used = False
            repair_body = ""
            log.info("dev_context_leak event=premium_full_draft_sanitized_retry_succeeded")
        final_grade_started = time.perf_counter()
        ok_final, final_reasons = _grade_draft(out)
        if server_timing is not None:
            server_timing.record(
                "backend_quality_grade",
                (time.perf_counter() - final_grade_started) * 1000,
                path="final",
                qualityOk=bool(ok_final),
                reasonCount=len(final_reasons),
            )
            server_timing.record(
                "backend_post_processing",
                (time.perf_counter() - post_processing_started) * 1000,
                repairUsed=bool(repair_used),
            )
        validation_started = time.perf_counter()
        generation_outcome: Literal["ok", "needs_details", "degraded"] = "ok" if ok_final else "needs_details"
        if not ok_final:
            log.info(
                "premium_full_draft event=validator_reject category=quality_or_intent_schema needs_details=1 "
                "reasons=%s",
                ",".join(final_reasons[:24]),
            )
            log.info(
                "premium_full_draft_quality_event event=premium_intent_or_quality_needs_details reasons=%s",
                ",".join(final_reasons[:16]),
            )
        if uga:
            needles = (
                "mediation",
                "binding arbitration",
                "fair market value",
                "trailing 6 months earnings",
                "independent appraisal",
            )
            hit = [n for n in needles if n in doc.lower()]
            log.info(
                "[gap-trace-backend] stage=premium_full_draft_response "
                "doc_len=%s contains_needles=%s needles_hit=%s",
                len(doc),
                1 if hit else 0,
                ",".join(hit),
            )
        else:
            log.info(
                "[gap-trace-backend] stage=premium_full_draft_response "
                "doc_len=%s contains_needles=0 needles_hit=",
                len(doc),
            )
        if not (out.document_text or "").strip():
            raise ValueError("empty_document_text")
        log.info(
            "premium_full_draft_quality_event event=premium_full_draft_render_source source=%s doc_len=%s",
            "server_repaired_accepted" if repair_used else "server_primary_accepted",
            len(doc),
        )
        log.info(
            "[CLAW] premium accepted outcome=%s doc_len=%s repair_used=%s sim_regen=%s",
            generation_outcome,
            len(doc),
            repair_used,
            sim_regen,
        )
        try:
            _pfd_sum = {
                "phase": "success",
                "model": str(llm_model or ""),
                "intake_len": len(intake_s),
                "prompt_len": len(json.dumps(user_payload, ensure_ascii=False)),
                "llm_response_len": len((llm_text or "").strip()),
                "parsed_doc_len": len((out_primary.document_text or "").strip()),
                "final_doc_len": len(doc),
                "generation_outcome": str(generation_outcome),
                "failure_code": None,
                "quality_gate_ok": bool(ok_final),
                "quality_gate_reasons": [str(x) for x in list(final_reasons)[:24]] if not ok_final else [],
                "repair_used": bool(repair_used),
            }
            log.info("claw_premium route=premium_full_draft json_summary=%s", json.dumps(_pfd_sum, default=str)[:12000])
        except Exception as jsum_e:
            log.info(
                "claw_premium route=premium_full_draft json_summary=unavailable jsum_err=%s",
                type(jsum_e).__name__,
            )
        primary_full = (out_primary.document_text or "").strip()
        agreement_validation = _validate_and_log_premium_agreement_draft(
            authoritative_draft=doc,
            agreement_intelligence=out.agreement_intelligence,
            original_intake=intake_s,
            stage="final",
        )
        if server_timing is not None:
            server_timing.record(
                "backend_validation",
                (time.perf_counter() - validation_started) * 1000,
                qualityOk=bool(ok_final),
            )
        ok_model = PremiumFullDraftResponse(
            title=out.title,
            agreement_family=out.agreement_family,
            document_text=doc,
            server_full_document_text=primary_full,
            server_repair_document_text=repair_body,
            authoritative_draft=doc,
            agreement_intelligence=out.agreement_intelligence,
            agreement_validation=agreement_validation,
            key_terms_found=out.key_terms_found,
            missing_material_info=out.missing_material_info,
            generation_outcome=generation_outcome,
            schema_validation_reasons=final_reasons,
            generation_ok=bool(doc.strip()),
            retryable=False,
        )
        return _premium_full_draft_finalize_http_response(
            ok_model,
            intake_len=len(intake_s),
            session_hint=session_hint,
            server_timing=server_timing,
            request=request,
        )
    except Exception as exc:
        code, log_detail = _classify_premium_full_draft_failure(exc)
        if code == "airlock_blocked":
            try:
                from backend.agreements.premium_airlock import (
                    assess_premium_agreement_outbound_airlock,
                    log_premium_airlock_decision_for_route,
                )

                prem = assess_premium_agreement_outbound_airlock(
                    airlock_wire_text, policy_profile="agreement_outbound"
                )
                if prem is not None:
                    log_premium_airlock_decision_for_route(
                        prem, airlock_route="premium_full_draft:failure"
                    )
                diag = first_privilege_airlock_block_diagnostic(
                    airlock_wire_text, policy_profile="agreement_outbound"
                )
                suffix = ""
                if diag is not None:
                    suffix = (
                        f" first_block_reason={diag.reason_code} first_block_category={diag.rule_category}"
                        f" first_block_rule_id={diag.matched_rule_id}"
                    )
                log.warning(
                    "[premium-full-draft] event=airlock_blocked airlock_profile=agreement_outbound "
                    "airlock_route=premium_full_draft:user_wire user_message_index=0 user_content_chars=%s%s",
                    len(airlock_wire_text),
                    suffix,
                )
            except Exception as adiag:
                log.warning(
                    "[premium-full-draft] event=airlock_blocked_diag_failed exc=%s",
                    type(adiag).__name__,
                )
        try:
            _pfd_degraded = {
                "phase": "degraded",
                "model": str(llm_model or ""),
                "intake_len": len(intake_s),
                "prompt_len": len(json.dumps(user_payload, ensure_ascii=False)),
                "failure_code": str(code),
                "error_detail": str(log_detail)[:800],
            }
            log.info(
                "claw_premium route=premium_full_draft json_summary=%s",
                json.dumps(_pfd_degraded, default=str)[:10000],
            )
        except Exception as jex:
            log.info("claw_premium route=premium_full_draft json_summary=unavailable ex=%s", type(jex).__name__)
        log.warning(
            "premium_full_draft event=model_path_failed category=%s detail=%s err_type=%s",
            code,
            log_detail,
            type(exc).__name__,
            exc_info=True,
        )
        log.warning(
            "[premium-full-draft] event=failure stage=model_path status=200 code=%s exc_type=%s session_hint=%s intake_len=%s detail=%s",
            code,
            type(exc).__name__,
            session_hint,
            len(intake_s),
            str(log_detail)[:500],
        )
        dm = _premium_full_draft_degraded_response(
            intake_s=intake_s,
            ctx_dict=ctx_dict,
            failure_code=code,
            failure_message=_degraded_user_message_for_code(code),
        )
        return _premium_full_draft_finalize_http_response(
            dm,
            intake_len=len(intake_s),
            session_hint=session_hint,
            server_timing=server_timing,
            request=request,
        )


@router.post("/premium-review", response_model=PremiumAgreementReviewResponse)
def premium_agreement_review(request: Request, body: PremiumAgreementReviewRequest) -> PremiumAgreementReviewResponse:
    """
    After premium full draft (or final premium body): light AI pass for clarity / completeness nudges.
    Does not replace the agreement; returns structured bullets for the UI.
    """
    require_claw_org_id_header(request)
    ok_txt, msg_txt = validate_negotiate_text(body.intake_text, "owner")
    if not ok_txt:
        raise HTTPException(status_code=400, detail=msg_txt)
    request_ip = request.client.host if request.client else "unknown"
    doc = (body.document_text or "").strip()
    if not doc:
        raise HTTPException(status_code=400, detail="document_text is required")
    trunc_note = ""
    if len(doc) > 200_000:
        doc = doc[:200_000]
        trunc_note = " [Draft truncated to first ~200k characters for review.]"
    max_out = max(500, int(os.environ.get("CLAW_PREMIUM_REVIEW_MAX_TOKENS", "2000")))
    llm_model = resolve_llm_model_for_access_class("premium")
    user_payload: Dict[str, Any] = {
        "intake": (body.intake_text or "").strip(),
        "generated_premium_draft": doc + trunc_note,
    }
    if body.context is not None:
        user_payload["extracted_context"] = body.context.model_dump(exclude_none=True, mode="json")
    plen = len(json.dumps(user_payload, ensure_ascii=False))
    if plen > 260_000:
        raise HTTPException(status_code=400, detail="Input too large for premium review")
    log.info(
        "claw_premium route=premium_review model=%s max_out=%d %s document_len=%d intake_len=%d payload_json_len=%d",
        (llm_model or ""),
        max_out,
        _openai_key_diagnostics(),
        len(doc),
        len((body.intake_text or "").strip()),
        plen,
    )
    try:
        llm_text = call_legal_llm(
            messages=[
                {"role": "system", "content": _premium_agreement_review_system_prompt()},
                {"role": "user", "content": json.dumps(user_payload, ensure_ascii=False)},
            ],
            model=llm_model,
            max_tokens=max_out,
            temperature=0.2,
            airlock_profile="agreement_outbound",
        )
        log.info("claw_premium route=premium_review openai_response_chars=%d", len((llm_text or "").strip()))
        parsed = _extract_json_object(llm_text)
        out = _normalize_premium_agreement_review_result(parsed)
        _safe_record_ai_call(request, request_ip)
        return out
    except Exception as exc:
        kind = _classify_premium_llm_failure(exc)
        log.warning(
            "claw_premium route=premium_review FAILED class=%s exc_type=%s model=%s %s err_snip=%s (fail_open_200)",
            kind,
            type(exc).__name__,
            (llm_model or ""),
            _openai_key_diagnostics(),
            (str(exc) or "")[:500].replace("\n", " "),
            exc_info=True,
        )
        return _fallback_premium_agreement_review_response()


@router.get("/usage/summary")
def get_agreement_usage_summary(request: Request) -> Dict[str, Any]:
    """User-facing usage (no internal Key units)."""
    require_claw_org_id_header(request)
    return usage_summary_for_subject(resolve_subject_from_request(request))


def _ensure_agreement_parties_have_ids(parties: List[AgreementParty]) -> List[AgreementParty]:
    sanitized = _sanitize_agreement_parties_in_order(list(parties or []))
    out: List[AgreementParty] = []
    for p in sanitized:
        pid = (p.id or "").strip() or str(uuid.uuid4())
        email = str(getattr(p, "email", None) or "").strip() or None
        phone = str(getattr(p, "phone", None) or "").strip() or None
        out.append(AgreementParty(name=p.name, role=p.role, id=pid, email=email, phone=phone))
    return out


@router.post("/draft")
def create_agreement_draft(body: AgreementDraftCreate, request: Request) -> Dict[str, Any]:
    require_claw_org_id_header(request)
    subject = resolve_subject_from_request(request)
    request_ip = request.client.host if request.client else "unknown"
    if not review_first_paid_pro_persist_bypass(request=request, purpose=body.purpose or ""):
        assert_can_create_draft(subject_ref=subject, request_ip=request_ip or "unknown")
    now = _utc_now_iso()
    agreement_id = str(uuid.uuid4())
    parties = _ensure_agreement_parties_have_ids(list(body.parties or []))
    draft = AgreementDraft(
        id=agreement_id,
        title=body.title,
        jurisdiction=body.jurisdiction,
        parties=parties,
        purpose=body.purpose,
        payment_terms=body.payment_terms,
        duration=body.duration,
        due_date=body.due_date,
        effective_date=body.effective_date,
        feed_visibility=body.feed_visibility,
        feed_party_anonymize=body.feed_party_anonymize,
        feed_show_financial_summary=body.feed_show_financial_summary,
        feed_anchor_network=body.feed_anchor_network,
        created_at=now,
        updated_at=now,
        versions=[],
        audit_log=[AuditEvent(event_type="created", at=now)],
    )
    dump = draft.model_dump()
    _save_draft_sync(dump, request)
    record_public_feed_event_if_applicable(draft_dict=dump, event_type="created", at=now)
    record_usage_ledger_event(
        subject_ref=subject,
        event_type="agreement_created",
        agreement_id=agreement_id,
        metadata={},
    )
    record_draft_created(
        agreement_id=agreement_id,
        subject_ref=subject,
        request_ip=request_ip or "unknown",
    )
    try:
        from backend.integrations.hooks_emit import claw_emit_integration_event_from_subject

        claw_emit_integration_event_from_subject(
            subject,
            "agreement.created",
            "agreement",
            agreement_id,
            {"title": body.title, "surface": "draft"},
        )
    except Exception:
        pass
    return {"id": agreement_id, "draft": dump, "canonical_json": canonicalize_agreement(dump)}


class AgreementDraftForkRequest(BaseModel):
    """Create a new draft seeded from an existing agreement (new id; proof does not carry over)."""

    source_agreement_id: str = Field(..., min_length=1)


@router.post("/draft-from-agreement")
def create_draft_from_prior_agreement(body: AgreementDraftForkRequest, request: Request) -> Dict[str, Any]:
    require_claw_org_id_header(request)
    subject = resolve_subject_from_request(request)
    request_ip = request.client.host if request.client else "unknown"
    assert_can_create_draft(subject_ref=subject, request_ip=request_ip or "unknown")
    src_id = body.source_agreement_id.strip()
    assert_registered_owner_matches(request, src_id)
    try:
        raw = load_draft(src_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="agreement_not_found")

    now = _utc_now_iso()
    new_id = str(uuid.uuid4())
    parties_acc: List[AgreementParty] = []
    for p in raw.get("parties") or []:
        if not isinstance(p, dict):
            continue
        name = str(p.get("name") or "").strip()
        role = str(p.get("role") or "party").strip() or "party"
        if name:
            parties_acc.append(AgreementParty(name=name, role=role, id=str(uuid.uuid4())))
    parties_out = _ensure_agreement_parties_have_ids(parties_acc)

    base_title = str(raw.get("title") or "").strip() or "Agreement"
    title = f"{base_title} (new from prior)"
    purpose = str(raw.get("purpose") or "").strip()
    fork_banner = (
        "Started from a prior workspace agreement as a new draft. "
        "Review all terms — prior signatures and proof do not apply to this copy."
    )
    purpose = f"{purpose}\n\n— {fork_banner}" if purpose else fork_banner

    fv = raw.get("feed_visibility") or "private"
    if fv not in ("private", "link_only", "public"):
        fv = "private"

    draft = AgreementDraft(
        id=new_id,
        title=title,
        jurisdiction=str(raw.get("jurisdiction") or ""),
        parties=parties_out,
        purpose=purpose,
        payment_terms=str(raw.get("payment_terms") or ""),
        duration=raw.get("duration"),
        due_date=raw.get("due_date"),
        effective_date=raw.get("effective_date"),
        feed_visibility=fv,
        feed_party_anonymize=bool(raw.get("feed_party_anonymize", False)),
        feed_show_financial_summary=bool(raw.get("feed_show_financial_summary", False)),
        feed_anchor_network=raw.get("feed_anchor_network"),
        payment_request=raw.get("payment_request") if isinstance(raw.get("payment_request"), dict) else None,
        payment_required=bool(raw.get("payment_required", False)),
        created_at=now,
        updated_at=now,
        versions=[],
        audit_log=[
            AuditEvent(event_type="created", at=now),
            AuditEvent(
                event_type="forked_from_agreement",
                at=now,
                field="source_agreement_id",
                value={
                    "source_agreement_id": src_id,
                    "note": "New draft; canonical proof records remain on the source agreement.",
                },
            ),
        ],
        review_sent_at=None,
        workspace_archived_at=None,
    )
    dump = draft.model_dump()
    _save_draft_sync(dump, request)
    record_public_feed_event_if_applicable(draft_dict=dump, event_type="created", at=now)
    record_usage_ledger_event(
        subject_ref=subject,
        event_type="agreement_created",
        agreement_id=new_id,
        metadata={"forked_from": src_id},
    )
    record_draft_created(
        agreement_id=new_id,
        subject_ref=subject,
        request_ip=request_ip or "unknown",
    )
    try:
        from backend.integrations.hooks_emit import claw_emit_integration_event_from_subject

        claw_emit_integration_event_from_subject(
            subject,
            "agreement.created",
            "agreement",
            new_id,
            {"title": title, "surface": "draft_from_agreement", "source_agreement_id": src_id},
        )
    except Exception:
        pass
    return {"id": new_id, "draft": dump, "canonical_json": canonicalize_agreement(dump)}


class AgreementSigningLockBody(BaseModel):
    locked_version_id: str
    locked_at: str
    locked_by: str = "owner"


class WorkspaceArchiveBody(BaseModel):
    archived: bool = True


class WorkspaceFolderAssignBody(BaseModel):
    """Set or clear which proof-layer folder this agreement is filed under (flat folders)."""

    folder_id: Optional[str] = None


class WorkspaceTagsReplaceBody(BaseModel):
    """Replace manual workspace tags (searchable labels; not required)."""

    tags: List[str] = Field(default_factory=list)


def _normalize_workspace_tags(raw: List[Any]) -> List[str]:
    seen: set[str] = set()
    out: List[str] = []
    for t in raw or []:
        s = str(t).strip()
        if not s:
            continue
        if len(s) > 40:
            s = s[:40]
        k = s.casefold()
        if k in seen:
            continue
        seen.add(k)
        out.append(s)
        if len(out) >= 12:
            break
    return out


def _folder_name_map_for_subject(subject: str) -> Dict[str, str]:
    st = ProofLayerStore()
    st.init_schema()
    return {str(r["folder_id"]): str(r["folder_name"]) for r in st.list_folders(subject)}


class RecipientAccessMintRequest(BaseModel):
    mode: Literal["sign", "review"] = "sign"
    role: Literal["recipient", "reviewer", "signer"] = "recipient"
    recipient_subject: Optional[str] = None
    recipient_party_id: Optional[str] = None
    inviter_display_name: Optional[str] = None
    single_use: bool = False
    ttl_seconds: int = 60 * 60 * 24 * 7
    # Paid Pro review-first only: exact frozen corpus used by the SPA to mint/share review links.
    review_first_document_text: Optional[str] = None
    review_first_document_source: Optional[str] = None


def _review_first_corpus_hash(text: str) -> str:
    body = (text or "").strip()
    return f"{len(body)}:{hashlib.sha256(body.encode('utf-8')).hexdigest()[:8]}"


def _review_first_final_corpus_from_draft(
    draft: AgreementDraft, *, include_field_fallbacks: bool = False
) -> Tuple[str, str]:
    pr = draft.pro_redline_v1 if isinstance(draft.pro_redline_v1, dict) else {}
    rf = pr.get("review_first_final_corpus") if isinstance(pr, dict) else None
    if isinstance(rf, dict):
        txt = str(rf.get("text") or "").strip()
        if txt:
            return txt, "review_first_final_corpus"
    if not include_field_fallbacks:
        return "", "none"
    for key in (
        "server_full_document_text",
        "premium_server_full_document_text",
        "premium_full_document_text",
        "document_text",
        "rendered_document_text",
    ):
        txt = str(getattr(draft, key, "") or "").strip()
        if txt:
            return txt, key
    return "", "none"


def _persist_review_first_final_corpus_if_supplied(
    agreement_id: str,
    body: RecipientAccessMintRequest,
    *,
    locked_version_id: str,
    subject_ref: Optional[str] = None,
) -> None:
    if body.mode != "review":
        return
    corpus = (body.review_first_document_text or "").strip()
    if len(corpus) < 80:
        return
    draft = _load_or_404(agreement_id)
    now = _utc_now_iso()
    corpus_hash = _review_first_corpus_hash(corpus)
    existing, _source = _review_first_final_corpus_from_draft(draft)
    existing_hash = _review_first_corpus_hash(existing) if existing else ""
    if existing_hash == corpus_hash:
        return

    pro_redline = dict(draft.pro_redline_v1 or {})
    pro_redline["review_first_final_corpus"] = {
        "text": corpus,
        "source": (body.review_first_document_source or "recipient_access_token").strip(),
        "hash": corpus_hash,
        "persisted_at": now,
        "locked_version_id": locked_version_id or None,
    }
    audit = [*(draft.audit_log or [])]
    audit.append(
        AuditEvent(
            event_type="review_first_final_corpus_persisted",
            at=now,
            field="review_first_final_corpus",
            value={
                "hash": corpus_hash,
                "len": len(corpus),
                "source": (body.review_first_document_source or "recipient_access_token").strip(),
                "locked_version_id": locked_version_id or None,
            },
        )
    )
    versions = [*(draft.versions or [])]
    versions.append(
        VersionSnapshot(
            version=len(versions) + 1,
            created_at=now,
            note="Review-first final corpus",
        )
    )
    next_draft = _merge_agreement_draft(
        draft,
        updated_at=now,
        review_sent_at=draft.review_sent_at or now,
        purpose=corpus,
        document_text=corpus,
        server_full_document_text=corpus,
        premium_server_full_document_text=corpus,
        premium_full_document_text=corpus,
        rendered_document_text=corpus,
        premium_render_source="review_first_final_corpus",
        pro_redline_v1=pro_redline,
        audit_log=audit,
        versions=versions,
    )
    _save_draft_sync(next_draft.model_dump(), subject_ref=subject_ref)
    log.info(
        "[review-first-backend-corpus-persisted] agreement_id_short=%s len=%s hash=%s source=%s",
        agreement_id[:8],
        len(corpus),
        corpus_hash,
        (body.review_first_document_source or "recipient_access_token").strip(),
    )


@router.get("/access/policy")
def recipient_access_policy() -> Dict[str, Any]:
    """Public: lets the SPA decide whether ``t=`` links are mandatory."""
    return {
        "recipient_link_token_required": recipient_access_token_required(),
        "mint_key_configured": bool(os.getenv("CLAW_RECIPIENT_LINK_MINT_KEY", "").strip()),
        "signing_token_configured": operator_signing_token_secret_configured(),
        "review_link_mint_enabled": review_link_mint_enabled(),
        "signing_token_env_var_detected": detected_signing_token_env_var(),
        "recipient_token_ttl_seconds": {
            "min": recipient_token_ttl_min_seconds(),
            "max": recipient_token_ttl_max_seconds(),
        },
    }


def _recipient_party_id_on_draft(draft: Dict[str, Any], party_id: str) -> bool:
    pid = (party_id or "").strip()
    if not pid:
        return False
    for p in draft.get("parties") or []:
        if not isinstance(p, dict):
            continue
        if str(p.get("id") or "").strip() == pid:
            return True
    return False


@router.get("/access/validate")
def recipient_access_validate(token: str = "", agreement_id: str = "") -> Dict[str, Any]:
    try:
        secret_raw = resolve_signing_token_secret_raw()
    except SigningTokenSecretMissingInProductionError as e:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "signing_token_secret_not_configured",
                "message": str(e),
            },
        ) from e
    log_ok = os.getenv("CLAW_RECIPIENT_TOKEN_LOG_VALIDATIONS", "1").strip().lower() not in (
        "0",
        "false",
        "no",
    )
    # Do not consume single-use JTI here — recipients need the same token for GET/render until signing flows
    # attach stricter consumption (see validate_recipient_access_token_for_agreement).
    out = validate_recipient_access_token_for_agreement(
        token=token,
        path_agreement_id="",
        query_agreement_id=(agreement_id or "").strip() or None,
        secret_raw=secret_raw,
        consume_single_use=False,
        log_validation=log_ok,
    )
    try:
        from backend.services.agreement_draft_store import load_draft, save_draft
        from backend.services.recipient_delivery_registry import extract_jti_from_token, record_invite_opened

        aid = str(out.get("agreement_id") or "").strip()
        pid = str(out.get("recipient_party_id") or "").strip()
        mode = str(out.get("mode") or "").strip()
        if aid and pid and mode in ("review", "sign"):
            phase = "review" if mode == "review" else "signing"
            raw = load_draft(aid)
            audit = list(raw.get("audit_log") or [])
            record_invite_opened(
                raw,
                phase=phase,
                participant_id=pid,
                jti=extract_jti_from_token(token),
                audit_log=audit,
            )
            raw["audit_log"] = audit
            save_draft({**raw, "id": aid})
    except Exception:
        pass
    return {
        "ok": True,
        "agreement_id": out["agreement_id"],
        "mode": out["mode"],
        "locked_version_id": out["locked_version_id"],
        "role": out["role"],
        "recipient_party_id": out["recipient_party_id"],
        "inviter_display_name": out["inviter_display_name"],
    }


@router.get("/workspace-index")
def get_agreements_workspace_index(request: Request) -> Dict[str, Any]:
    """Lightweight list for Agreement Workspace landing (local / single-tenant style)."""
    require_claw_org_id_header(request)
    subject = resolve_subject_from_request(request)
    folder_names = _folder_name_map_for_subject(subject)
    summaries: List[Dict[str, Any]] = []
    skipped: List[Dict[str, str]] = []
    local_ids = list_draft_agreement_ids_newest_first()
    agreement_ids = merge_workspace_index_agreement_ids(
        subject_ref=subject,
        local_ids_newest_first=local_ids,
    )
    supabase_rows = supabase_rows_by_id_for_subject(subject)
    for aid in agreement_ids:
        if aid not in supabase_rows and not workspace_lists_agreement_for_subject(aid, subject):
            continue
        try:
            d = load_draft(aid)
        except Exception as exc:
            sb_row = supabase_rows.get(aid)
            if sb_row:
                summaries.append(fallback_summary_from_supabase_row(sb_row))
                log.info(
                    "[dashboard-workspace-index-row] agreement_id=%s source=supabase_fallback "
                    "load_draft_failed=%s skipped=false",
                    aid,
                    type(exc).__name__,
                )
                continue
            reason = f"load_draft_failed:{type(exc).__name__}"
            skipped.append({"id": aid, "reason": reason})
            log.warning("[dashboard-workspace-index-row] skipped agreement_id=%s reason=%s", aid, reason)
            continue
        try:
            parties = d.get("parties") or []
            signers = sum(
                1 for p in parties if str((p or {}).get("role") or "").lower() == "signer"
            )
            audit = d.get("audit_log") or []
            signed = any(
                isinstance(e, dict) and str(e.get("event_type") or "") == "signed"
                for e in audit
            )
            reviewer_approved = any(
                isinstance(e, dict)
                and str(e.get("event_type") or "")
                in ("recipient_approved", "participant_approved")
                for e in audit
            )
            appr_done, appr_req, all_reviewers_approved = _workspace_review_approval_rollups(d)
            lock = read_signing_lock(aid)
            lv_raw = str((lock or {}).get("locked_version_id") or "").strip()
            lv: Optional[str] = lv_raw or None
            wfid_raw = str(d.get("workspace_folder_id") or "").strip()
            wfid: Optional[str] = wfid_raw or None
            tags_raw = d.get("workspace_tags")
            tags_out: List[str] = []
            if isinstance(tags_raw, list):
                tags_out = _normalize_workspace_tags(list(tags_raw))
            review_status = "review_approved" if all_reviewers_approved else (
                "in_review" if reviewer_approved or d.get("review_sent_at") else "draft"
            )
            signing_status = "fully_signed" if signed else (
                "signing_locked" if lock is not None and bool(lv) else "unsigned"
            )
            log.info(
                "[dashboard-workspace-index-row] agreement_id=%s source=workspace_index "
                "review_status=%s signing_status=%s skipped=false",
                aid,
                review_status,
                signing_status,
            )
            summaries.append(
                {
                    "id": aid,
                    "title": (str(d.get("title") or "").strip() or "Untitled agreement"),
                    "created_at": str(d.get("created_at") or ""),
                    "updated_at": str(d.get("updated_at") or d.get("created_at") or ""),
                    "party_count": len(parties),
                    "signer_count": signers,
                    "version_ledger_count": len(d.get("versions") or []),
                    "completed_signed": signed,
                    "has_server_signing_lock": lock is not None and bool(lv),
                    "locked_version_id": lv,
                    "workspace_archived_at": d.get("workspace_archived_at"),
                    "review_sent_at": d.get("review_sent_at"),
                    "reviewer_approved": reviewer_approved,
                    "review_approvals_completed": appr_done,
                    "review_approvals_required": appr_req,
                    "all_reviewers_approved": all_reviewers_approved,
                    "workspace_folder_id": wfid,
                    "workspace_folder_name": (folder_names.get(wfid) if wfid else None),
                    "workspace_tags": tags_out,
                    "dashboard_source": "draft",
                    "content_unavailable": False,
                }
            )
        except Exception as exc:
            reason = f"summary_build_failed:{type(exc).__name__}"
            skipped.append({"id": aid, "reason": reason})
            log.exception(
                "[dashboard-workspace-index-row] skipped agreement_id=%s reason=%s",
                aid,
                reason,
            )
            continue
    return {"agreements": summaries, "skipped": skipped}


@router.put("/{agreement_id}/signing-lock")
def put_agreement_signing_lock(
    agreement_id: str, body: AgreementSigningLockBody, request: Request
) -> Dict[str, Any]:
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="signing_lock")
    try:
        assert_draft_exists(agreement_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="agreement_not_found")
    draft_full = _load_or_404(agreement_id)
    draft_full = _persist_party_id_backfill(draft_full)
    missing_signer_approvals = _signing_approval_gate_errors(draft_full)
    if missing_signer_approvals:
        raise HTTPException(
            status_code=400,
            detail={
                "code": "approvals_incomplete",
                "missing_signer_approvals": missing_signer_approvals,
            },
        )
    content_sha256 = _draft_locked_content_sha256(draft_full)
    payload = {
        "locked_version_id": body.locked_version_id,
        "locked_at": body.locked_at,
        "locked_by": body.locked_by,
        "content_sha256": content_sha256,
    }
    write_signing_lock(agreement_id, payload)
    record_public_feed_event_if_applicable(
        draft_dict=draft_full.model_dump(),
        event_type="finalized",
        at=_utc_now_iso(),
    )
    return {"ok": True, "signing_lock": payload}


@router.delete("/{agreement_id}/signing-lock")
def delete_agreement_signing_lock(agreement_id: str, request: Request) -> Dict[str, Any]:
    """Clear signing lock so the owner can edit the draft again (invalidates prior sign-only tokens)."""
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="signing_unlock")
    try:
        assert_draft_exists(agreement_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="agreement_not_found")
    clear_signing_lock(agreement_id)
    return {"ok": True}


@router.post("/{agreement_id}/signing-ceremony/start")
def post_signing_ceremony_start(
    agreement_id: str, body: SigningCeremonyStartBody, request: Request
) -> Dict[str, Any]:
    """Record signature_initiated + return the version hash for the locked signing snapshot."""
    assert_free_incomplete_draft_not_expired(agreement_id, surface="signing_ceremony_start")
    draft = _load_or_404(agreement_id)
    lv, draft = _signing_ceremony_guards_lv(draft, agreement_id)
    part_id, sp = _resolve_signing_participant_for_ceremony(draft, body.participant_id)
    assert_agreement_recipient_write_allowed(
        request,
        agreement_id,
        allowed_modes=("sign",),
        bind_participant_id=part_id,
    )
    now = _utc_now_iso()
    fp = _agreement_version_hash(agreement_id, lv, draft)
    audit = list(draft.audit_log or [])
    audit.append(
        AuditEvent(
            event_type="signature_initiated",
            at=now,
            field="signing",
            value={
                "participant_id": part_id or None,
                "participant_display_name": sp.name,
                "locked_version_id": lv,
                "agreement_version_hash": fp,
            },
        ).model_dump()
    )
    next_draft = _merge_agreement_draft(draft, updated_at=now, audit_log=audit)
    _save_draft_sync(next_draft.model_dump(), request)
    return {
        "ok": True,
        "locked_version_id": lv,
        "agreement_version_hash": fp,
        "participant_display_name": sp.name,
    }


@router.post("/{agreement_id}/signing-ceremony/complete")
def post_signing_ceremony_complete(
    agreement_id: str, body: SigningCeremonyCompleteBody, request: Request
) -> Dict[str, Any]:
    """Record signature_completed; append legacy signed when every signer has completed."""
    assert_free_incomplete_draft_not_expired(agreement_id, surface="signing_ceremony_complete")
    draft = _load_or_404(agreement_id)
    lv, draft = _signing_ceremony_guards_lv(draft, agreement_id)
    req_lv = str(body.locked_version_id or "").strip()
    if req_lv != lv:
        raise HTTPException(status_code=400, detail="locked_version_mismatch")
    lock_row = read_signing_lock(agreement_id)
    stored_sha = str((lock_row or {}).get("content_sha256") or "").strip()
    if stored_sha and _draft_locked_content_sha256(draft) != stored_sha:
        raise HTTPException(status_code=409, detail="stale_locked_version")
    part_id, sp = _resolve_signing_participant_for_ceremony(draft, body.participant_id)
    assert_agreement_recipient_write_allowed(
        request,
        agreement_id,
        allowed_modes=("sign",),
        bind_participant_id=part_id,
    )
    if part_id and part_id in _signature_completed_participant_ids(draft.audit_log):
        raise HTTPException(status_code=409, detail="already_signed")
    if not part_id and _has_legacy_signature_without_participant(draft.audit_log):
        raise HTTPException(status_code=409, detail="already_signed")
    now = _utc_now_iso()
    fp = _agreement_version_hash(agreement_id, lv, draft)
    typed = (body.typed_name or "").strip()
    audit = list(draft.audit_log or [])
    audit.append(
        AuditEvent(
            event_type="signature_completed",
            at=now,
            field="signing",
            value={
                "participant_id": part_id or None,
                "participant_display_name": sp.name,
                "typed_name": typed or None,
                "locked_version_id": lv,
                "agreement_version_hash": fp,
            },
        ).model_dump()
    )
    fully = _all_signers_signed_from_audit(draft, audit)
    if fully:
        assert_can_complete_agreement(agreement_id=agreement_id)
        audit.append(
            AuditEvent(
                event_type="signed",
                at=now,
                field="signing",
                value={"fully_executed": True, "agreement_version_hash": fp},
            ).model_dump()
        )
    next_draft = _merge_agreement_draft(draft, updated_at=now, audit_log=audit)
    dump = next_draft.model_dump()
    _save_draft_sync(dump, request)
    if fully:
        record_agreement_finalized(agreement_id=agreement_id)
        record_public_feed_event_if_applicable(draft_dict=dump, event_type="signed", at=now)
        try:
            from backend.integrations.hooks_emit import (
                claw_emit_integration_event,
                claw_org_id_for_registered_agreement,
            )

            oid = claw_org_id_for_registered_agreement(agreement_id)
            if oid:
                fp_short = fp[:24] if fp else ""
                claw_emit_integration_event(
                    oid,
                    "agreement.signed",
                    "agreement",
                    agreement_id,
                    {"locked_version_id": lv, "agreement_version_hash_prefix": fp_short},
                )
                claw_emit_integration_event(
                    oid,
                    "agreement.completed",
                    "agreement",
                    agreement_id,
                    {"locked_version_id": lv, "lifecycle": "fully_executed"},
                )
        except Exception:
            pass
    return {
        "ok": True,
        "participant_display_name": sp.name,
        "signed_at": now,
        "agreement_version_hash": fp,
        "fully_executed": fully,
    }


@router.post("/{agreement_id}/recipient-access-token")
def post_recipient_access_token(
    agreement_id: str, request: Request, body: RecipientAccessMintRequest
) -> Dict[str, Any]:
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="recipient_access_token")
    if not _recipient_link_mint_key_ok(request):
        raise HTTPException(status_code=403, detail="recipient_link_mint_key_invalid")
    try:
        assert_draft_exists(agreement_id)
    except KeyError:
        raise HTTPException(status_code=404, detail="agreement_not_found")

    if body.mode == "sign" and body.role == "reviewer":
        raise HTTPException(
            status_code=400,
            detail="reviewer_role_incompatible_with_sign_mode",
        )

    lock = read_signing_lock(agreement_id)
    env_ttl = os.getenv("CLAW_RECIPIENT_TOKEN_TTL_SECONDS", "").strip()
    raw_ttl = int(env_ttl) if env_ttl else int(body.ttl_seconds)
    ttl = clamp_recipient_token_ttl_seconds(raw_ttl)

    if body.mode == "sign":
        if not lock or not lock.get("locked_version_id"):
            raise HTTPException(status_code=409, detail="signing_not_finalized_server_side")
        lv = str(lock["locked_version_id"])
    else:
        lv = str((lock or {}).get("locked_version_id") or "")

    secret = _signing_token_secret_bytes(agreement_id=agreement_id)
    token: Optional[str] = None
    last_mint_error: Optional[BaseException] = None
    for attempt in range(3):  # initial + 2 retries (short backoff)
        try:
            token = mint_recipient_access_token(
                secret=secret,
                agreement_id=agreement_id,
                locked_version_id=lv,
                mode=body.mode,
                role=body.role,
                ttl_seconds=ttl,
                recipient_subject=body.recipient_subject,
                recipient_party_id=body.recipient_party_id,
                inviter_display_name=body.inviter_display_name,
                single_use=body.single_use,
            )
            break
        except HTTPException:
            raise
        except Exception as exc:  # noqa: BLE001 — bounded retries; never log token material
            last_mint_error = exc
            if attempt < 2:
                time.sleep(0.06 * (2**attempt))
                continue
            break
    if not token:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "recipient_token_mint_unavailable",
                "message": "Unable to mint recipient access token after retries. Try again shortly.",
            },
        ) from last_mint_error
    u_ev = "signature_request_sent" if body.mode == "sign" else "recipient_invite_sent"
    record_usage_ledger_event(
        subject_ref=resolve_subject_from_request(request),
        event_type=u_ev,
        agreement_id=agreement_id,
        metadata={"role": body.role},
    )
    _persist_review_first_final_corpus_if_supplied(
        agreement_id,
        body,
        locked_version_id=lv,
        subject_ref=resolve_subject_from_request(request),
    )
    return {
        "token": token,
        "expires_in_seconds": ttl,
        "mode": body.mode,
        "locked_version_id": lv,
    }


@router.post("/{agreement_id}/review-delivery-dry-run")
def post_agreement_review_delivery_dry_run(agreement_id: str, request: Request) -> Dict[str, Any]:
    """Owner-only: count + structured email payloads without minting tokens (``review_url`` always null)."""
    require_claw_org_id_header(request)
    subject = resolve_subject_from_request(request)
    aid = (agreement_id or "").strip()
    if not aid or not workspace_lists_agreement_for_subject(aid, subject):
        raise HTTPException(status_code=404, detail={"code": "agreement_not_found", "message": "Not found"})
    draft_full = _load_or_404(aid)
    d = draft_full.model_dump(mode="json")
    from backend.config.runtime_environment import review_delivery_mode as _review_delivery_mode

    mode = _review_delivery_mode()
    payloads = _review_delivery_email_payload_rows_from_draft_dict(d)
    aid_short = aid[:12] if len(aid) >= 12 else aid
    log.info(
        "[review-delivery-dry-run] agreement_id_short=%s payload_count=%s mode=%s",
        aid_short,
        len(payloads),
        mode,
    )
    return {"review_delivery_mode": mode, "payload_count": len(payloads), "payloads": payloads}


@router.post("/{agreement_id}/review-sent")
def post_agreement_review_sent(agreement_id: str, request: Request) -> Dict[str, Any]:
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="review_sent")
    draft = _load_or_404(agreement_id)
    invite_emails_already_sent = bool((draft.review_invite_emails_sent_at or "").strip())
    log.info(
        "[review-sent] start agreement_id=%s org_id=%s review_sent_at_present=%s "
        "review_invite_emails_sent_at_present=%s",
        agreement_id,
        resolve_subject_from_request(request),
        bool((draft.review_sent_at or "").strip()),
        invite_emails_already_sent,
    )
    now = _utc_now_iso()
    next_data = draft.model_dump()
    next_data["review_sent_at"] = now
    next_data["updated_at"] = now
    next_draft = AgreementDraft.model_validate(next_data)
    _save_draft_sync(next_draft.model_dump(), request)
    try:
        from backend.integrations.hooks_emit import claw_emit_integration_event_from_subject

        claw_emit_integration_event_from_subject(
            resolve_subject_from_request(request),
            "agreement.sent",
            "agreement",
            agreement_id,
            {"review_sent_at": now},
        )
    except Exception:
        pass
    try:
        from backend.affiliates.influence import maybe_record_agreement_sent_influence
        from backend.utils.enforce import org_id_from_subject

        subj = resolve_subject_from_request(request)
        maybe_record_agreement_sent_influence(org_id_from_subject(subj))
    except Exception:
        pass
    try:
        from backend.services.email.review_delivery import maybe_send_review_invites_after_review_sent

        if invite_emails_already_sent:
            log.info(
                "[review-email-delivery] skipped agreement_id=%s org_id=%s skip_reason=invite_emails_already_sent",
                agreement_id,
                resolve_subject_from_request(request),
            )
        else:
            email_draft = _load_or_404(agreement_id).model_dump(mode="json")
            delivery_marker = maybe_send_review_invites_after_review_sent(
                agreement_id=agreement_id,
                draft=email_draft,
                org_id=resolve_subject_from_request(request),
            )
            if delivery_marker:
                marked = next_draft.model_dump()
                marked["review_invite_emails_sent_at"] = delivery_marker
                marked["updated_at"] = _utc_now_iso()
                if email_draft.get("recipient_delivery_v1"):
                    marked["recipient_delivery_v1"] = email_draft["recipient_delivery_v1"]
                if email_draft.get("audit_log"):
                    marked["audit_log"] = email_draft["audit_log"]
                next_draft = AgreementDraft.model_validate(marked)
                _save_draft_sync(next_draft.model_dump(), request)
    except Exception:
        pass
    return {"ok": True, "draft": next_draft.model_dump()}


@router.post("/{agreement_id}/signing-links-sent")
def post_agreement_signing_links_sent(
    agreement_id: str,
    request: Request,
    body: SigningLinksSentBody = SigningLinksSentBody(),
) -> Dict[str, Any]:
    """Owner-triggered VS01 signing invite emails after packet prepare (parallel signing default)."""
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="signing_links_sent")
    draft = _load_or_404(agreement_id)
    sent_count = 0
    skip_reason: str | None = None
    portable = body.portable_packet if isinstance(body.portable_packet, dict) else None
    document_id = (body.document_id or "").strip() or None
    packet_revision = (body.packet_revision or "").strip() or None
    if portable and document_id:
        try:
            stored = {
                "v": 1,
                "document_id": document_id,
                "packet_revision": packet_revision,
                "portable": portable,
                "stored_at": _utc_now_iso(),
            }
            draft = _merge_agreement_draft(draft, vs01_signing_packet_v1=stored, updated_at=_utc_now_iso())
            _save_draft_sync(draft.model_dump(), request)
        except Exception:
            logging.getLogger(__name__).exception(
                "vs01_signing_packet_persist_failed agreement_id=%s",
                agreement_id,
            )
    try:
        from backend.services.email.signing_delivery import maybe_send_signing_invites_after_packet_prepared

        email_draft = draft.model_dump(mode="json")
        notify_audit = maybe_send_signing_invites_after_packet_prepared(
            agreement_id=agreement_id,
            draft=email_draft,
            targets=[t.model_dump() for t in (body.targets or [])],
            packet_revision=(body.packet_revision or "").strip() or None,
            org_id=resolve_subject_from_request(request),
        )
        if notify_audit:
            value = notify_audit.get("value") if isinstance(notify_audit.get("value"), dict) else {}
            sent_count = int(value.get("sent_count") or 0)
            audit = [*(draft.audit_log or []), AuditEvent.model_validate(notify_audit)]
            merge_fields: Dict[str, Any] = {"updated_at": _utc_now_iso(), "audit_log": audit}
            if email_draft.get("recipient_delivery_v1"):
                merge_fields["recipient_delivery_v1"] = email_draft["recipient_delivery_v1"]
            if len(email_draft.get("audit_log") or []) > len(audit) - 1:
                merge_fields["audit_log"] = email_draft["audit_log"]
            next_draft = _merge_agreement_draft(draft, **merge_fields)
            _save_draft_sync(next_draft.model_dump(), request)
            return {"ok": True, "sent_count": sent_count, "skip_reason": None, "draft": next_draft.model_dump()}
        skip_reason = "not_sent"
    except Exception:
        logging.getLogger(__name__).exception(
            "signing_invite_delivery_hook_failed agreement_id=%s",
            agreement_id,
        )
        skip_reason = "delivery_error"
    return {
        "ok": True,
        "sent_count": sent_count,
        "skip_reason": skip_reason,
        "draft": draft.model_dump(),
    }


@router.post("/{agreement_id}/vs01-signer-complete")
def post_vs01_signer_complete(
    agreement_id: str,
    request: Request,
    body: Vs01SignerCompleteBody = Vs01SignerCompleteBody(),
) -> Dict[str, Any]:
    """Record VS01 signer completion; append fully-executed + completion emails when all signers finish."""
    aid = (agreement_id or "").strip()
    signer_role_id = (body.signer_role_id or "").strip()
    if not signer_role_id:
        raise HTTPException(status_code=400, detail="signer_role_id_required")

    auth_mode: Optional[str] = None
    if _agreements_write_allowed():
        try:
            _owner_mutation_guards(request, aid, surface="vs01_signer_complete")
            auth_mode = "owner"
        except HTTPException:
            auth_mode = None
    if not auth_mode:
        from backend.security.agreement_read_scope import recipient_access_token_from_request
        from backend.services.vs01_signer_completion import vs01_open_signing_link_completion_allowed

        tok = recipient_access_token_from_request(request)
        if tok:
            assert_agreement_recipient_write_allowed(request, aid, allowed_modes=("sign",))
            auth_mode = "recipient"
        else:
            draft_for_auth = _load_or_404(aid)
            if vs01_open_signing_link_completion_allowed(
                draft_for_auth.model_dump(),
                signer_role_id=signer_role_id,
                document_id=(body.document_id or "").strip(),
            ):
                auth_mode = "signing_link"
            else:
                assert_agreement_recipient_write_allowed(request, aid, allowed_modes=("sign",))
                auth_mode = "recipient"

    from backend.services.vs01_signer_completion import (
        completion_emails_already_sent,
        merge_fresh_audit_for_vs01_signer,
        orchestrate_vs01_signer_complete,
        resolve_participant_id_for_signer_role,
    )

    draft = _load_or_404(aid)
    now = (body.signed_at or "").strip() or _utc_now_iso()
    signed_date_iso = (body.signed_date_iso or "").strip() or now[:10]
    signed_date_display = (body.signed_date_display or "").strip()
    if not signed_date_display and signed_date_iso:
        try:
            from backend.services.vs01_fully_executed_snapshot import _format_signing_date_display

            signed_date_display = _format_signing_date_display(signed_date_iso)
        except Exception:
            signed_date_display = signed_date_iso
    participant_id = resolve_participant_id_for_signer_role(
        draft.model_dump(),
        signer_role_id,
        body.participant_id or "",
    )
    lock_row = read_signing_lock(aid)
    lv = str((lock_row or {}).get("locked_version_id") or "").strip() or None
    fp = _agreement_version_hash(aid, lv, draft) if lv else None
    display_name = (body.display_name or "").strip()
    if not display_name and participant_id:
        sp = _signer_party_by_participant_id(draft, participant_id)
        if sp:
            display_name = (sp.name or "").strip()

    portable_packet = body.portable_packet if isinstance(body.portable_packet, dict) else None

    pending = orchestrate_vs01_signer_complete(
        draft.model_dump(),
        signer_role_id=signer_role_id,
        participant_id=participant_id,
        display_name=display_name,
        document_id=(body.document_id or "").strip(),
        signed_at=now,
        signed_date_iso=signed_date_iso,
        signed_date_display=signed_date_display,
        locked_version_id=lv,
        agreement_version_hash=fp,
        portable_packet=portable_packet,
    )

    fresh = _load_or_404(aid)
    outcome = merge_fresh_audit_for_vs01_signer(
        fresh.model_dump(),
        pending,
        signer_role_id=signer_role_id,
        portable_packet=portable_packet,
    )

    completion_emails_sent = completion_emails_already_sent(outcome.audit)

    if outcome.audit_mutated:
        next_draft = _merge_agreement_draft(
            fresh,
            updated_at=now,
            audit_log=outcome.audit,
            vs01_signing_packet_v1=outcome.draft_dict.get("vs01_signing_packet_v1"),
        )
        _save_draft_sync(next_draft.model_dump(), request)

        if outcome.newly_finalized:
            assert_can_complete_agreement(agreement_id=aid)
            record_agreement_finalized(agreement_id=aid)
            record_public_feed_event_if_applicable(
                draft_dict=next_draft.model_dump(),
                event_type="signed",
                at=now,
            )
            try:
                from backend.integrations.hooks_emit import (
                    claw_emit_integration_event,
                    claw_org_id_for_registered_agreement,
                )

                oid = claw_org_id_for_registered_agreement(aid)
                if oid:
                    fp_short = fp[:24] if fp else ""
                    claw_emit_integration_event(
                        oid,
                        "agreement.signed",
                        "agreement",
                        aid,
                        {"locked_version_id": lv, "agreement_version_hash_prefix": fp_short},
                    )
                    claw_emit_integration_event(
                        oid,
                        "agreement.completed",
                        "agreement",
                        aid,
                        {"locked_version_id": lv, "lifecycle": "fully_executed"},
                    )
            except Exception:
                pass

    if outcome.fully_executed:
        from backend.services.vs01_fully_executed_snapshot import ensure_fully_executed_snapshot_on_draft

        reloaded_for_snap = _load_or_404(aid)
        ensured = ensure_fully_executed_snapshot_on_draft(
            reloaded_for_snap.model_dump(),
            agreement_id=aid,
        )
        if ensured.mutated:
            snap_draft = _merge_agreement_draft(
                reloaded_for_snap,
                updated_at=now,
                vs01_signing_packet_v1=ensured.draft_dict.get("vs01_signing_packet_v1"),
            )
            _save_draft_sync(snap_draft.model_dump(), request)

    if outcome.fully_executed:
        from backend.services.vs01_signer_completion import vs01_completion_email_lock

        with vs01_completion_email_lock(aid):
            reloaded = _load_or_404(aid)
            reloaded_audit = list(reloaded.model_dump().get("audit_log") or [])
            if completion_emails_already_sent(reloaded_audit):
                completion_emails_sent = True
            else:
                try:
                    from backend.services.email.signing_completion_delivery import (
                        maybe_send_signing_completion_emails,
                    )

                    notify_audit = maybe_send_signing_completion_emails(
                        agreement_id=aid,
                        draft={**reloaded.model_dump(), "audit_log": reloaded_audit},
                        org_id=resolve_subject_from_request(request),
                    )
                    if notify_audit:
                        fresh_for_email = _load_or_404(aid)
                        fresh_audit = list(fresh_for_email.model_dump().get("audit_log") or [])
                        if not completion_emails_already_sent(fresh_audit):
                            email_audit = list(fresh_audit)
                            email_audit.append(AuditEvent.model_validate(notify_audit).model_dump())
                            next_email = _merge_agreement_draft(
                                fresh_for_email,
                                updated_at=now,
                                audit_log=email_audit,
                            )
                            _save_draft_sync(next_email.model_dump(), request)
                            completion_emails_sent = True
                        else:
                            completion_emails_sent = True
                except Exception:
                    logging.getLogger(__name__).exception(
                        "vs01_signing_completion_email_failed agreement_id=%s",
                        aid,
                    )

    return {
        "ok": True,
        "already_signed": outcome.already_signed,
        "fully_executed": outcome.fully_executed,
        "completion_emails_sent": completion_emails_sent,
        "auth_mode": auth_mode,
    }


@router.post("/{agreement_id}/vs01-ensure-signed-snapshot")
def post_vs01_ensure_signed_snapshot(agreement_id: str, request: Request) -> Dict[str, Any]:
    """Ensure fully-executed signed snapshot exists; retry completion emails when safe."""
    aid = (agreement_id or "").strip()
    if _agreements_write_allowed():
        _owner_mutation_guards(request, aid, surface="vs01_ensure_signed_snapshot")
    else:
        raise HTTPException(status_code=403, detail="forbidden")

    from backend.services.vs01_fully_executed_snapshot import ensure_fully_executed_snapshot_on_draft
    from backend.services.vs01_signer_completion import (
        all_signers_signed_from_audit,
        completion_emails_already_sent,
        vs01_completion_email_lock,
    )

    draft = _load_or_404(aid)
    audit = list(draft.model_dump().get("audit_log") or [])
    if not all_signers_signed_from_audit(draft.model_dump(), audit):
        raise HTTPException(status_code=409, detail="agreement_not_fully_executed")

    now = _utc_now_iso()
    ensured = ensure_fully_executed_snapshot_on_draft(draft.model_dump(), agreement_id=aid)
    if ensured.mutated:
        next_draft = _merge_agreement_draft(
            draft,
            updated_at=now,
            vs01_signing_packet_v1=ensured.draft_dict.get("vs01_signing_packet_v1"),
        )
        _save_draft_sync(next_draft.model_dump(), request)
        draft = next_draft

    completion_emails_sent = completion_emails_already_sent(draft.model_dump().get("audit_log") or [])
    if ensured.snapshot_ready and not completion_emails_sent:
        with vs01_completion_email_lock(aid):
            reloaded = _load_or_404(aid)
            reloaded_audit = list(reloaded.model_dump().get("audit_log") or [])
            if not completion_emails_already_sent(reloaded_audit):
                try:
                    from backend.services.email.signing_completion_delivery import (
                        maybe_send_signing_completion_emails,
                    )

                    notify_audit = maybe_send_signing_completion_emails(
                        agreement_id=aid,
                        draft={**reloaded.model_dump(), "audit_log": reloaded_audit},
                        org_id=resolve_subject_from_request(request),
                    )
                    if notify_audit:
                        fresh_for_email = _load_or_404(aid)
                        fresh_audit = list(fresh_for_email.model_dump().get("audit_log") or [])
                        if not completion_emails_already_sent(fresh_audit):
                            email_audit = list(fresh_audit)
                            email_audit.append(AuditEvent.model_validate(notify_audit).model_dump())
                            next_email = _merge_agreement_draft(
                                fresh_for_email,
                                updated_at=now,
                                audit_log=email_audit,
                            )
                            _save_draft_sync(next_email.model_dump(), request)
                            completion_emails_sent = True
                except Exception:
                    logging.getLogger(__name__).exception(
                        "vs01_ensure_signed_snapshot_email_failed agreement_id=%s",
                        aid,
                    )

    return {
        "ok": True,
        "snapshot_ready": ensured.snapshot_ready,
        "snapshot_source": ensured.source,
        "completion_emails_sent": completion_emails_sent,
    }


@router.get("/{agreement_id}/recipient-delivery-status")
def get_recipient_delivery_status(agreement_id: str, request: Request) -> JSONResponse:
    """Owner-facing per-recipient review/signing delivery rows — never HTTP 500."""
    aid = (agreement_id or "").strip()
    from backend.build_info import RECIPIENT_DELIVERY_STATUS_HANDLER_REV, git_commit_short
    from backend.services.recipient_delivery_status import (
        build_recipient_delivery_status,
        degraded_recipient_delivery_payload,
        draft_diagnostic_types,
        recipient_delivery_json_response,
        _log_stage,
    )

    _log_stage(
        agreement_id=aid,
        stage="handler_enter",
        extra={
            "handler_rev": RECIPIENT_DELIVERY_STATUS_HANDLER_REV,
            "git_commit": git_commit_short(),
        },
    )

    draft: Optional[Dict[str, Any]] = None
    try:
        if not _agreements_write_allowed():
            raise HTTPException(status_code=403, detail="verifier_only")
        _log_stage(agreement_id=aid, stage="guards_start")
        _owner_mutation_guards(request, aid, surface="recipient_delivery_status")
        _log_stage(agreement_id=aid, stage="load_draft")
        raw = _load_draft_dict_or_404(aid)
        draft = raw
        _log_stage(agreement_id=aid, stage="load_draft_ok", draft=raw, extra=draft_diagnostic_types(raw))
        payload = build_recipient_delivery_status(raw, agreement_id=aid)
        payload["agreement_id"] = aid
        payload["degraded"] = False
        _log_stage(
            agreement_id=aid,
            stage="serialize_response",
            draft=raw,
            extra={"recipient_row_count": len(payload.get("recipients") or [])},
        )
        return recipient_delivery_json_response(payload, agreement_id=aid)
    except HTTPException as exc:
        if exc.status_code in (401, 403):
            raise
        _agreements_log.warning(
            "[recipient-delivery-status-error] agreement_id=%s exception_type=HTTPException "
            "status_code=%s stage=route_http detail=%s",
            aid,
            exc.status_code,
            str(exc.detail)[:300],
        )
        return recipient_delivery_json_response(
            degraded_recipient_delivery_payload(aid, draft, error="recipient_status_degraded"),
            agreement_id=aid,
        )
    except Exception as exc:
        _agreements_log.error(
            "[recipient-delivery-status-error] agreement_id=%s exception_type=%s exception_message=%s "
            "stage=route traceback=%s",
            aid,
            type(exc).__name__,
            str(exc)[:500],
            traceback.format_exc(),
        )
        if draft is None:
            try:
                from backend.services.agreement_draft_store import load_draft

                raw = load_draft(aid)
                if isinstance(raw, dict):
                    draft = raw
            except Exception:
                draft = None
        try:
            if isinstance(draft, dict):
                payload = build_recipient_delivery_status(draft, agreement_id=aid)
                payload["agreement_id"] = aid
                payload["degraded"] = False
                return recipient_delivery_json_response(payload, agreement_id=aid)
        except Exception as retry_exc:
            _agreements_log.error(
                "[recipient-delivery-status-error] agreement_id=%s exception_type=%s exception_message=%s "
                "stage=route_retry traceback=%s",
                aid,
                type(retry_exc).__name__,
                str(retry_exc)[:500],
                traceback.format_exc(),
            )
        return recipient_delivery_json_response(
            degraded_recipient_delivery_payload(aid, draft, error="recipient_status_degraded"),
            agreement_id=aid,
        )


@router.post("/{agreement_id}/recipient-invite-resend")
def post_recipient_invite_resend(
    agreement_id: str,
    request: Request,
    body: RecipientInviteResendBody = RecipientInviteResendBody(),
) -> Dict[str, Any]:
    """Resend a review or signing invite without changing email or agreement corpus."""
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="recipient_invite_resend")
    draft = _load_or_404(agreement_id)
    _assert_negotiation_not_locked(agreement_id)
    from backend.services.recipient_invite_resend import resend_recipient_invite

    next_data, meta = resend_recipient_invite(
        agreement_id=agreement_id,
        draft=draft.model_dump(mode="json"),
        phase=body.phase,
        participant_id=body.participant_id,
        signing_url=body.signing_url,
        signer_role_id=body.signer_role_id,
        org_id=resolve_subject_from_request(request),
    )
    next_draft = AgreementDraft.model_validate(next_data)
    _save_draft_sync(next_draft.model_dump(), request)
    return {"ok": True, "draft": next_draft.model_dump(), **meta}


@router.post("/{agreement_id}/review-recipient-email")
def post_review_recipient_email_correction(
    agreement_id: str,
    request: Request,
    body: ReviewRecipientEmailCorrectBody = ReviewRecipientEmailCorrectBody(),
) -> Dict[str, Any]:
    """Owner corrects a reviewer's mistyped email before approval; optionally resends review invite."""
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="review_recipient_email")
    draft = _load_or_404(agreement_id)
    _assert_negotiation_not_locked(agreement_id)
    _assert_draft_mutable_after_signatures(draft)
    from backend.services.recipient_email_correction import correct_review_recipient_email

    next_data, meta = correct_review_recipient_email(
        agreement_id=agreement_id,
        draft=draft.model_dump(mode="json"),
        participant_id=body.participant_id,
        new_email=body.new_email,
        resend_invite=body.resend_invite,
        org_id=resolve_subject_from_request(request),
    )
    next_draft = AgreementDraft.model_validate(next_data)
    _save_draft_sync(next_draft.model_dump(), request)
    return {"ok": True, "draft": next_draft.model_dump(), **meta}


@router.post("/{agreement_id}/signing-recipient-email")
def post_signing_recipient_email_correction(
    agreement_id: str,
    request: Request,
    body: SigningRecipientEmailCorrectBody = SigningRecipientEmailCorrectBody(),
) -> Dict[str, Any]:
    """Owner corrects a signer's mistyped email; optionally resends that signer's signing invite."""
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="signing_recipient_email")
    draft = _load_or_404(agreement_id)
    _assert_negotiation_not_locked(agreement_id)
    from backend.services.recipient_email_correction import correct_signing_recipient_email

    next_data, meta = correct_signing_recipient_email(
        agreement_id=agreement_id,
        draft=draft.model_dump(mode="json"),
        participant_id=body.participant_id,
        new_email=body.new_email,
        signer_role_id=body.signer_role_id,
        signing_url=body.signing_url,
        resend_invite=body.resend_invite,
        org_id=resolve_subject_from_request(request),
    )
    next_draft = AgreementDraft.model_validate(next_data)
    _save_draft_sync(next_draft.model_dump(), request)
    return {"ok": True, "draft": next_draft.model_dump(), **meta}


@router.patch("/{agreement_id}/workspace-archive")
def patch_agreement_workspace_archive(
    agreement_id: str, body: WorkspaceArchiveBody, request: Request
) -> Dict[str, Any]:
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="workspace_archive")
    draft = _load_or_404(agreement_id)
    now = _utc_now_iso()
    next_data = draft.model_dump()
    next_data["workspace_archived_at"] = now if body.archived else None
    next_data["updated_at"] = now
    next_draft = AgreementDraft.model_validate(next_data)
    _save_draft_sync(next_draft.model_dump(), request)
    return {"ok": True, "draft": next_draft.model_dump()}


@router.patch("/{agreement_id}/workspace-folder")
def patch_agreement_workspace_folder(
    agreement_id: str, body: WorkspaceFolderAssignBody, request: Request
) -> Dict[str, Any]:
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="workspace_folder")
    draft = _load_or_404(agreement_id)
    subject = resolve_subject_from_request(request)
    fid = (body.folder_id or "").strip() or None
    if fid:
        st = ProofLayerStore()
        st.init_schema()
        if not st.get_folder(subject, fid):
            raise HTTPException(status_code=400, detail="unknown_workspace_folder")
    now = _utc_now_iso()
    next_data = draft.model_dump()
    next_data["workspace_folder_id"] = fid
    next_data["updated_at"] = now
    next_draft = AgreementDraft.model_validate(next_data)
    _save_draft_sync(next_draft.model_dump(), request)
    return {"ok": True, "draft": next_draft.model_dump()}


@router.patch("/{agreement_id}/workspace-tags")
def patch_agreement_workspace_tags(
    agreement_id: str, body: WorkspaceTagsReplaceBody, request: Request
) -> Dict[str, Any]:
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")
    _owner_mutation_guards(request, agreement_id, surface="workspace_tags")
    draft = _load_or_404(agreement_id)
    now = _utc_now_iso()
    next_data = draft.model_dump()
    next_data["workspace_tags"] = _normalize_workspace_tags(list(body.tags or []))
    next_data["updated_at"] = now
    next_draft = AgreementDraft.model_validate(next_data)
    _save_draft_sync(next_draft.model_dump(), request)
    return {"ok": True, "draft": next_draft.model_dump()}


def _public_agreement_verify_payload(aid: str, draft: AgreementDraft) -> Dict[str, Any]:
    """Full public verify JSON; callers wrap in try/except for graceful degradation."""
    overview_hash = _public_agreement_overview_hash(aid, draft)
    sig_status = _public_signature_status(aid, draft)
    claw_feed: Optional[Dict[str, Any]] = None
    if str(getattr(draft, "feed_visibility", "private") or "private").strip().lower() == "public":
        row = get_claw_feed_store().get_feed_anchor_summary_for_agreement(aid)
        if row:
            claw_feed = {
                "event_type": row.get("event_type"),
                "at": row.get("at"),
                "summary": row.get("summary"),
                "anchor_network": row.get("anchor_network"),
                "anchor_status": row.get("anchor_status"),
                "anchor_txid": row.get("anchor_txid"),
                "batch_id": row.get("batch_id"),
            }
    settlement_net = settlement_anchor_network_hint()
    return {
        "agreement_id": aid,
        "summary": {
            "title": draft.title,
            "jurisdiction": draft.jurisdiction,
            "created_at": draft.created_at,
            "updated_at": draft.updated_at,
            "status": _public_lifecycle_label(draft, aid),
            "review_sent_at": draft.review_sent_at,
        },
        "participants": [{"name": p.name, "role": p.role} for p in draft.parties or []],
        "version_history": _public_version_history(draft),
        "signature_status": sig_status,
        "signature_events": _public_signature_events(draft.audit_log),
        "verification": {
            "agreement_hash": overview_hash,
            "signing_commitment_hash": sig_status.get("signing_commitment_hash"),
            "schema": "claw.agreement.public_verify/v1",
        },
        "claw_feed": claw_feed,
        "settlement_anchor": {
            "network_hint": settlement_net,
            "note": (
                "Periodic settlement anchoring may additionally be recorded on Bitcoin-class networks "
                f"(hint: {settlement_net}) for stronger checkpointing; feed-visible events prefer "
                "lower-cost anchors when enabled."
            ),
        },
    }


def _public_agreement_verify_pending_payload(aid: str, draft: AgreementDraft) -> Dict[str, Any]:
    """Minimal safe bundle when hashes, feed, or anchors cannot be assembled (never 500)."""
    try:
        settlement_net = settlement_anchor_network_hint()
    except Exception:
        settlement_net = ""
    try:
        lifecycle = _public_lifecycle_label(draft, aid)
    except Exception:
        lifecycle = "in_negotiation"
    signer_party_count = 0
    try:
        signer_party_count = len(
            [p for p in draft.parties or [] if _normalize_workflow_role(p.role) == "signer"]
        )
    except Exception:
        signer_party_count = 0
    return {
        "agreement_id": aid,
        "record_status": "pending",
        "record_status_reason": "verification_bundle_incomplete",
        "summary": {
            "title": draft.title,
            "jurisdiction": draft.jurisdiction,
            "created_at": draft.created_at,
            "updated_at": draft.updated_at,
            "status": lifecycle,
            "review_sent_at": draft.review_sent_at,
        },
        "participants": [{"name": p.name, "role": p.role} for p in draft.parties or []],
        "version_history": [],
        "signature_status": {
            "fully_executed": False,
            "signatures_recorded": 0,
            "signer_party_count": signer_party_count,
            "locked_version_id": None,
            "signing_commitment_hash": None,
        },
        "signature_events": [],
        "verification": {
            "agreement_hash": "",
            "signing_commitment_hash": None,
            "schema": "claw.agreement.public_verify/v1",
            "record_note": (
                "Public verification details are not available yet. The agreement exists, but proof "
                "or anchor metadata may still be pending."
            ),
        },
        "claw_feed": None,
        "settlement_anchor": {
            "network_hint": settlement_net or "unknown",
            "note": (
                "Settlement anchor details may be unavailable until the record is finalized."
            ),
        },
    }


@router.get("/public/{agreement_id}/verify")
def get_public_agreement_verify(agreement_id: str) -> Dict[str, Any]:
    """Public read-only verification bundle (no agreement body text, purpose, or payment terms)."""
    if not public_agreement_verify_enabled():
        raise HTTPException(status_code=404, detail="not_found")
    aid = (agreement_id or "").strip()
    if not aid:
        raise HTTPException(status_code=404, detail="not_found")
    try:
        raw = load_draft(aid)
    except KeyError:
        raise HTTPException(
            status_code=404,
            detail={"code": "agreement_not_found", "message": "No agreement exists for this id."},
        )
    except ValueError:
        raise HTTPException(
            status_code=404,
            detail={
                "code": "agreement_not_found",
                "message": "Agreement record is missing or not readable.",
            },
        )
    try:
        draft = AgreementDraft.model_validate(raw)
    except ValidationError:
        raise HTTPException(
            status_code=404,
            detail={
                "code": "agreement_not_found",
                "message": "Agreement record is missing or not readable.",
            },
        )
    try:
        return _public_agreement_verify_payload(aid, draft)
    except Exception:
        return _public_agreement_verify_pending_payload(aid, draft)


@router.get("/public/{agreement_id}/vs01-signing-packet")
def get_public_vs01_signing_packet(
    agreement_id: str,
    document_id: str,
    packet_revision: Optional[str] = None,
    recipient_email: Optional[str] = None,
    participant_id: Optional[str] = None,
) -> Dict[str, Any]:
    """Public VS01 signing packet for emailed recipient links (no auth; fields + corpus seed only)."""
    aid = (agreement_id or "").strip()
    did = (document_id or "").strip()
    if not aid or not did:
        raise HTTPException(status_code=404, detail="not_found")
    try:
        raw = load_draft(aid)
    except (KeyError, ValueError):
        raise HTTPException(status_code=404, detail="not_found")
    try:
        draft = AgreementDraft.model_validate(raw)
    except ValidationError:
        raise HTTPException(status_code=404, detail="not_found")

    pid = (participant_id or "").strip()
    remail = (recipient_email or "").strip()
    if pid and remail:
        from backend.services.recipient_delivery_registry import is_signing_email_superseded
        from backend.security.recipient_access_token import RECIPIENT_INVITE_SUPERSEDED

        if is_signing_email_superseded(raw, pid, remail):
            raise HTTPException(
                status_code=403,
                detail={
                    "code": "invite_superseded",
                    "message": RECIPIENT_INVITE_SUPERSEDED,
                },
            )

    stored = draft.vs01_signing_packet_v1 if isinstance(draft.vs01_signing_packet_v1, dict) else None
    if not stored:
        raise HTTPException(status_code=404, detail="packet_not_found")
    if (stored.get("document_id") or "").strip() != did:
        raise HTTPException(status_code=404, detail="packet_document_mismatch")
    rev = (packet_revision or "").strip()
    stored_rev = (stored.get("packet_revision") or "").strip()
    if rev and stored_rev and rev != stored_rev:
        raise HTTPException(status_code=404, detail="packet_revision_mismatch")
    portable = stored.get("portable")
    if not isinstance(portable, dict):
        raise HTTPException(status_code=404, detail="packet_invalid")
    return {"ok": True, "portable": portable}


@router.get("/{agreement_id}")
def get_agreement_draft(agreement_id: str, request: Request) -> Dict[str, Any]:
    assert_agreement_full_draft_read_allowed(request, agreement_id)
    draft = _load_or_404(agreement_id)
    lock = read_signing_lock(agreement_id)
    lv = str((lock or {}).get("locked_version_id") or "").strip()
    signing_lock_out: Optional[Dict[str, Any]] = None
    if lock and lv:
        signing_lock_out = {
            "locked_version_id": lv,
            "locked_at": lock.get("locked_at"),
            "locked_by": lock.get("locked_by"),
            "content_sha256": lock.get("content_sha256"),
        }
    return {
        "id": agreement_id,
        "draft": _draft_with_sanitized_parties(draft).model_dump(),
        "economics": economics_overlay_for_agreement(agreement_id),
        "signing_lock": signing_lock_out,
    }


@router.post("/{agreement_id}/update-field")
def update_agreement_field(
    agreement_id: str, body: AgreementFieldUpdateRequest, request: Request
) -> Dict[str, Any]:
    _owner_mutation_guards(request, agreement_id, surface="update_field")
    draft = _load_or_404(agreement_id)
    _assert_negotiation_not_locked(agreement_id)
    _assert_draft_mutable_after_signatures(draft)
    if not body.field:
        raise HTTPException(status_code=400, detail="field_required")
    if body.field not in {
        "title",
        "jurisdiction",
        "parties",
        "purpose",
        "payment_terms",
        "duration",
        "due_date",
        "effective_date",
        "feed_visibility",
        "feed_party_anonymize",
        "feed_show_financial_summary",
        "feed_anchor_network",
        "payment_request",
        "payment_required",
    }:
        raise HTTPException(status_code=400, detail="unsupported_field")

    next_data = draft.model_dump()
    now = _utc_now_iso()
    audit_log = list(next_data.get("audit_log") or [])
    if body.field == "parties":
        prior_ids = {(p.id or "").strip() for p in (draft.parties or []) if (p.id or "").strip()}
        parties_raw: List[AgreementParty] = []
        if isinstance(body.value, list):
            prior_by_id = {(p.id or "").strip(): p for p in (draft.parties or []) if (p.id or "").strip()}
            for p in body.value:
                if not isinstance(p, dict):
                    continue
                name = str(p.get("name") or "").strip()
                role = str(p.get("role") or "party").strip() or "party"
                pid = str(p.get("id") or "").strip() or str(uuid.uuid4())
                prior = prior_by_id.get(pid)
                email_raw = str(p.get("email") or "").strip()
                email = email_raw or (str(prior.email or "").strip() if prior else "") or None
                phone_raw = str(p.get("phone") or "").strip()
                phone = phone_raw or (str(prior.phone or "").strip() if prior else "") or None
                parties_raw.append(
                    AgreementParty(name=name, role=role, id=pid, email=email, phone=phone)
                )
        parties = _ensure_agreement_parties_have_ids(parties_raw)
        next_data["parties"] = [p.model_dump() for p in parties]
        for ap in parties:
            aid = (ap.id or "").strip()
            if aid and aid not in prior_ids:
                audit_log.append(
                    AuditEvent(
                        event_type="participant_added",
                        at=now,
                        field="parties",
                        value={
                            "participant_id": aid,
                            "name": ap.name,
                            "role": ap.role,
                        },
                    ).model_dump()
                )
    elif body.field == "payment_request":
        if body.value is None or body.value == "":
            next_data["payment_request"] = None
        elif isinstance(body.value, dict):
            next_data["payment_request"] = body.value
        else:
            raise HTTPException(status_code=400, detail="invalid_payment_request")
    elif body.field == "payment_required":
        next_data["payment_required"] = bool(body.value)
    elif body.field in ("feed_party_anonymize", "feed_show_financial_summary"):
        next_data[body.field] = bool(body.value)
    elif body.field == "feed_visibility":
        v = str(body.value or "private").strip().lower()
        if v not in ("private", "link_only", "public"):
            raise HTTPException(status_code=400, detail="invalid_feed_visibility")
        next_data[body.field] = v
    elif body.field == "feed_anchor_network":
        if body.value is None or body.value == "":
            next_data[body.field] = None
        else:
            net = str(body.value).strip()
            if net not in ALLOWED_AGREEMENT_ANCHOR_NETWORKS:
                raise HTTPException(status_code=400, detail="invalid_feed_anchor_network")
            next_data[body.field] = net
    else:
        if body.value is None:
            next_data[body.field] = None
        else:
            next_data[body.field] = str(body.value).strip()

    next_data["updated_at"] = now
    audit_log.append(
        AuditEvent(
            event_type="field_updated",
            at=now,
            field=body.field,
            value=next_data.get(body.field),
        ).model_dump()
    )
    next_data["audit_log"] = audit_log
    next_draft = AgreementDraft.model_validate(next_data)
    _save_draft_sync(next_draft.model_dump(), request)
    try:
        from backend.integrations.hooks_emit import claw_emit_integration_event_from_subject

        claw_emit_integration_event_from_subject(
            resolve_subject_from_request(request),
            "agreement.updated",
            "agreement",
            agreement_id,
            {"field": body.field},
        )
    except Exception:
        pass
    return {
        "id": agreement_id,
        "draft": next_draft.model_dump(),
        "canonical_json": canonicalize_agreement(next_draft.model_dump()),
        "economics": economics_overlay_for_agreement(agreement_id),
    }


@router.post("/{agreement_id}/render", response_model=AgreementRenderResponse)
def render_agreement(agreement_id: str, request: Request) -> AgreementRenderResponse:
    # Read-only preview: do not persist audit_log/updated_at (render is deterministic from draft + watermark).
    assert_agreement_full_draft_read_allowed(request, agreement_id)
    draft = _load_or_404(agreement_id)
    wm = _watermark_active_for_agreement(agreement_id)
    return AgreementRenderResponse(
        id=agreement_id,
        rendered_html=_render_html(draft, watermark=wm),
    )


@router.post("/{agreement_id}/recipient-preview-export-pdf")
def post_recipient_preview_export_pdf(
    agreement_id: str,
    request: Request,
    body: RecipientPreviewPdfExportBody,
) -> Response:
    """
    Turn preview HTML into a downloadable PDF (same Story pipeline as VS01 seed, recipient typography profile).
    Caller must pass the HTML currently shown for that export variant (original baseline, proposed, or redline).
    """
    aid = (agreement_id or "").strip()
    if not aid:
        raise HTTPException(status_code=400, detail="missing_agreement_id")
    assert_agreement_full_draft_read_allowed(request, aid)
    cap = assess_agreement_pdf_story_capability()
    if not cap.get("available"):
        log.warning(
            "[recipient-pdf-export] rejected agreement_id=%s reason=%s",
            aid,
            (cap.get("reason") or "")[:400],
        )
        raise HTTPException(
            status_code=503,
            detail={
                "code": "recipient_pdf_export_unavailable",
                "message": _RECIPIENT_PDF_EXPORT_UNAVAILABLE_USER,
            },
        )

    draft_pdf = _load_or_404(aid)
    party_names_pdf = [str(p.name or "").strip() for p in (draft_pdf.parties or []) if str(p.name or "").strip()]
    scan_plain = strip_html_agreement_scan_text(body.html or "")
    ok_ph_pdf, _, ph_diag_pdf = validate_user_visible_agreement_text(
        scan_plain,
        party_names=party_names_pdf,
        intake_raw="",
        surface="recipient_preview_export_pdf",
        agreement_family="",
    )
    if not ok_ph_pdf:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "agreement_placeholder_blocked",
                "message": "This export still contains drafting placeholders. Resolve them in the agreement before creating a PDF.",
                "placeholder": ph_diag_pdf,
            },
        )

    built = agreement_rendered_html_to_pdf_bytes(
        body.html,
        title="Agreement",
        story_css_profile="recipient",
    )
    if built.render_mode not in RECIPIENT_PREVIEW_PDF_STORY_RENDER_MODES:
        log.warning(
            "[recipient-pdf-export] rejected_non_story_render agreement_id=%s render_mode=%s",
            aid,
            built.render_mode,
        )
        raise HTTPException(
            status_code=503,
            detail={
                "code": "recipient_pdf_export_unavailable",
                "message": _RECIPIENT_PDF_EXPORT_UNAVAILABLE_USER,
            },
        )

    fn = _RECIPIENT_PREVIEW_PDF_FILENAMES.get(body.export_kind) or "lawdog-agreement.pdf"
    return Response(
        content=built.pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'},
    )


class CompletedSignedExportBody(BaseModel):
    """Optional HTML matching the owner signed view; server falls back to stored signed snapshot corpus."""

    html: str = Field(default="", max_length=1_200_000)

    @field_validator("html", mode="before")
    @classmethod
    def _trim_completed_signed_html(cls, v: object) -> str:
        return ("" if v is None else str(v)).strip()


def _signed_corpus_plain_to_export_html(corpus_plain: str) -> str:
    body = html.escape((corpus_plain or "").strip())
    return (
        "<article style='max-width:720px;margin:0 auto'>"
        "<pre style='white-space:pre-wrap;font-family:Georgia,serif;font-size:15px;line-height:1.65;"
        "color:#0f172a;margin:0;padding:0;border:0;background:transparent'>"
        f"{body}</pre></article>"
    )


def _completed_signed_pdf_filename(draft: AgreementDraft) -> str:
    title = (draft.title or "").strip() or "agreement"
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")[:80] or "agreement"
    return f"{slug}-signed.pdf"


def _completed_signed_export_pdf_response(
    *,
    agreement_id: str,
    draft: AgreementDraft,
    html_for_export: str,
) -> Response:
    cap = assess_agreement_pdf_story_capability()
    if not cap.get("available"):
        log.warning(
            "[completed-signed-pdf-export] rejected agreement_id=%s reason=%s",
            agreement_id,
            (cap.get("reason") or "")[:400],
        )
        raise HTTPException(
            status_code=503,
            detail={
                "code": "completed_signed_pdf_export_unavailable",
                "message": _RECIPIENT_PDF_EXPORT_UNAVAILABLE_USER,
            },
        )

    party_names_pdf = [str(p.name or "").strip() for p in (draft.parties or []) if str(p.name or "").strip()]
    scan_plain = strip_html_agreement_scan_text(html_for_export or "")
    ok_ph_pdf, _, ph_diag_pdf = validate_user_visible_agreement_text(
        scan_plain,
        party_names=party_names_pdf,
        intake_raw=_draft_placeholder_intake_corpus(draft),
        surface="completed_signed_export_pdf",
        agreement_family="",
    )
    if not ok_ph_pdf:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "agreement_placeholder_blocked",
                "message": "This export still contains drafting placeholders. Resolve them before creating a PDF.",
                "placeholder": ph_diag_pdf,
            },
        )

    built = agreement_rendered_html_to_pdf_bytes(
        html_for_export,
        title=(draft.title or "Agreement").strip() or "Agreement",
        story_css_profile="recipient",
    )
    if built.render_mode not in RECIPIENT_PREVIEW_PDF_STORY_RENDER_MODES:
        log.warning(
            "[completed-signed-pdf-export] rejected_non_story_render agreement_id=%s render_mode=%s",
            agreement_id,
            built.render_mode,
        )
        raise HTTPException(
            status_code=503,
            detail={
                "code": "completed_signed_pdf_export_unavailable",
                "message": _RECIPIENT_PDF_EXPORT_UNAVAILABLE_USER,
            },
        )

    fn = _completed_signed_pdf_filename(draft)
    return Response(
        content=built.pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{fn}"'},
    )


@router.post("/{agreement_id}/completed-signed-export-pdf")
def post_completed_signed_export_pdf(
    agreement_id: str,
    request: Request,
    body: CompletedSignedExportBody,
) -> Response:
    """Download PDF for a fully executed signed agreement (owner/read auth; signed snapshot or supplied HTML)."""
    aid = (agreement_id or "").strip()
    if not aid:
        raise HTTPException(status_code=400, detail="missing_agreement_id")
    assert_agreement_full_draft_read_allowed(request, aid)
    draft = _load_or_404(aid)
    if not _agreement_draft_fully_executed(draft):
        raise HTTPException(status_code=403, detail="agreement_not_fully_executed")

    html_for_export = (body.html or "").strip()
    if not html_for_export:
        from backend.services.vs01_signer_completion import read_fully_executed_snapshot_from_draft

        snap = read_fully_executed_snapshot_from_draft(draft.model_dump())
        corpus_plain = str((snap or {}).get("corpus_plain") or "").strip()
        if len(corpus_plain) < 80:
            raise HTTPException(status_code=409, detail="signed_snapshot_unavailable")
        html_for_export = _signed_corpus_plain_to_export_html(corpus_plain)

    return _completed_signed_export_pdf_response(
        agreement_id=aid,
        draft=draft,
        html_for_export=html_for_export,
    )


@router.get("/public/{agreement_id}/completed-signed-export-pdf")
def get_public_completed_signed_export_pdf(agreement_id: str) -> Response:
    """Public PDF download for fully executed agreements only (no auth; uses stored signed snapshot)."""
    if not public_agreement_verify_enabled():
        raise HTTPException(status_code=404, detail="not_found")
    aid = (agreement_id or "").strip()
    if not aid:
        raise HTTPException(status_code=404, detail="not_found")
    try:
        raw = load_draft(aid)
    except (KeyError, ValueError):
        raise HTTPException(status_code=404, detail="not_found")
    try:
        draft = AgreementDraft.model_validate(raw)
    except ValidationError:
        raise HTTPException(status_code=404, detail="not_found")
    if not _agreement_draft_fully_executed(draft):
        raise HTTPException(status_code=403, detail="agreement_not_fully_executed")

    from backend.services.vs01_signer_completion import read_fully_executed_snapshot_from_draft

    snap = read_fully_executed_snapshot_from_draft(draft.model_dump())
    corpus_plain = str((snap or {}).get("corpus_plain") or "").strip()
    if len(corpus_plain) < 80:
        raise HTTPException(status_code=409, detail="signed_snapshot_unavailable")
    html_for_export = _signed_corpus_plain_to_export_html(corpus_plain)
    return _completed_signed_export_pdf_response(
        agreement_id=aid,
        draft=draft,
        html_for_export=html_for_export,
    )


def _draft_placeholder_intake_corpus(draft: AgreementDraft) -> str:
    """Best-effort intake allowlist for placeholder validation (party names, emails, purpose)."""
    parts: List[str] = []
    for p in draft.parties or []:
        nm = str(p.name or "").strip()
        em = str(p.email or "").strip()
        if nm:
            parts.append(nm)
        if em:
            parts.append(em)
    for key in ("purpose", "payment_terms", "title", "jurisdiction"):
        seg = str(getattr(draft, key, None) or "").strip()
        if seg:
            parts.append(seg)
    return "\n".join(parts)


class Vs01SigningSeedBody(BaseModel):
    """Optional authoritative signing corpus for paid Pro VS01 seed (must exceed stored draft preview)."""

    signing_corpus_plain: Optional[str] = Field(default=None, max_length=1_200_000)


_VS01_SIGNING_CORPUS_OVERRIDE_MIN_LEN = 1500


def _vs01_signing_seed_error_detail(
    *,
    agreement_id: str,
    stage: str,
    code: str,
    message: str,
    exc: Optional[BaseException] = None,
    extra: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    out: Dict[str, Any] = {
        "code": code,
        "message": (message or "")[:2000],
        "agreement_id": agreement_id,
        "stage": stage,
    }
    if exc is not None:
        out["exc_type"] = type(exc).__name__
    if extra:
        for k, v in extra.items():
            if v is not None:
                out[k] = v
    return out


@router.post("/{agreement_id}/vs01-signing-seed")
def post_agreement_vs01_signing_seed(
    agreement_id: str,
    request: Request,
    body: Vs01SigningSeedBody = Vs01SigningSeedBody(),
) -> Dict[str, Any]:
    """
    Owner-only: render locked agreement HTML to PDF and finalize as a VS01 /v1/documents body
    (used by paid Pro sender-first → /app/esign/:documentId bridge).
    """
    aid = (agreement_id or "").strip()
    if not aid:
        log.info(
            "[agreement-vs01-seed] event=rejected agreement_id= stage=validate status=400 code=missing_agreement_id"
        )
        raise HTTPException(
            status_code=400,
            detail=_vs01_signing_seed_error_detail(
                agreement_id="",
                stage="validate",
                code="missing_agreement_id",
                message="agreement_id is required",
            ),
        )

    if not _agreements_write_allowed():
        log.info(
            "[agreement-vs01-seed] event=rejected agreement_id=%s stage=policy status=403 code=verifier_only",
            aid,
        )
        raise HTTPException(status_code=403, detail="verifier_only")

    # --- auth_guards ---
    try:
        _owner_mutation_guards(request, aid, surface="vs01_signing_seed")
    except HTTPException:
        raise
    except Exception as exc:
        log.exception(
            "[agreement-vs01-seed] event=failure agreement_id=%s stage=auth_guards status=503 code=auth_guards",
            aid,
        )
        raise HTTPException(
            status_code=503,
            detail=_vs01_signing_seed_error_detail(
                agreement_id=aid,
                stage="auth_guards",
                code="vs01_signing_seed_auth_failed",
                message=str(exc) or type(exc).__name__,
                exc=exc,
            ),
        ) from exc

    # --- load_agreement ---
    try:
        draft = _load_or_404(aid)
    except HTTPException:
        raise
    except Exception as exc:
        log.exception(
            "[agreement-vs01-seed] event=failure agreement_id=%s stage=load_agreement status=422 code=load_failed",
            aid,
        )
        raise HTTPException(
            status_code=422,
            detail=_vs01_signing_seed_error_detail(
                agreement_id=aid,
                stage="load_agreement",
                code="vs01_signing_seed_load_failed",
                message=str(exc) or type(exc).__name__,
                exc=exc,
            ),
        ) from exc

    signing_plain = (body.signing_corpus_plain or "").strip()
    if len(signing_plain) >= _VS01_SIGNING_CORPUS_OVERRIDE_MIN_LEN:
        field_key_override, corpus_before = primary_agreement_plain_field_and_value(draft)
        if len(signing_plain) > len(corpus_before or ""):
            merge_fields: Dict[str, str] = {field_key_override: signing_plain}
            for alt_key in (
                "premium_full_document_text",
                "server_full_document_text",
                "document_text",
            ):
                if alt_key != field_key_override:
                    merge_fields[alt_key] = signing_plain
            draft = _merge_agreement_draft(draft, **merge_fields)
            log.info(
                "[vs01-signing-seed-corpus-override] agreement_id=%s len=%s field=%s prev_len=%s",
                aid,
                len(signing_plain),
                field_key_override,
                len(corpus_before or ""),
            )

    # --- placeholder_template_safety (pre-render) ---
    party_names_vs = [str(p.name or "").strip() for p in (draft.parties or []) if str(p.name or "").strip()]
    field_key_vs, corpus_vs = primary_agreement_plain_field_and_value(draft)
    intake_corpus_vs = _draft_placeholder_intake_corpus(draft)
    ok_ph_vs, fixed_corpus_vs, ph_diag_vs = validate_user_visible_agreement_text(
        corpus_vs,
        party_names=party_names_vs,
        intake_raw=intake_corpus_vs,
        surface="vs01_signing_seed",
        agreement_family="",
    )
    if not ok_ph_vs:
        log.warning(
            "[agreement-vs01-seed] event=rejected agreement_id=%s stage=placeholder_safety status=422",
            aid,
        )
        raise HTTPException(
            status_code=422,
            detail=_vs01_signing_seed_error_detail(
                agreement_id=aid,
                stage="placeholder_safety",
                code="vs01_signing_seed_placeholder_blocked",
                message="Signing seed is blocked until drafting placeholders are resolved.",
                extra={"placeholder": ph_diag_vs},
            ),
        )
    if (corpus_vs or "").strip() and fixed_corpus_vs.strip() != corpus_vs.strip():
        draft = _merge_agreement_draft(draft, **{field_key_vs: fixed_corpus_vs})

    # --- economics_watermark (fail-open: seed must not depend on usage-economics DB uptime) ---
    try:
        wm = _watermark_active_for_agreement(aid)
    except Exception as exc:
        log.warning(
            "[agreement-vs01-seed] event=warning agreement_id=%s stage=economics_watermark status=200 "
            "code=economics_overlay_skipped exc_type=%s msg=%s — using watermark=false for VS01 seed",
            aid,
            type(exc).__name__,
            (str(exc) or "")[:400],
            exc_info=True,
        )
        wm = False

    # --- render_html ---
    try:
        html = _render_html(draft, watermark=wm)
    except HTTPException:
        raise
    except Exception as exc:
        log.exception(
            "[agreement-vs01-seed] event=failure agreement_id=%s stage=render_html status=503 code=render_failed",
            aid,
        )
        raise HTTPException(
            status_code=503,
            detail=_vs01_signing_seed_error_detail(
                agreement_id=aid,
                stage="render_html",
                code="vs01_signing_seed_render_failed",
                message=str(exc) or type(exc).__name__,
                exc=exc,
            ),
        ) from exc

    html_len = len(html or "")
    log.info(
        "[agreement-vs01-seed] event=progress agreement_id=%s stage=render_html status=200 html_len=%s watermark=%s",
        aid,
        html_len,
        bool(wm),
    )

    # --- html_to_pdf ---
    title = (draft.title or "").strip() or "Agreement"
    try:
        built = agreement_rendered_html_to_pdf_bytes(html, title=title)
    except Exception as exc:
        log.exception(
            "[agreement-vs01-seed] event=failure agreement_id=%s stage=html_to_pdf status=503 "
            "html_len=%s exc_type=%s",
            aid,
            html_len,
            type(exc).__name__,
        )
        raise HTTPException(
            status_code=503,
            detail=_vs01_signing_seed_error_detail(
                agreement_id=aid,
                stage="html_to_pdf",
                code="vs01_signing_seed_pdf_failed",
                message=str(exc) or type(exc).__name__,
                exc=exc,
            ),
        ) from exc

    pdf_len = len(built.pdf_bytes or b"")
    # --- finalize_document ---
    try:
        meta = document_service.finalize_document(
            built.pdf_bytes, content_type="application/pdf", agreement_id=aid
        )
    except Exception as exc:
        _seed_store_ctx = document_service.document_storage_seed_error_context()
        log.exception(
            "[agreement-vs01-seed] event=failure agreement_id=%s stage=finalize_document status=503 "
            "code=vs01_finalize_failed html_len=%s pdf_len=%s render_mode=%s exc_type=%s msg=%s "
            "documents_candidates=%s unified_artifact_store=%s",
            aid,
            html_len,
            pdf_len,
            getattr(built, "render_mode", None),
            type(exc).__name__,
            (str(exc) or "")[:500],
            _seed_store_ctx.get("documents_candidates"),
            _seed_store_ctx.get("unified_artifact_store_enabled"),
        )
        raise HTTPException(
            status_code=503,
            detail=_vs01_signing_seed_error_detail(
                agreement_id=aid,
                stage="finalize_document",
                code="vs01_finalize_failed",
                message=str(exc) or type(exc).__name__,
                exc=exc,
                extra=_seed_store_ctx,
            ),
        ) from exc

    # --- response_serialization (defensive) ---
    doc_id = meta.get("document_id") if isinstance(meta, dict) else None
    if not doc_id or not isinstance(doc_id, str) or not doc_id.strip():
        log.error(
            "[agreement-vs01-seed] event=failure agreement_id=%s stage=response_serialization status=503 "
            "html_len=%s pdf_len=%s reason=missing_document_id_in_meta",
            aid,
            html_len,
            pdf_len,
        )
        raise HTTPException(
            status_code=503,
            detail=_vs01_signing_seed_error_detail(
                agreement_id=aid,
                stage="response_serialization",
                code="vs01_finalize_incomplete",
                message="finalize_document returned no document_id",
            ),
        )

    hsh = meta.get("content_sha256") if isinstance(meta, dict) else None
    log.info(
        "[agreement-vs01-seed] event=success agreement_id=%s stage=complete status=200 html_len=%s pdf_len=%s "
        "render_mode=%s document_id=%s content_sha256=%s",
        aid,
        html_len,
        pdf_len,
        getattr(built, "render_mode", None),
        doc_id,
        hsh if isinstance(hsh, str) else None,
    )
    try:
        return {
            "ok": True,
            "document_id": doc_id,
            "content_sha256": hsh if isinstance(hsh, str) else None,
        }
    except Exception as exc:
        log.exception(
            "[agreement-vs01-seed] event=failure agreement_id=%s stage=response_serialization status=500",
            aid,
        )
        raise HTTPException(
            status_code=500,
            detail=_vs01_signing_seed_error_detail(
                agreement_id=aid,
                stage="response_serialization",
                code="vs01_signing_seed_response_failed",
                message=str(exc) or type(exc).__name__,
                exc=exc,
            ),
        ) from exc


PRO_REDLINE_SNAPSHOT_MAX = 256_000


def _agreement_id_short(aid: str) -> str:
    s = (aid or "").strip()
    return s[:12] if len(s) >= 12 else s


def _canonical_agreement_plain_from_raw(raw: Dict[str, Any]) -> str:
    best = ""
    for k in (
        "server_full_document_text",
        "premium_server_full_document_text",
        "premium_full_document_text",
        "document_text",
        "rendered_document_text",
    ):
        seg = raw.get(k)
        if not isinstance(seg, str):
            continue
        t = seg.strip()
        if len(t) > len(best):
            best = t
    if not best:
        p = str(raw.get("purpose") or "").strip()
        pay = str(raw.get("payment_terms") or "").strip()
        best = (p + "\n\n" + pay).strip()
    return best


def _truncate_redline_snapshot(s: str) -> Tuple[str, bool]:
    if len(s) <= PRO_REDLINE_SNAPSHOT_MAX:
        return s, False
    return s[:PRO_REDLINE_SNAPSHOT_MAX], True


def _default_pro_redline() -> Dict[str, Any]:
    return {"version_counter": 0, "version_events": [], "pending_import": None, "suggestions": []}


def _pro_redline_get(raw: Dict[str, Any]) -> Dict[str, Any]:
    pr = raw.get("pro_redline_v1")
    if isinstance(pr, dict):
        out = _default_pro_redline()
        out.update(pr)
        if not isinstance(out.get("version_events"), list):
            out["version_events"] = []
        if not isinstance(out.get("suggestions"), list):
            out["suggestions"] = []
        try:
            out["version_counter"] = int(out.get("version_counter") or 0)
        except Exception:
            out["version_counter"] = 0
        return out
    return _default_pro_redline()


def _pro_redline_set(raw: Dict[str, Any], pr: Dict[str, Any]) -> None:
    raw["pro_redline_v1"] = pr


def _build_docx_bytes_from_plain(text: str) -> bytes:
    import docx

    doc = docx.Document()
    for line in (text or "").replace("\r\n", "\n").split("\n"):
        doc.add_paragraph(line or " ")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


class ProRedlineImportTextBody(BaseModel):
    imported_text: str = Field(..., min_length=1, max_length=920_000)


class ProRedlineReviewerSuggestionBody(BaseModel):
    participant_id: str = Field(..., min_length=1)
    suggestion_text: str = Field(..., min_length=1, max_length=48_000)
    reviewer_display_name: str = ""
    reviewer_email: str = ""


@router.get("/{agreement_id}/export-draft.txt")
def export_draft_txt(agreement_id: str, request: Request) -> Response:
    assert_agreement_full_draft_read_allowed(request, agreement_id)
    raw = load_draft(agreement_id)
    text = _canonical_agreement_plain_from_raw(raw)
    if not (text or "").strip():
        raise HTTPException(status_code=400, detail="empty_document_export")
    body = text.encode("utf-8")
    short = _agreement_id_short(agreement_id)
    return Response(
        content=body,
        media_type="text/plain; charset=utf-8",
        headers={"Content-Disposition": f'attachment; filename="lawdog-agreement-{short}.txt"'},
    )


@router.get("/{agreement_id}/export-draft.docx")
def export_draft_docx(agreement_id: str, request: Request) -> Response:
    assert_agreement_full_draft_read_allowed(request, agreement_id)
    raw = load_draft(agreement_id)
    text = _canonical_agreement_plain_from_raw(raw)
    if not (text or "").strip():
        raise HTTPException(status_code=400, detail="empty_document_export")
    blob = _build_docx_bytes_from_plain(text)
    short = _agreement_id_short(agreement_id)
    return Response(
        content=blob,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="lawdog-agreement-{short}.docx"'},
    )


@router.post("/{agreement_id}/pro-redline/import-text")
def pro_redline_import_text(agreement_id: str, body: ProRedlineImportTextBody, request: Request) -> Dict[str, Any]:
    _owner_mutation_guards(request, agreement_id, surface="pro_redline_import_text")
    raw = load_draft(agreement_id)
    base = _canonical_agreement_plain_from_raw(raw)
    imported = (body.imported_text or "").strip()
    if not imported:
        raise HTTPException(status_code=400, detail="imported_text_required")
    blocks, changed = compute_pro_redline_block_diff(base, imported)
    diff_summary = {"blocks": blocks, "changed_block_count": changed}
    pending_id = str(uuid.uuid4())
    now = _utc_now_iso()
    base_snap, base_trunc = _truncate_redline_snapshot(base)
    imp_snap, imp_trunc = _truncate_redline_snapshot(imported)
    pr = _pro_redline_get(raw)
    pr["pending_import"] = {
        "id": pending_id,
        "created_at": now,
        "base_len": len(base),
        "imported_len": len(imported),
        "imported_text": imported,
        "diff_summary_json": diff_summary,
        "base_snapshot_truncated": base_trunc,
        "imported_snapshot_truncated": imp_trunc,
        "base_document_text": base_snap,
        "imported_document_text": imp_snap,
    }
    pr.setdefault("version_events", []).append(
        {
            "version_number": int(pr.get("version_counter") or 0),
            "source": "imported_revision",
            "actor_label": "Owner",
            "created_at": now,
            "pending_revision_id": pending_id,
            "diff_summary_json": diff_summary,
        }
    )
    _pro_redline_set(raw, pr)
    raw["updated_at"] = now
    _save_draft_sync(AgreementDraft.model_validate(raw).model_dump(), request)
    log.info(
        "[pro-redline-import] agreementIdShort=%s baseLen=%s importedLen=%s changedBlockCount=%s",
        _agreement_id_short(agreement_id),
        len(base),
        len(imported),
        changed,
    )
    return {
        "ok": True,
        "pending_id": pending_id,
        "changed_block_count": changed,
        "no_changes": changed == 0,
        "diff_summary": diff_summary,
    }


@router.post("/{agreement_id}/pro-redline/import-file")
async def pro_redline_import_file(agreement_id: str, request: Request, file: UploadFile = File(...)) -> Dict[str, Any]:
    _owner_mutation_guards(request, agreement_id, surface="pro_redline_import_file")
    raw_bytes = await file.read()
    fname = (file.filename or "").strip().lower() or "upload.txt"
    if not fname.endswith(".txt"):
        raise HTTPException(status_code=400, detail="pro_redline_import_txt_only_v1")
    try:
        extracted = (raw_bytes or b"").decode("utf-8").strip()
    except UnicodeDecodeError:
        extracted = (raw_bytes or b"").decode("utf-8", errors="replace").strip()
    if not extracted:
        raise HTTPException(status_code=400, detail="imported_text_empty")
    body = ProRedlineImportTextBody(imported_text=extracted)
    return pro_redline_import_text(agreement_id, body, request)


@router.post("/{agreement_id}/pro-redline/accept-import")
def pro_redline_accept_import(agreement_id: str, request: Request) -> Dict[str, Any]:
    _owner_mutation_guards(request, agreement_id, surface="pro_redline_accept_import")
    _assert_negotiation_not_locked(agreement_id)
    raw = load_draft(agreement_id)
    pr = _pro_redline_get(raw)
    pending = pr.get("pending_import")
    if not isinstance(pending, dict):
        raise HTTPException(status_code=400, detail="no_pending_import")
    pid = str(pending.get("id") or "").strip()
    if not pid:
        raise HTTPException(status_code=400, detail="no_pending_import")
    imported = str(pending.get("imported_text") or "").strip()
    if not imported:
        raise HTTPException(status_code=400, detail="invalid_pending_import")
    party_names_imp: List[str] = []
    for p in raw.get("parties") or []:
        if isinstance(p, dict):
            nm = str(p.get("name") or "").strip()
        else:
            nm = str(getattr(p, "name", "") or "").strip()
        if nm:
            party_names_imp.append(nm)
    ok_imp, fixed_imp, diag_imp = validate_user_visible_agreement_text(
        imported,
        party_names=party_names_imp,
        intake_raw="",
        surface="pro_redline_accept_import",
        agreement_family=str(raw.get("title") or "")[:120],
    )
    if not ok_imp:
        raise HTTPException(
            status_code=422,
            detail={
                "code": "agreement_placeholder_blocked",
                "message": "Imported text still contains drafting placeholders. Fix them before accepting the import.",
                "placeholder": diag_imp,
            },
        )
    imported = fixed_imp.strip()
    base_before = _canonical_agreement_plain_from_raw(raw)
    now = _utc_now_iso()
    from_v = int(pr.get("version_counter") or 0)
    pr["version_counter"] = from_v + 1
    base_snap, base_trunc = _truncate_redline_snapshot(base_before)
    imp_snap, imp_trunc = _truncate_redline_snapshot(imported)
    acc_snap, acc_trunc = _truncate_redline_snapshot(imported)
    ev: Dict[str, Any] = {
        "version_number": pr["version_counter"],
        "source": "owner_accepted_revision",
        "actor_label": "Owner",
        "actor_email": None,
        "created_at": now,
        "base_document_text": base_snap,
        "imported_document_text": imp_snap,
        "accepted_document_text": acc_snap,
        "base_snapshot_truncated": base_trunc,
        "imported_snapshot_truncated": imp_trunc,
        "accepted_snapshot_truncated": acc_trunc,
        "diff_summary_json": pending.get("diff_summary_json"),
        "pending_revision_id": pid,
    }
    pr.setdefault("version_events", []).append(ev)
    pr["pending_import"] = None
    _pro_redline_set(raw, pr)
    raw["server_full_document_text"] = imported
    raw["premium_server_full_document_text"] = imported
    raw["premium_full_document_text"] = imported
    raw["document_text"] = imported
    # Drop stale rendered/plain cache so export + canonical pick the accepted corpus (longest-field merge).
    raw["rendered_document_text"] = None
    raw["updated_at"] = now
    audit = list(raw.get("audit_log") or [])
    audit.append(
        AuditEvent(
            event_type="pro_redline_import_accepted",
            at=now,
            field="pro_redline_v1",
            value={"pending_id": pid, "version_number": pr["version_counter"]},
        ).model_dump()
    )
    raw["audit_log"] = audit
    _save_draft_sync(AgreementDraft.model_validate(raw).model_dump(), request)
    log.info(
        "[pro-redline-accept] agreementIdShort=%s fromVersion=%s toVersion=%s",
        _agreement_id_short(agreement_id),
        from_v,
        pr["version_counter"],
    )
    return {"ok": True, "version_number": pr["version_counter"], "draft": AgreementDraft.model_validate(raw).model_dump()}


@router.post("/{agreement_id}/pro-redline/reject-import")
def pro_redline_reject_import(agreement_id: str, request: Request) -> Dict[str, Any]:
    _owner_mutation_guards(request, agreement_id, surface="pro_redline_reject_import")
    raw = load_draft(agreement_id)
    pr = _pro_redline_get(raw)
    pending = pr.get("pending_import")
    pid = str((pending or {}).get("id") or "").strip() if isinstance(pending, dict) else ""
    now = _utc_now_iso()
    pr["pending_import"] = None
    pr.setdefault("version_events", []).append(
        {
            "version_number": int(pr.get("version_counter") or 0),
            "source": "owner_rejected_revision",
            "actor_label": "Owner",
            "created_at": now,
            "pending_revision_id": pid or None,
            "rejection_kind": "import",
        }
    )
    _pro_redline_set(raw, pr)
    raw["updated_at"] = now
    audit = list(raw.get("audit_log") or [])
    audit.append(
        AuditEvent(
            event_type="pro_redline_import_rejected",
            at=now,
            field="pro_redline_v1",
            value={"pending_revision_id": pid or None},
        ).model_dump()
    )
    raw["audit_log"] = audit
    _save_draft_sync(AgreementDraft.model_validate(raw).model_dump(), request)
    log.info(
        "[pro-redline-reject] agreementIdShort=%s pendingRevisionId=%s",
        _agreement_id_short(agreement_id),
        pid or "-",
    )
    return {"ok": True, "draft": AgreementDraft.model_validate(raw).model_dump()}


@router.post("/{agreement_id}/pro-redline/reviewer-suggestion")
def pro_redline_reviewer_suggestion(
    agreement_id: str, body: ProRedlineReviewerSuggestionBody, request: Request
) -> Dict[str, Any]:
    assert_free_incomplete_draft_not_expired(agreement_id, surface="pro_redline_reviewer_suggestion")
    draft = _load_or_404(agreement_id)
    assert_agreement_recipient_write_allowed(
        request,
        agreement_id,
        allowed_modes=("review",),
        bind_participant_id=body.participant_id,
    )
    lock = read_signing_lock(agreement_id)
    if lock and bool((lock or {}).get("locked_version_id")):
        raise HTTPException(status_code=400, detail="negotiation_locked")
    draft = _persist_party_id_backfill(draft)
    _validate_nonowner_proposer(draft, body.participant_id)
    raw = load_draft(agreement_id)
    pr = _pro_redline_get(raw)
    sid = str(uuid.uuid4())
    now = _utc_now_iso()
    row = {
        "id": sid,
        "created_at": now,
        "participant_id": (body.participant_id or "").strip(),
        "reviewer_label": (body.reviewer_display_name or "").strip(),
        "reviewer_email": (body.reviewer_email or "").strip(),
        "suggestion_text": (body.suggestion_text or "").strip(),
        "status": "pending",
    }
    pr.setdefault("suggestions", []).append(row)
    sug_preview = row["suggestion_text"][:4000] if len(row["suggestion_text"]) > 4000 else row["suggestion_text"]
    pr.setdefault("version_events", []).append(
        {
            "version_number": int(pr.get("version_counter") or 0),
            "source": "reviewer_suggestion",
            "actor_label": row["reviewer_label"] or "Reviewer",
            "actor_email": row["reviewer_email"] or None,
            "created_at": now,
            "suggestion_id": sid,
            "suggestion_text": sug_preview,
        }
    )
    _pro_redline_set(raw, pr)
    raw["updated_at"] = now
    audit = list(raw.get("audit_log") or [])
    audit.append(
        AuditEvent(
            event_type="pro_redline_reviewer_suggestion",
            at=now,
            field="pro_redline_v1",
            value={"suggestion_id": sid, "participant_id": row["participant_id"]},
        ).model_dump()
    )
    raw["audit_log"] = audit
    _save_draft_sync(AgreementDraft.model_validate(raw).model_dump(), request)
    log.info(
        "[pro-redline-suggestion-submit] agreementIdShort=%s actorType=%s suggestionLen=%s",
        _agreement_id_short(agreement_id),
        "reviewer",
        len(row["suggestion_text"]),
    )
    return {"ok": True, "suggestion_id": sid}


@router.post("/{agreement_id}/pro-redline/suggestions/{suggestion_id}/reject")
def pro_redline_suggestion_reject(agreement_id: str, suggestion_id: str, request: Request) -> Dict[str, Any]:
    _owner_mutation_guards(request, agreement_id, surface="pro_redline_suggestion_reject")
    raw = load_draft(agreement_id)
    pr = _pro_redline_get(raw)
    found = False
    hit_sug: Optional[str] = None
    hit_label: Optional[str] = None
    for s in pr.get("suggestions") or []:
        if not isinstance(s, dict):
            continue
        if str(s.get("id") or "").strip() != (suggestion_id or "").strip():
            continue
        if str(s.get("status") or "") != "pending":
            raise HTTPException(status_code=400, detail="suggestion_not_pending")
        hit_sug = str(s.get("suggestion_text") or "")
        hit_label = (str(s.get("reviewer_label") or "").strip() or None)
        s["status"] = "rejected"
        found = True
        break
    if not found:
        raise HTTPException(status_code=404, detail="suggestion_not_found")
    now = _utc_now_iso()
    sug_snap, _sug_trunc = _truncate_redline_snapshot(hit_sug or "")
    pr.setdefault("version_events", []).append(
        {
            "version_number": int(pr.get("version_counter") or 0),
            "source": "owner_rejected_revision",
            "actor_label": "Owner",
            "created_at": now,
            "suggestion_text": sug_snap or None,
            "suggestion_id": suggestion_id,
            "rejection_kind": "suggestion",
            "reviewer_label": hit_label,
        }
    )
    _pro_redline_set(raw, pr)
    raw["updated_at"] = now
    _save_draft_sync(AgreementDraft.model_validate(raw).model_dump(), request)
    return {"ok": True, "draft": AgreementDraft.model_validate(raw).model_dump()}


class ProRedlineMarkAppliedBody(BaseModel):
    """When `applied_document_text` is non-empty, persist as the new authoritative Pro corpus (post refine)."""

    applied_document_text: str = Field(default="", max_length=920_000)


@router.post("/{agreement_id}/pro-redline/suggestions/{suggestion_id}/mark-applied")
def pro_redline_suggestion_mark_applied(
    agreement_id: str,
    suggestion_id: str,
    request: Request,
    body: ProRedlineMarkAppliedBody = ProRedlineMarkAppliedBody(),
) -> Dict[str, Any]:
    """Owner-only: mark a reviewer suggestion row applied after they merged it (e.g. via LawDog Pro refine)."""
    _owner_mutation_guards(request, agreement_id, surface="pro_redline_suggestion_mark_applied")
    raw = load_draft(agreement_id)
    pr = _pro_redline_get(raw)
    hit: Optional[Dict[str, Any]] = None
    for s in pr.get("suggestions") or []:
        if isinstance(s, dict) and str(s.get("id") or "").strip() == (suggestion_id or "").strip():
            hit = s
            break
    if not hit:
        raise HTTPException(status_code=404, detail="suggestion_not_found")
    if str(hit.get("status") or "") != "pending":
        raise HTTPException(status_code=400, detail="suggestion_not_pending")
    now = _utc_now_iso()
    hit["status"] = "applied"
    pr["version_counter"] = int(pr.get("version_counter") or 0) + 1
    applied_doc = (body.applied_document_text or "").strip()
    if applied_doc:
        party_names_pr: List[str] = []
        for p in raw.get("parties") or []:
            if isinstance(p, dict):
                nm = str(p.get("name") or "").strip()
            else:
                nm = str(getattr(p, "name", "") or "").strip()
            if nm:
                party_names_pr.append(nm)
        ok_ph, fixed_doc, ph_diag = validate_user_visible_agreement_text(
            applied_doc,
            party_names=party_names_pr,
            intake_raw="",
            surface="pro_redline_mark_applied",
            agreement_family=str(raw.get("title") or "")[:120],
        )
        if not ok_ph:
            raise HTTPException(
                status_code=422,
                detail={
                    "code": "agreement_placeholder_blocked",
                    "message": "Applied text still contains drafting placeholders. Resolve them before saving.",
                    "placeholder": ph_diag,
                },
            )
        applied_doc = fixed_doc.strip()
        raw["server_full_document_text"] = applied_doc
        raw["premium_server_full_document_text"] = applied_doc
        raw["premium_full_document_text"] = applied_doc
        raw["document_text"] = applied_doc
        raw["rendered_document_text"] = None
    pr.setdefault("version_events", []).append(
        {
            "version_number": pr["version_counter"],
            "source": "owner_accepted_revision",
            "actor_label": "Owner",
            "created_at": now,
            "suggestion_text": str(hit.get("suggestion_text") or "")[:PRO_REDLINE_SNAPSHOT_MAX],
            "suggestion_id": suggestion_id,
            "note": "reviewer_suggestion_marked_applied",
        }
    )
    _pro_redline_set(raw, pr)
    raw["updated_at"] = now
    _save_draft_sync(AgreementDraft.model_validate(raw).model_dump(), request)
    return {"ok": True, "draft": AgreementDraft.model_validate(raw).model_dump()}


@router.post("/{agreement_id}/export-docx")
def export_agreement_docx(agreement_id: str, request: Request) -> Dict[str, Any]:
    _owner_mutation_guards(request, agreement_id, surface="export_docx")
    draft = _load_or_404(agreement_id)
    _assert_negotiation_not_locked(agreement_id)
    next_data = draft.model_dump()
    now = _utc_now_iso()
    next_data["updated_at"] = now
    audit_log = list(next_data.get("audit_log") or [])
    audit_log.append(AuditEvent(event_type="export_docx", at=now).model_dump())
    next_data["audit_log"] = audit_log
    next_draft = AgreementDraft.model_validate(next_data)
    _save_draft_sync(next_draft.model_dump(), request)
    wm = _watermark_active_for_agreement(agreement_id)
    return {
        "id": agreement_id,
        "status": "stub",
        "message": "DOCX export pipeline not yet enabled in this build.",
        "download_path": None,
        "watermark_required": wm,
        "watermark_document_note": WATERMARK_LABEL if wm else None,
    }


def _negotiate_assist_tag_options_posture(
    opts: List[Dict[str, str]], posture: NegotiationPosture, guidance: str
) -> None:
    prefix = f"{guidance} — Apply this strategic lens to the following edit instruction. "
    for o in opts:
        o["posture"] = posture
        raw_inst = o.get("instruction") or ""
        o["instruction"] = (prefix + str(raw_inst)).strip()


def _negotiate_assist_fallback(
    recipient_instruction: str,
    prior: Optional[Dict[str, Any]],
    curr: Optional[Dict[str, Any]],
    posture: NegotiationPosture = "cooperative",
) -> Dict[str, Any]:
    instr = (recipient_instruction or "").strip() or "unspecified changes"
    pay_before = ""
    pay_after = ""
    dur_before = ""
    dur_after = ""
    if isinstance(prior, dict):
        pay_before = str(prior.get("payment_terms") or "").strip()
        dur_before = str(prior.get("duration") or "").strip()
    if isinstance(curr, dict):
        pay_after = str(curr.get("payment_terms") or "").strip()
        dur_after = str(curr.get("duration") or "").strip()
    parts: List[str] = []
    if pay_before != pay_after and pay_after:
        parts.append("payment terms appear to differ from your prior draft")
    if dur_before != dur_after and dur_after:
        parts.append("term or duration may have changed")
    impact = "; ".join(parts) if parts else "review payment, term, and scope against your last version"
    what = (
        f"The other party asked for an adjustment. Requested change (summary): {instr[:400]}. "
        f"Impact (high level): {impact}. This is not legal advice—review the draft yourself."
    )
    if posture == "cooperative":
        opt_rows: List[Tuple[str, str, str]] = [
            ("opt_accept", "Accept with relationship intact", "Meet their ask; keep tone collaborative."),
            (
                "opt_counter",
                "Counter with shared compromise",
                "Split the difference on economics or timing while preserving core scope.",
            ),
            (
                "opt_scope",
                "Trade scope for speed",
                "Offer a narrower deliverable set in exchange for favorable terms.",
            ),
        ]
    elif posture == "firm":
        opt_rows = [
            ("opt_accept", "Accept only if terms are explicit", "Codify their request narrowly; no silent extras."),
            ("opt_counter", "Hold the line on economics", "Reject overreach; restate your prior commercial position."),
            (
                "opt_scope",
                "Reject vague expansion",
                "Decline broad or open-ended obligations; keep original boundaries.",
            ),
        ]
    elif posture == "protective":
        opt_rows = [
            (
                "opt_accept",
                "Accept with safeguards",
                "If you meet their ask, tighten termination, payment, and liability caps.",
            ),
            (
                "opt_counter",
                "Protect payment and narrow scope",
                "Keep stronger payment protections; reduce scope to offset their request.",
            ),
            (
                "opt_reject_partial",
                "Decline high-risk asks",
                "Push back on uncapped liability, auto-renew without exit, or fuzzy deliverables.",
            ),
        ]
    elif posture == "fast_close":
        opt_rows = [
            ("opt_accept", "Accept to sign", "Take their change if it’s low-risk and unblocks signature."),
            ("opt_counter", "Simple middle ground", "One clean compromise; avoid new layers."),
            ("opt_minimal", "Minimal wording tweak", "Smallest edit that gets to yes this week."),
        ]
    elif posture == "founder_friendly":
        opt_rows = [
            ("opt_accept", "Accept cash-friendly terms", "If you concede, preserve runway and cap exposure."),
            ("opt_counter", "Protect IP and control", "Hold founder IP and operating flexibility; trade minor economics."),
            ("opt_scope", "Cap obligations", "Limit personal guarantees, broad non-competes, and open-ended duties."),
        ]
    else:  # investor_friendly
        opt_rows = [
            (
                "opt_accept",
                "Accept with reporting hooks",
                "If aligned, add light reporting and milestone clarity.",
            ),
            (
                "opt_counter",
                "Structured compromise",
                "Formalize economics with schedules and clear remedies.",
            ),
            (
                "opt_protect",
                "Strengthen auditability",
                "Add notice, inspection, and default mechanics without exploding length.",
            ),
        ]

    opts: List[Dict[str, str]] = []
    for idx, (oid, label, summary) in enumerate(opt_rows):
        if idx == 0:
            instruction = f"Accept the counterparty's proposed revision and apply it fully as stated: {instr}"
        elif idx == 1:
            instruction = (
                "Counter the last proposed change: propose a balanced revision addressing their request "
                f"while protecting your side’s stated priorities. They asked: {instr[:280]}"
            )
        else:
            instruction = (
                "Counter the last proposed change: adjust term, scope, or risk allocation to match this strategy; "
                f"context — they asked: {instr[:280]}"
            )
        opts.append(
            {
                "id": oid,
                "label": label,
                "summary": summary,
                "instruction": instruction,
            }
        )

    guidance = NEGOTIATION_POSTURE_GUIDANCE.get(posture, NEGOTIATION_POSTURE_GUIDANCE["cooperative"])
    _negotiate_assist_tag_options_posture(opts, posture, guidance)
    risk = _negotiate_risk_heuristic(instr)
    return {"what_changed": what, "options": opts, "risk_assessment": risk}


def _negotiate_assist_llm(
    body: NegotiateAssistRequest,
    draft: AgreementDraft,
    trace_context: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    prior = body.prior_snapshot or {}
    curr = body.current_snapshot or draft.model_dump()
    ins = (body.recipient_instruction or "").strip()
    posture = body.negotiation_posture
    lens = NEGOTIATION_POSTURE_GUIDANCE.get(posture, NEGOTIATION_POSTURE_GUIDANCE["cooperative"])
    heur_risk = _negotiate_risk_heuristic(ins)
    system = (
        "You are a structured negotiation assistant for CLAW.\n"
        "You do NOT give legal advice. You help summarize proposed edits, triage risk (informational only), "
        "and suggest response strategies.\n"
        f"Negotiation posture (strategic lens only): {posture}.\n"
        f"Lens meaning: {lens}\n"
        "When `mode` includes summary: keep `what_changed` factual in 2-4 short sentences; match the posture "
        "lightly in tone only (no legal conclusions).\n"
        "When `mode` includes options: all three options MUST reflect this posture; also reflect `risk_assessment.tier` "
        "in tone: low_risk → simpler/faster paths; economic_impact → comparisons and tradeoffs; "
        "manual_legal_review → caution, narrowing, or deferral (still compact).\n"
        "Always include `risk_assessment` with: tier (low_risk | economic_impact | manual_legal_review), "
        "label (badge: Low risk | Economic impact | Manual legal review), "
        "explanation (one short sentence), rationale (1-2 lines max, plain English), "
        "helper_text (one line), confidence (low|medium|high). "
        "Triage is decision support only. Do not state outcomes are illegal/unenforceable.\n"
        "Return ONLY valid JSON with this shape:\n"
        '{ "what_changed": "string or empty if mode is options-only", '
        '"risk_assessment": { "tier": "", "label": "", "explanation": "", "rationale": "", '
        '"helper_text": "", "confidence": "" }, '
        '"options": [ { "id": "string", "label": "short title", '
        '"summary": "one line", "instruction": "plain-language edit instruction for an agreement editor", '
        '"posture": "<same as negotiation_posture in request>" } ] }\n'
        "Rules:\n"
        "- Each option must include `posture` matching the request negotiation_posture.\n"
        "- If mode is summary: options may be an empty array [].\n"
        "- If mode is options: what_changed may be an empty string.\n"
        "- Provide exactly 3 options when mode is options or both.\n"
        "- Each `instruction` must be self-contained for the agreement editor and **surgical**: minimal edits, "
        "preserve unchanged language, prefer short insertions over replacing whole clauses.\n"
        "- Do not claim the other party accepted anything.\n"
        "- No markdown, no prose outside JSON.\n"
    )
    payload: Dict[str, Any] = {
        "recipient_request": ins,
        "prior_draft_fields": prior,
        "current_draft_fields": curr,
        "mode": body.mode,
        "negotiation_posture": posture,
        "risk_triage_rules_hint": {
            "tier": heur_risk["tier"],
            "rationale": heur_risk["rationale"],
        },
    }
    try:
        llm_model = resolve_llm_model_for_access_class(body.ai_model_class)
        text = call_legal_llm(
            messages=[
                {"role": "system", "content": system},
                {"role": "user", "content": json.dumps(payload, ensure_ascii=False)},
            ],
            model=llm_model,
            max_tokens=768,
            temperature=0.2,
            trace_context=trace_context,
            airlock_profile="agreement_outbound",
            airlock_log_context="negotiation_risk_triage",
        )
        parsed = _extract_json_object(text)
        out: Dict[str, Any] = {"options": []}
        if body.mode in ("summary", "both"):
            wc = parsed.get("what_changed")
            out["what_changed"] = str(wc).strip() if wc else ""
        raw_ra = parsed.get("risk_assessment")
        out["risk_assessment"] = _merge_risk_assessment(
            heur_risk, raw_ra if isinstance(raw_ra, dict) else None
        )
        if body.mode in ("options", "both"):
            raw_opts = parsed.get("options")
            opts: List[Dict[str, str]] = []
            if isinstance(raw_opts, list):
                for i, o in enumerate(raw_opts[:3]):
                    if not isinstance(o, dict):
                        continue
                    opts.append(
                        {
                            "id": str(o.get("id") or f"opt_{i}"),
                            "label": str(o.get("label") or f"Option {i + 1}"),
                            "summary": str(o.get("summary") or "").strip(),
                            "instruction": str(o.get("instruction") or "").strip(),
                            "posture": str(o.get("posture") or posture),
                        }
                    )
            guidance_line = NEGOTIATION_POSTURE_GUIDANCE.get(
                posture, NEGOTIATION_POSTURE_GUIDANCE["cooperative"]
            )
            for o in opts:
                if o.get("posture") != posture:
                    o["posture"] = posture
                inst = str(o.get("instruction") or "").strip()
                if inst and guidance_line[: min(24, len(guidance_line))] not in inst[:160]:
                    o["instruction"] = (f"{guidance_line} — Apply this lens. {inst}").strip()
            out["options"] = opts
        if body.mode == "summary" and "what_changed" not in out:
            out["what_changed"] = ""
        if body.mode == "options" and not out["options"]:
            raise ValueError("missing_options")
        if body.mode == "both":
            if not out.get("what_changed"):
                raise ValueError("missing_summary")
            if len(out["options"]) < 3:
                raise ValueError("need_three_options")
        if "risk_assessment" not in out or not out["risk_assessment"]:
            out["risk_assessment"] = dict(heur_risk)
        return out
    except Exception:
        return _negotiate_assist_fallback(
            ins,
            prior if isinstance(prior, dict) else None,
            curr if isinstance(curr, dict) else None,
            body.negotiation_posture,
        )


@router.post("/{agreement_id}/negotiate-assist")
def negotiate_assist(
    agreement_id: str,
    body: NegotiateAssistRequest,
    request: Request,
    response: Response,
):
    """
    Read-only assist: what_changed text + 3 strategy options.
    Applying edits stays on POST /revise (same drafting pipeline).
    """
    session_type = body.session_type
    if session_type == "owner":
        require_claw_org_id_header(request)
        assert_registered_owner_matches(request, agreement_id)
    elif session_type == "recipient":
        assert_agreement_recipient_write_allowed(request, agreement_id, allowed_modes=("review",))
    draft = _load_or_404(agreement_id)
    ins = (body.recipient_instruction or "").strip()
    ok_txt, msg_txt = validate_negotiate_text(ins, session_type)
    if not ok_txt:
        return JSONResponse(
            status_code=400,
            content={"error": "input_too_large", "message": msg_txt},
        )
    if session_type == "recipient" and not recipient_prompt_allowed(ins):
        return JSONResponse(
            status_code=400,
            content={
                "error": "invalid_request",
                "message": "This request could not be processed for this agreement.",
            },
        )

    curr = body.current_snapshot or draft.model_dump()
    prior = body.prior_snapshot or {}
    payload_chars = len(
        json.dumps(
            {"recipient_instruction": ins, "prior_snapshot": prior, "current_snapshot": curr},
            ensure_ascii=False,
        )
    )
    ok_pay, msg_pay = validate_negotiate_payload_size(payload_chars, session_type)
    if not ok_pay:
        return JSONResponse(
            status_code=400,
            content={"error": "input_too_large", "message": msg_pay},
        )

    trace = build_llm_trace_context(
        session_type=session_type,
        agreement_id=agreement_id,
        request=request,
    )
    fallback = _negotiate_assist_fallback(
        body.recipient_instruction,
        body.prior_snapshot,
        curr,
        body.negotiation_posture,
    )
    if not OPENAI_API_KEY:
        rem = (
            peek_recipient_remaining(agreement_id, client_fingerprint(request))
            if session_type == "recipient"
            else 9999
        )
        for hk, hv in usage_response_header(rem if session_type == "recipient" else 9999).items():
            response.headers[hk] = hv
        return fallback

    remaining_after = 9999
    if session_type == "recipient":
        acquired, remaining_after = recipient_try_acquire_llm_slot(
            agreement_id, client_fingerprint(request)
        )
        if not acquired:
            return JSONResponse(
                status_code=429,
                content={
                    "error": "rate_limited",
                    "message": "AI usage limit reached for this session",
                },
                headers=usage_response_header(0),
            )
    try:
        out = _negotiate_assist_llm(body, draft, trace_context=trace)
        for hk, hv in usage_response_header(
            remaining_after if session_type == "recipient" else 9999
        ).items():
            response.headers[hk] = hv
        return out
    except Exception:
        for hk, hv in usage_response_header(
            remaining_after if session_type == "recipient" else 9999
        ).items():
            response.headers[hk] = hv
        return fallback


@router.post("/{agreement_id}/revise")
def revise_agreement(
    agreement_id: str,
    body: AgreementReviseRequest,
    request: Request,
    response: Response,
):
    if body.session_type == "owner":
        require_claw_org_id_header(request)
        assert_registered_owner_matches(request, agreement_id)
    elif body.session_type == "recipient":
        assert_agreement_recipient_write_allowed(request, agreement_id, allowed_modes=("review",))
    if body.persist:
        assert_free_incomplete_draft_not_expired(agreement_id, surface="revise")
        _assert_negotiation_not_locked(agreement_id)
    draft = _load_or_404(agreement_id)
    instruction = (body.instruction or "").strip()
    if not instruction:
        raise HTTPException(status_code=400, detail="instruction_required")
    session_type = body.session_type
    ok_ins, msg_ins = validate_instruction_size(instruction, session_type)
    if not ok_ins:
        return JSONResponse(
            status_code=400,
            content={"error": "input_too_large", "message": msg_ins},
        )
    if session_type == "recipient" and not recipient_prompt_allowed(instruction):
        return JSONResponse(
            status_code=400,
            content={
                "error": "invalid_request",
                "message": "This request could not be processed for this agreement.",
            },
        )

    trace = build_llm_trace_context(
        session_type=session_type,
        agreement_id=agreement_id,
        request=request,
    )
    remaining = 9999
    fp = client_fingerprint(request)
    if session_type == "recipient" and OPENAI_API_KEY:
        acquired, remaining = recipient_try_acquire_llm_slot(agreement_id, fp)
        if not acquired:
            return JSONResponse(
                status_code=429,
                content={
                    "error": "rate_limited",
                    "message": "AI usage limit reached for this session",
                },
                headers=usage_response_header(0),
            )
    elif session_type == "recipient":
        remaining = peek_recipient_remaining(agreement_id, fp)

    revised = _revise_with_instruction(
        draft,
        instruction,
        trace_context=trace,
        ai_model_class=body.ai_model_class,
    )
    revised = _coalesce_revision_draft_with_base(draft, revised)
    if session_type == "recipient":
        revised = _maybe_apply_recipient_deterministic_no_op_patch(draft, instruction, revised)
    revision_validation = _validate_revision_expectations(draft, revised, instruction)
    now = _utc_now_iso()
    persist = bool(getattr(body, "persist", True))
    next_draft = _merge_agreement_draft(
        draft,
        title=revised.title,
        jurisdiction=revised.jurisdiction,
        parties=_ensure_agreement_parties_have_ids(list(revised.parties or [])),
        purpose=revised.purpose,
        payment_terms=revised.payment_terms,
        duration=revised.duration,
        due_date=revised.due_date,
        effective_date=revised.effective_date,
        updated_at=now,
        audit_log=(
            [
                *(draft.audit_log or []),
                AuditEvent(event_type="field_updated", at=now, field="chat_revise", value=instruction),
            ]
            if persist
            else list(draft.audit_log or [])
        ),
    )
    if persist:
        _save_draft_sync(next_draft.model_dump(), request)
    else:
        record_usage_ledger_event(
            subject_ref=resolve_subject_from_request(request),
            event_type="revision_preview_used",
            agreement_id=agreement_id,
            metadata={"session_type": session_type},
        )
    for hk, hv in usage_response_header(remaining if session_type == "recipient" else 9999).items():
        response.headers[hk] = hv
    wm = _watermark_active_for_agreement(agreement_id)
    if session_type == "recipient":
        try:
            tok = (request.headers.get("X-Claw-Recipient-Access-Token") or "").strip()
            blob_before = _revision_comparison_blob(draft)
            blob_after = _revision_comparison_blob(next_draft)
            fields_changed = blob_before != blob_after
            html_before = _render_html(draft, watermark=wm)
            html_after = _render_html(next_draft, watermark=wm)
            rendered_changed = _norm_revision_comparison_text(html_before) != _norm_revision_comparison_text(
                html_after
            )
            log.info(
                "[recipient-revise] agreement_id=%s instruction_len=%s recipient_token_present=%s "
                "fields_changed=%s rendered_changed=%s",
                agreement_id,
                len(instruction or ""),
                bool(tok),
                fields_changed,
                rendered_changed,
            )
        except Exception as exc:
            log.warning(
                "[recipient-revise] diag_failed agreement_id=%s exc=%s",
                agreement_id,
                type(exc).__name__,
            )
    return {
        "id": agreement_id,
        "draft": next_draft.model_dump(),
        "rendered_html": _render_html(next_draft, watermark=wm),
        "canonical_json": canonicalize_agreement(next_draft.model_dump()),
        "revision_validation": revision_validation,
    }


@router.post("/{agreement_id}/refine")
def refine_agreement(
    agreement_id: str,
    body: AgreementRefineRequest,
    request: Request,
    response: Response,
):
    """Create-flow alias: same behavior as POST /revise with owner session and persist."""
    return revise_agreement(
        agreement_id,
        AgreementReviseRequest(
            instruction=body.instruction,
            session_type="owner",
            persist=True,
        ),
        request,
        response,
    )


@router.post("/{agreement_id}/commit-revision")
def commit_agreement_revision(agreement_id: str, body: AgreementCommitRevisionRequest, request: Request):
    """Write a draft produced during preview; avoids a second LLM call."""
    _owner_mutation_guards(request, agreement_id, surface="commit_revision")
    _assert_negotiation_not_locked(agreement_id)
    current = _load_or_404(agreement_id)
    _assert_draft_mutable_after_signatures(current)
    now = _utc_now_iso()
    instruction = (body.instruction or "").strip()
    coalesced = _coalesce_revision_draft_with_base(current, body.draft)
    revision_validation = _validate_revision_expectations(current, coalesced, instruction)
    next_draft = _merge_agreement_draft(
        current,
        updated_at=now,
        audit_log=[
            *(current.audit_log or []),
            AuditEvent(
                event_type="field_updated",
                at=now,
                field="chat_revise",
                value=instruction or "revision_committed",
            ),
        ],
        title=coalesced.title,
        jurisdiction=coalesced.jurisdiction,
        parties=_ensure_agreement_parties_have_ids(list(coalesced.parties or [])),
        purpose=coalesced.purpose,
        payment_terms=coalesced.payment_terms,
        duration=coalesced.duration,
        due_date=coalesced.due_date,
        effective_date=coalesced.effective_date,
        feed_visibility=coalesced.feed_visibility,
        feed_party_anonymize=coalesced.feed_party_anonymize,
        feed_show_financial_summary=coalesced.feed_show_financial_summary,
        feed_anchor_network=coalesced.feed_anchor_network,
    )
    nd = next_draft.model_dump()
    _save_draft_sync(nd, request)
    record_public_feed_event_if_applicable(draft_dict=nd, event_type="revision_applied", at=now)
    wm = _watermark_active_for_agreement(agreement_id)
    return {
        "id": agreement_id,
        "draft": nd,
        "rendered_html": _render_html(next_draft, watermark=wm),
        "canonical_json": canonicalize_agreement(nd),
        "revision_validation": revision_validation,
    }


def _audit_event_dict(e: Any) -> Dict[str, Any]:
    if hasattr(e, "model_dump"):
        return e.model_dump()
    if isinstance(e, dict):
        return e
    return {}


def _recipient_proposal_closed_from_index(audit: List[Any], proposal_id: str, pending_index: int) -> bool:
    pid = (proposal_id or "").strip()
    for j in range(pending_index + 1, len(audit)):
        d = _audit_event_dict(audit[j])
        et = str(d.get("event_type") or "")
        val = d.get("value") or {}
        if str(val.get("proposal_id") or "").strip() != pid:
            continue
        if et in (
            "recipient_proposal_applied",
            "recipient_proposal_rejected",
            "recipient_proposal_superseded",
        ):
            return True
    return False


def _open_recipient_proposal_payloads(audit: Any) -> List[Dict[str, Any]]:
    """FIFO-ordered open pending proposals (each value dict includes proposal_id, submitted_at, …)."""
    entries = list(audit or [])
    open_vals: List[Dict[str, Any]] = []
    for i, e in enumerate(entries):
        d = _audit_event_dict(e)
        if (d.get("event_type") or "") != "recipient_proposal_pending":
            continue
        val = d.get("value") or {}
        if not isinstance(val, dict):
            continue
        pid = str(val.get("proposal_id") or "").strip()
        if not pid:
            continue
        if _recipient_proposal_closed_from_index(entries, pid, i):
            continue
        open_vals.append(val)
    open_vals.sort(key=lambda v: str(v.get("submitted_at") or ""))
    return open_vals


def _open_recipient_proposal_ids_set(audit: Any) -> set:
    return {str(v.get("proposal_id") or "").strip() for v in _open_recipient_proposal_payloads(audit) if v}


def _normalize_workflow_role(role: str) -> str:
    r = (role or "").strip().lower()
    if r in ("owner", "sender", "landlord"):
        return "owner"
    if r in ("signer", "signatory"):
        return "signer"
    if r in ("reviewer",):
        return "reviewer"
    if r in ("viewer", "counterparty", "fyi", "copy", "read_only", "readonly"):
        return "viewer"
    return r or "party"


def _approved_participant_ids(audit: Any) -> set:
    out: set = set()
    for e in audit or []:
        d = _audit_event_dict(e)
        et = str(d.get("event_type") or "")
        if et not in ("participant_approved", "recipient_approved"):
            continue
        val = d.get("value") or {}
        if not isinstance(val, dict):
            continue
        pid = str(val.get("participant_id") or "").strip()
        if pid:
            out.add(pid)
    return out


def _review_delivery_email_payload_rows_from_draft_dict(d: Dict[str, Any]) -> List[Dict[str, Any]]:
    """Shape for future email sends; ``review_url`` is always None here (mint happens client-side today)."""
    parties = d.get("parties") or []
    if not isinstance(parties, list) or not parties:
        return []
    oi = next(
        (
            i
            for i, p in enumerate(parties)
            if isinstance(p, dict) and _normalize_workflow_role(str(p.get("role") or "")) == "owner"
        ),
        0,
    )
    title = str(d.get("title") or "").strip() or "Untitled agreement"
    out: List[Dict[str, Any]] = []
    for i, party in enumerate(parties):
        if i == oi:
            continue
        if not isinstance(party, dict):
            continue
        name = str(party.get("name") or "").strip()
        email = str(party.get("email") or "").strip().lower()
        if not name or not email or "@" not in email:
            continue
        role = _normalize_workflow_role(str(party.get("role") or ""))
        if role == "owner":
            continue
        out.append(
            {
                "to": email,
                "party_name": name,
                "reviewer_name": name,
                "agreement_title": title,
                "review_url": None,
            }
        )
    return out


def _workspace_review_approval_rollups(d: Dict[str, Any]) -> tuple[int, int, bool]:
    """
    Returns (approved_reviewer_count, required_reviewer_count, all_reviewers_approved).

    ``required`` is max(len(reviewer party ids), 1). ``approved`` counts reviewer party ids present in
    ``_approved_participant_ids``. Legacy approvals without ``participant_id`` count as one approval when
    no reviewer ids matched (single-recipient flows).
    """
    parties = d.get("parties") or []
    audit = d.get("audit_log") or []
    approved_ids = _approved_participant_ids(audit)
    reviewer_ids: List[str] = []
    for p in parties:
        if _normalize_workflow_role(str((p or {}).get("role") or "")) != "reviewer":
            continue
        pid = str((p or {}).get("id") or "").strip()
        if pid:
            reviewer_ids.append(pid)
    required = len(reviewer_ids)
    if required < 1:
        required = 1
    approved = sum(1 for pid in reviewer_ids if pid in approved_ids)
    if approved == 0:
        for e in audit or []:
            if not isinstance(e, dict):
                continue
            et = str(e.get("event_type") or "")
            if et not in ("recipient_approved", "participant_approved"):
                continue
            val = e.get("value") or {}
            if isinstance(val, dict) and str(val.get("participant_id") or "").strip():
                continue
            approved = 1
            break
    open_props = len(_open_recipient_proposal_payloads(audit)) > 0
    proposal_applied = any(
        isinstance(e, dict) and str(e.get("event_type") or "") == "recipient_proposal_applied"
        for e in (audit or [])
    )
    all_done = (not open_props) and approved >= required and (approved > 0 or proposal_applied)
    return approved, required, all_done


def _signing_approval_gate_errors(draft: AgreementDraft) -> List[str]:
    """Each signer must have participant_approved / recipient_approved with matching participant_id once IDs exist."""
    parties = draft.parties or []
    if not any((p.id or "").strip() for p in parties):
        return []
    approved = _approved_participant_ids(draft.audit_log)
    missing: List[str] = []
    for p in parties:
        if _normalize_workflow_role(p.role) != "signer":
            continue
        pid = (p.id or "").strip()
        if pid and pid not in approved:
            missing.append(p.name or pid)
    return missing


def _signing_snapshot_dict(draft: AgreementDraft) -> Dict[str, Any]:
    return {
        "title": draft.title,
        "jurisdiction": draft.jurisdiction,
        "parties": [p.model_dump() for p in draft.parties or []],
        "purpose": draft.purpose,
        "payment_terms": draft.payment_terms,
        "duration": draft.duration,
        "due_date": draft.due_date,
        "effective_date": draft.effective_date,
    }


def _draft_locked_content_sha256(draft: AgreementDraft) -> str:
    """SHA-256 of canonical signing snapshot (stable across audit-only draft changes)."""
    return canon_sha256_hex(_signing_snapshot_dict(draft))


def _agreement_version_hash(agreement_id: str, locked_version_id: str, draft: AgreementDraft) -> str:
    return canon_sha256_hex(
        {
            "agreement_id": agreement_id,
            "locked_version_id": locked_version_id,
            "snapshot": _signing_snapshot_dict(draft),
        }
    )


def _signature_completed_participant_ids(audit: Any) -> set:
    out: set = set()
    for e in audit or []:
        d = _audit_event_dict(e)
        if str(d.get("event_type") or "") != "signature_completed":
            continue
        val = d.get("value") or {}
        if not isinstance(val, dict):
            continue
        pid = str(val.get("participant_id") or "").strip()
        if pid:
            out.add(pid)
    return out


def _has_legacy_signature_without_participant(audit: Any) -> bool:
    for e in audit or []:
        d = _audit_event_dict(e)
        if str(d.get("event_type") or "") != "signature_completed":
            continue
        val = d.get("value") or {}
        if isinstance(val, dict) and not str(val.get("participant_id") or "").strip():
            return True
    return False


def _assert_draft_mutable_after_signatures(draft: AgreementDraft) -> None:
    """Block canonical draft edits once any signer has completed the ceremony."""
    audit = draft.audit_log or []
    if _signature_completed_participant_ids(audit):
        raise HTTPException(status_code=400, detail="agreement_immutable_after_signature")
    if _has_legacy_signature_without_participant(audit):
        raise HTTPException(status_code=400, detail="agreement_immutable_after_signature")


def _signer_party_by_participant_id(draft: AgreementDraft, participant_id: str) -> Optional[AgreementParty]:
    pid = (participant_id or "").strip()
    if not pid:
        return None
    for p in draft.parties or []:
        if (p.id or "").strip() != pid:
            continue
        if _normalize_workflow_role(p.role) == "signer":
            return p
    return None


def _signing_ceremony_guards_lv(draft: AgreementDraft, agreement_id: str) -> Tuple[str, AgreementDraft]:
    draft = _persist_party_id_backfill(draft)
    if _agreement_draft_fully_executed(draft):
        raise HTTPException(status_code=400, detail="signing_already_complete")
    lock = read_signing_lock(agreement_id)
    lv = str((lock or {}).get("locked_version_id") or "").strip()
    if not lv:
        raise HTTPException(status_code=400, detail="not_ready_for_signature")
    if _open_recipient_proposal_payloads(draft.audit_log):
        raise HTTPException(status_code=400, detail="open_proposals_block_signing")
    return lv, draft


def _resolve_signing_participant_for_ceremony(
    draft: AgreementDraft, participant_id: str
) -> Tuple[str, AgreementParty]:
    """Returns (participant_id_for_audit, signer_party)."""
    part_id = (participant_id or "").strip()
    parties_have_ids = any((p.id or "").strip() for p in draft.parties or [])
    signer_parties = [p for p in draft.parties or [] if _normalize_workflow_role(p.role) == "signer"]
    if not signer_parties:
        raise HTTPException(status_code=400, detail="no_signers_on_agreement")
    if parties_have_ids:
        if not part_id:
            raise HTTPException(status_code=400, detail="participant_id_required")
        sp = _signer_party_by_participant_id(draft, part_id)
        if not sp:
            raise HTTPException(status_code=403, detail="signer_not_found")
        approved = _approved_participant_ids(draft.audit_log)
        if part_id not in approved:
            raise HTTPException(status_code=403, detail="participant_not_approved")
        return part_id, sp
    miss = _signing_approval_gate_errors(draft)
    if miss:
        raise HTTPException(
            status_code=403,
            detail={"code": "approvals_incomplete", "missing_signer_approvals": miss},
        )
    if len(signer_parties) != 1:
        raise HTTPException(status_code=400, detail="participant_id_required")
    sp = signer_parties[0]
    return (sp.id or "").strip(), sp


def _all_signers_signed_from_audit(draft: AgreementDraft, audit: List[Any]) -> bool:
    signers = [p for p in draft.parties or [] if _normalize_workflow_role(p.role) == "signer"]
    if not signers:
        return False
    done = _signature_completed_participant_ids(audit)
    ids = [(p.id or "").strip() for p in signers]
    if all(ids):
        return bool(ids) and all(i in done for i in ids)
    if len(signers) == 1:
        return _has_legacy_signature_without_participant(audit)
    return False


def public_agreement_verify_enabled() -> bool:
    v = os.getenv("CLAW_PUBLIC_AGREEMENT_VERIFY", "1").strip().lower()
    return v not in ("0", "false", "no", "off")


def _public_agreement_overview_hash(agreement_id: str, draft: AgreementDraft) -> str:
    """SHA-256 of publishable metadata + version index (no purpose, payment, or body text)."""
    party_pub = [{"name": p.name, "role": p.role} for p in draft.parties or []]
    ver_rows = [{"version": v.version, "created_at": v.created_at, "note": v.note} for v in draft.versions or []]
    payload = {
        "agreement_id": agreement_id,
        "title": draft.title,
        "jurisdiction": draft.jurisdiction,
        "created_at": draft.created_at,
        "updated_at": draft.updated_at,
        "parties": party_pub,
        "versions": ver_rows,
    }
    return canon_sha256_hex(payload)


def _public_version_history(draft: AgreementDraft) -> List[Dict[str, Any]]:
    out: List[Dict[str, Any]] = []
    for v in draft.versions or []:
        vd = {"version": v.version, "created_at": v.created_at, "note": v.note}
        out.append({**vd, "version_hash": canon_sha256_hex(vd)})
    return out


def _public_signature_events(audit: Any) -> List[Dict[str, Any]]:
    events: List[Dict[str, Any]] = []
    for e in audit or []:
        d = _audit_event_dict(e)
        et = str(d.get("event_type") or "")
        if et not in ("signature_initiated", "signature_completed", "signed"):
            continue
        val = d.get("value") or {}
        if not isinstance(val, dict):
            val = {}
        row: Dict[str, Any] = {
            "event_type": et,
            "at": d.get("at"),
            "agreement_version_hash": val.get("agreement_version_hash"),
            "locked_version_id": val.get("locked_version_id"),
        }
        if et in ("signature_initiated", "signature_completed"):
            row["participant_display_name"] = val.get("participant_display_name")
        if et == "signature_completed" and val.get("typed_name"):
            row["typed_name"] = val.get("typed_name")
        if et == "signed" and val.get("fully_executed"):
            row["fully_executed"] = True
        events.append(row)
    return events


def _public_lifecycle_label(draft: AgreementDraft, agreement_id: str) -> str:
    audit = list(draft.audit_log or [])
    if any(_audit_event_dict(e).get("event_type") == "signed" for e in audit):
        return "fully_executed"
    if _all_signers_signed_from_audit(draft, audit):
        return "fully_executed"
    if _signature_completed_participant_ids(audit) or _has_legacy_signature_without_participant(audit):
        return "partially_signed"
    lock = read_signing_lock(agreement_id)
    if lock and str(lock.get("locked_version_id") or "").strip():
        return "locked_for_signing"
    return "in_negotiation"


def _public_signature_status(agreement_id: str, draft: AgreementDraft) -> Dict[str, Any]:
    audit = list(draft.audit_log or [])
    fully = any(_audit_event_dict(e).get("event_type") == "signed" for e in audit) or _all_signers_signed_from_audit(
        draft, audit
    )
    n_completed = sum(1 for e in audit if _audit_event_dict(e).get("event_type") == "signature_completed")
    lock = read_signing_lock(agreement_id)
    lv = str((lock or {}).get("locked_version_id") or "").strip() or None
    commitment = _agreement_version_hash(agreement_id, lv, draft) if lv else None
    signers = [p for p in draft.parties or [] if _normalize_workflow_role(p.role) == "signer"]
    return {
        "fully_executed": fully,
        "signatures_recorded": n_completed,
        "signer_party_count": len(signers),
        "locked_version_id": lv,
        "signing_commitment_hash": commitment,
    }


def _validate_nonowner_proposer(draft: AgreementDraft, proposer_id: str) -> AgreementParty:
    pid = (proposer_id or "").strip()
    if not pid:
        raise HTTPException(status_code=400, detail="proposer_id_required")
    for p in draft.parties or []:
        if (p.id or "").strip() != pid:
            continue
        wr = _normalize_workflow_role(p.role)
        if wr == "owner":
            raise HTTPException(status_code=403, detail="owner_cannot_submit_recipient_proposal")
        if wr == "viewer":
            raise HTTPException(status_code=403, detail="viewer_cannot_submit_recipient_proposal")
        return p
    raise HTTPException(status_code=400, detail="proposer_not_found")


def _persist_party_id_backfill(draft: AgreementDraft) -> AgreementDraft:
    """Assign UUIDs to parties missing ids and persist (safe migration)."""
    if not any(not (p.id or "").strip() for p in (draft.parties or [])):
        return draft
    parties = _ensure_agreement_parties_have_ids(list(draft.parties or []))
    now = _utc_now_iso()
    next_draft = _merge_agreement_draft(draft, updated_at=now, parties=parties)
    _save_draft_sync(next_draft.model_dump())
    return next_draft


def _resolve_assumed_owner_party_index(parties: List[AgreementParty]) -> int:
    """Mirror review-link owner index: first explicit owner, else index 0."""
    for i, p in enumerate(parties or []):
        if _normalize_workflow_role(p.role) == "owner":
            return i
    return 0


def _review_link_eligible_party_ids(draft: AgreementDraft) -> List[str]:
    """Non-owner counterparty rows eligible for review-link / staged proposal."""
    parties = list(draft.parties or [])
    owner_idx = _resolve_assumed_owner_party_index(parties)
    eligible: List[str] = []
    for i, p in enumerate(parties):
        if i == owner_idx:
            continue
        pid = (p.id or "").strip()
        if not pid or pid.startswith("legacy_"):
            continue
        if _normalize_workflow_role(p.role) == "viewer":
            continue
        eligible.append(pid)
    return eligible


def _parse_recipient_access_token_context(request: Request, agreement_id: str) -> Dict[str, Any]:
    tok = recipient_access_token_from_request(request)
    if not tok:
        return {
            "has_auth_token": False,
            "token_valid": False,
            "token_pid": "",
            "token_role": "",
            "token_mode": "",
        }
    try:
        secret_raw = resolve_signing_token_secret_raw()
    except SigningTokenSecretMissingInProductionError:
        return {
            "has_auth_token": True,
            "token_valid": False,
            "token_pid": "",
            "token_role": "",
            "token_mode": "",
        }
    try:
        out = validate_recipient_access_token_for_agreement(
            token=tok,
            path_agreement_id=agreement_id,
            query_agreement_id=None,
            secret_raw=secret_raw,
            consume_single_use=False,
            log_validation=False,
        )
        return {
            "has_auth_token": True,
            "token_valid": True,
            "token_pid": str(out.get("recipient_party_id") or "").strip(),
            "token_role": str(out.get("role") or "").strip().lower(),
            "token_mode": str(out.get("mode") or "").strip().lower(),
        }
    except HTTPException:
        return {
            "has_auth_token": True,
            "token_valid": False,
            "token_pid": "",
            "token_role": "",
            "token_mode": "",
        }


def _infer_proposer_from_eligible_parties(
    draft: AgreementDraft, eligible: List[str], token_role: str
) -> Tuple[str, str]:
    if not eligible:
        return "", ""
    if len(eligible) == 1:
        return eligible[0], "inferred_single_reviewer"
    reviewer_ids: List[str] = []
    signer_ids: List[str] = []
    for p in draft.parties or []:
        pid = (p.id or "").strip()
        if pid not in eligible:
            continue
        wr = _normalize_workflow_role(p.role)
        if wr == "reviewer":
            reviewer_ids.append(pid)
        elif wr == "signer":
            signer_ids.append(pid)
    role = (token_role or "reviewer").strip().lower()
    if role in ("reviewer", "recipient") and len(reviewer_ids) == 1:
        return reviewer_ids[0], "inferred_reviewer_role"
    if role == "signer" and len(signer_ids) == 1:
        return signer_ids[0], "inferred_signer_role"
    if len(reviewer_ids) == 1:
        return reviewer_ids[0], "inferred_reviewer_role"
    return "", ""


def _log_recipient_proposal_stage_proposer_resolution(
    agreement_id: str,
    token_ctx: Dict[str, Any],
    body_proposer_id: str,
    inferred_party_id: str,
    inference_source: str,
    party_count: int,
    recipient_count: int,
    failure_reason: str,
) -> None:
    log.info(
        "[recipient-proposal-stage-proposer-resolution] agreement_id=%s has_auth_token=%s token_valid=%s token_pid=%s body_proposer_id=%s inferred_party_id=%s inference_source=%s party_count=%s recipient_count=%s failure_reason=%s",
        agreement_id,
        token_ctx.get("has_auth_token"),
        token_ctx.get("token_valid"),
        token_ctx.get("token_pid") or "",
        (body_proposer_id or "").strip(),
        inferred_party_id or "",
        inference_source or "",
        party_count,
        recipient_count,
        failure_reason or "",
    )


def _resolve_recipient_proposer_with_source(
    request: Request, agreement_id: str, draft: AgreementDraft, body_proposer_id: str
) -> Tuple[AgreementParty, str]:
    parties = list(draft.parties or [])
    party_count = len(parties)
    eligible = _review_link_eligible_party_ids(draft)
    recipient_count = len(eligible)
    token_ctx = _parse_recipient_access_token_context(request, agreement_id)
    body_pid = (body_proposer_id or "").strip()

    if body_pid:
        party = _validate_nonowner_proposer(draft, body_pid)
        _log_recipient_proposal_stage_proposer_resolution(
            agreement_id,
            token_ctx,
            body_pid,
            body_pid,
            "body",
            party_count,
            recipient_count,
            "",
        )
        return party, "body"

    token_pid = str(token_ctx.get("token_pid") or "").strip()
    if token_pid and token_ctx.get("token_valid"):
        party = _validate_nonowner_proposer(draft, token_pid)
        _log_recipient_proposal_stage_proposer_resolution(
            agreement_id,
            token_ctx,
            body_pid,
            token_pid,
            "token",
            party_count,
            recipient_count,
            "",
        )
        return party, "token"

    if not token_ctx.get("has_auth_token"):
        _log_recipient_proposal_stage_proposer_resolution(
            agreement_id,
            token_ctx,
            body_pid,
            "",
            "",
            party_count,
            recipient_count,
            "no_auth_token",
        )
        raise HTTPException(status_code=400, detail="proposer_id_required")

    if not token_ctx.get("token_valid"):
        _log_recipient_proposal_stage_proposer_resolution(
            agreement_id,
            token_ctx,
            body_pid,
            "",
            "",
            party_count,
            recipient_count,
            "invalid_token",
        )
        raise HTTPException(status_code=400, detail="proposer_id_required")

    inferred, source = _infer_proposer_from_eligible_parties(
        draft, eligible, str(token_ctx.get("token_role") or "")
    )
    if inferred:
        party = _validate_nonowner_proposer(draft, inferred)
        _log_recipient_proposal_stage_proposer_resolution(
            agreement_id,
            token_ctx,
            body_pid,
            inferred,
            source,
            party_count,
            recipient_count,
            "",
        )
        return party, source

    failure_reason = "ambiguous_reviewer_party" if recipient_count > 1 else "no_eligible_reviewer_party"
    _log_recipient_proposal_stage_proposer_resolution(
        agreement_id,
        token_ctx,
        body_pid,
        "",
        "",
        party_count,
        recipient_count,
        failure_reason,
    )
    raise HTTPException(status_code=400, detail="proposer_id_required")


def _resolve_recipient_proposer(
    request: Request, agreement_id: str, draft: AgreementDraft, body_proposer_id: str
) -> AgreementParty:
    party, _source = _resolve_recipient_proposer_with_source(
        request, agreement_id, draft, body_proposer_id
    )
    return party


def _staged_recipient_proposals_map(draft: AgreementDraft) -> Dict[str, Any]:
    pr = draft.pro_redline_v1 if isinstance(draft.pro_redline_v1, dict) else {}
    staged = pr.get(STAGED_RECIPIENT_PROPOSALS_KEY)
    return dict(staged) if isinstance(staged, dict) else {}


def _persist_staged_recipient_proposal(
    draft: AgreementDraft,
    proposal_id: str,
    payload: Dict[str, Any],
    request: Optional[Request] = None,
) -> AgreementDraft:
    now = _utc_now_iso()
    pro_redline = dict(draft.pro_redline_v1 or {})
    staged = _staged_recipient_proposals_map(draft)
    staged[proposal_id] = payload
    pro_redline[STAGED_RECIPIENT_PROPOSALS_KEY] = staged
    next_draft = _merge_agreement_draft(draft, updated_at=now, pro_redline_v1=pro_redline)
    _save_draft_sync(next_draft.model_dump(), request)
    return next_draft


def _pop_staged_recipient_proposal(
    draft: AgreementDraft,
    proposal_id: str,
    request: Optional[Request] = None,
) -> Optional[Dict[str, Any]]:
    pid = (proposal_id or "").strip()
    if not pid:
        return None
    staged = _staged_recipient_proposals_map(draft)
    payload = staged.pop(pid, None)
    if payload is None:
        return None
    now = _utc_now_iso()
    pro_redline = dict(draft.pro_redline_v1 or {})
    pro_redline[STAGED_RECIPIENT_PROPOSALS_KEY] = staged
    next_draft = _merge_agreement_draft(draft, updated_at=now, pro_redline_v1=pro_redline)
    _save_draft_sync(next_draft.model_dump(), request)
    return payload if isinstance(payload, dict) else None


def _proposal_value_for_id(audit: Any, proposal_id: str) -> Optional[Dict[str, Any]]:
    pid = (proposal_id or "").strip()
    if not pid:
        return None
    for e in reversed(list(audit or [])):
        d = _audit_event_dict(e)
        if (d.get("event_type") or "") != "recipient_proposal_pending":
            continue
        val = d.get("value") or {}
        if str(val.get("proposal_id") or "").strip() == pid:
            return val if isinstance(val, dict) else None
    return None


def _queue_recipient_proposal_from_payload(
    draft: AgreementDraft,
    payload: Dict[str, Any],
    request: Optional[Request] = None,
) -> AgreementDraft:
    proposal_id = str(payload.get("proposal_id") or "").strip()
    proposer_key = str(payload.get("proposer_id") or "").strip()
    for v in _open_recipient_proposal_payloads(draft.audit_log):
        if str(v.get("proposer_id") or "").strip() == proposer_key:
            raise HTTPException(
                status_code=409,
                detail="recipient_proposal_already_pending_from_participant",
            )
    now = _utc_now_iso()
    audit = [*(draft.audit_log or [])]
    audit.append(
        AuditEvent(
            event_type="recipient_proposal_pending",
            at=now,
            field="recipient_proposal",
            value=payload,
        )
    )
    audit.append(
        AuditEvent(
            event_type="participant_proposed_revision",
            at=now,
            field="recipient_proposal",
            value={
                "proposal_id": proposal_id,
                "participant_id": proposer_key,
                "display_name": str(payload.get("proposer_display_name") or "").strip(),
                "instruction": str(payload.get("instruction") or "")[:512],
            },
        )
    )
    next_draft = _merge_agreement_draft(draft, updated_at=now, audit_log=audit)
    _save_draft_sync(next_draft.model_dump(), request)
    return next_draft


@router.post("/{agreement_id}/recipient-proposal/stage")
def stage_recipient_proposal(
    agreement_id: str, body: RecipientProposalRequest, request: Request
) -> Dict[str, Any]:
    """Stage a recipient proposal payload; finalize with POST /recipient-proposal + proposal_id."""
    log.info("[recipient-proposal-stage] start agreement_id=%s", agreement_id)
    try:
        assert_free_incomplete_draft_not_expired(agreement_id, surface="recipient_proposal_stage")
        draft = _load_or_404(agreement_id)
        lock = read_signing_lock(agreement_id)
        if lock and bool((lock or {}).get("locked_version_id")):
            raise HTTPException(status_code=400, detail="negotiation_locked")
        draft = _persist_party_id_backfill(draft)
        proposer, proposer_source = _resolve_recipient_proposer_with_source(
            request, agreement_id, draft, body.proposer_id
        )
        proposer_id = (proposer.id or "").strip()
        log.info(
            "[recipient-proposal-stage] participant resolved agreement_id=%s proposer_id=%s proposer_id_source=%s",
            agreement_id,
            proposer_id,
            proposer_source,
        )
        assert_agreement_recipient_write_allowed(
            request,
            agreement_id,
            allowed_modes=("review",),
            bind_participant_id=proposer_id,
        )
        instruction = (body.instruction or "").strip()
        if not instruction:
            raise HTTPException(status_code=400, detail="instruction_required")
        proposed_purpose = (body.draft.purpose or "").strip()
        if not proposed_purpose:
            raise HTTPException(status_code=400, detail="proposed_draft_purpose_required")
        proposal_id = str(uuid.uuid4())
        now = _utc_now_iso()
        dname = (body.proposer_display_name or "").strip() or proposer.name
        payload: Dict[str, Any] = {
            "proposal_id": proposal_id,
            "instruction": instruction,
            "draft": body.draft.model_dump(),
            "rendered_html": (body.rendered_html or "").strip(),
            "staged_at": now,
            "proposer_id": proposer_id,
            "proposer_display_name": dname,
        }
        log.info(
            "[recipient-proposal-stage] payload summary agreement_id=%s proposal_id=%s instruction_len=%s proposed_purpose_len=%s rendered_html_len=%s canonical_purpose_len=%s",
            agreement_id,
            proposal_id,
            len(instruction),
            len(proposed_purpose),
            len(payload["rendered_html"]),
            len((draft.purpose or "").strip()),
        )
        _persist_staged_recipient_proposal(draft, proposal_id, payload, request)
        log.info(
            "[recipient-proposal-stage] stored agreement_id=%s proposal_id=%s proposer_id=%s",
            agreement_id,
            proposal_id,
            proposer_id,
        )
        return {
            "ok": True,
            "proposal_id": proposal_id,
            "staged": True,
            "proposer_id": proposer_id,
            "proposer_id_source": proposer_source,
        }
    except HTTPException:
        raise
    except Exception as exc:
        log.exception(
            "[recipient-proposal-stage] failed agreement_id=%s error=%s",
            agreement_id,
            exc,
        )
        raise HTTPException(status_code=500, detail="recipient_proposal_stage_failed") from exc


@router.post("/{agreement_id}/recipient-proposal")
def submit_recipient_proposal(
    agreement_id: str, body: RecipientProposalFinalizeBody, request: Request
) -> Dict[str, Any]:
    """Finalize a staged recipient proposal (audit-only pending until owner applies)."""
    proposal_id = (body.proposal_id or "").strip()
    if not proposal_id:
        raise HTTPException(status_code=400, detail="proposal_id_required")
    assert_free_incomplete_draft_not_expired(agreement_id, surface="recipient_proposal_submit")
    draft = _load_or_404(agreement_id)
    lock = read_signing_lock(agreement_id)
    if lock and bool((lock or {}).get("locked_version_id")):
        raise HTTPException(status_code=400, detail="negotiation_locked")
    draft = _persist_party_id_backfill(draft)
    staged = _pop_staged_recipient_proposal(draft, proposal_id, request)
    if not staged:
        raise HTTPException(status_code=400, detail="proposal_not_staged")
    proposer_id = str(staged.get("proposer_id") or "").strip()
    assert_agreement_recipient_write_allowed(
        request,
        agreement_id,
        allowed_modes=("review",),
        bind_participant_id=proposer_id,
    )
    staged["submitted_at"] = _utc_now_iso()
    next_draft = _queue_recipient_proposal_from_payload(draft, staged, request)
    return {"ok": True, "proposal_id": proposal_id, "draft": next_draft.model_dump()}


def _corpus_fingerprint(text: str) -> str:
    body = (text or "").strip()
    digest = hashlib.sha256(body.encode("utf-8")).hexdigest()[:16]
    return f"{len(body)}:{digest}"


@router.post("/{agreement_id}/recipient-proposal/{proposal_id}/reject")
def reject_recipient_proposal(
    agreement_id: str, proposal_id: str, request: Request
) -> Dict[str, Any]:
    _owner_mutation_guards(request, agreement_id, surface="recipient_proposal_reject")
    draft = _load_or_404(agreement_id)
    open_ids = _open_recipient_proposal_ids_set(draft.audit_log)
    pid = (proposal_id or "").strip()
    if pid not in open_ids:
        raise HTTPException(status_code=400, detail="proposal_not_pending")
    previous_hash = _corpus_fingerprint(draft.purpose or "")
    now = _utc_now_iso()
    audit = [*(draft.audit_log or [])]
    audit.append(
        AuditEvent(
            event_type="recipient_proposal_rejected",
            at=now,
            field="recipient_proposal",
            value={"proposal_id": pid},
        )
    )
    next_draft = _merge_agreement_draft(draft, updated_at=now, audit_log=audit)
    _save_draft_sync(next_draft.model_dump(), request)
    rejected_hash = _corpus_fingerprint(next_draft.purpose or "")
    log.info(
        "[owner-proposal-decline] agreement_id=%s proposal_id=%s previous_corpus_hash=%s declined_corpus_hash=%s",
        agreement_id,
        pid,
        previous_hash,
        rejected_hash,
    )
    log.info(
        "[owner-proposal-rejected] agreement_id=%s proposal_id=%s proposal_status=rejected previous_corpus_hash=%s rejected_corpus_hash=%s",
        agreement_id,
        pid,
        previous_hash,
        rejected_hash,
    )
    return {"ok": True, "draft": next_draft.model_dump()}


@router.post("/{agreement_id}/recipient-proposal/{proposal_id}/apply")
def apply_recipient_proposal(
    agreement_id: str, proposal_id: str, request: Request
) -> Dict[str, Any]:
    """Merge a pending recipient proposal into the canonical draft (owner action)."""
    _owner_mutation_guards(request, agreement_id, surface="recipient_proposal_apply")
    _assert_negotiation_not_locked(agreement_id)
    current = _load_or_404(agreement_id)
    open_payloads = _open_recipient_proposal_payloads(current.audit_log)
    open_ids = {str(v.get("proposal_id") or "").strip() for v in open_payloads}
    pid_apply = (proposal_id or "").strip()
    if pid_apply not in open_ids:
        raise HTTPException(status_code=400, detail="proposal_not_pending")
    pending = _proposal_value_for_id(current.audit_log, proposal_id)
    if not pending:
        raise HTTPException(status_code=404, detail="proposal_not_found")
    inner = pending.get("draft")
    if not isinstance(inner, dict):
        raise HTTPException(status_code=400, detail="invalid_proposal_payload")
    instruction = str(pending.get("instruction") or "").strip() or "recipient_proposal_applied"
    try:
        body_create = AgreementDraftCreate.model_validate(inner)
    except Exception:
        raise HTTPException(status_code=400, detail="invalid_proposal_draft")
    merged_parties = _ensure_agreement_parties_have_ids(list(body_create.parties or []))
    now = _utc_now_iso()
    other_ids = [str(v.get("proposal_id") or "").strip() for v in open_payloads if str(v.get("proposal_id") or "").strip() != pid_apply]
    tail_events: List[AuditEvent] = [
        AuditEvent(
            event_type="field_updated",
            at=now,
            field="chat_revise",
            value=instruction,
        ),
        AuditEvent(
            event_type="recipient_proposal_applied",
            at=now,
            field="recipient_proposal",
            value={"proposal_id": proposal_id},
        ),
    ]
    for oid in other_ids:
        tail_events.append(
            AuditEvent(
                event_type="recipient_proposal_superseded",
                at=now,
                field="recipient_proposal",
                value={
                    "proposal_id": oid,
                    "reason": "another_proposal_applied",
                    "applied_proposal_id": pid_apply,
                },
            )
        )
    accepted_corpus = (body_create.purpose or "").strip()
    pro_redline = (
        dict(current.pro_redline_v1 or {})
        if isinstance(current.pro_redline_v1, dict)
        else {}
    )
    if len(accepted_corpus) >= 80:
        pro_redline["review_first_final_corpus"] = {
            "text": accepted_corpus,
            "source": "recipient_proposal_applied",
            "hash": _corpus_fingerprint(accepted_corpus),
            "persisted_at": now,
        }
    next_draft = AgreementDraft(
        id=current.id,
        created_at=current.created_at,
        updated_at=now,
        versions=list(current.versions or []),
        review_sent_at=current.review_sent_at,
        workspace_archived_at=current.workspace_archived_at,
        workspace_folder_id=current.workspace_folder_id,
        workspace_tags=list(current.workspace_tags or []),
        audit_log=[*(current.audit_log or []), *[e for e in tail_events]],
        title=body_create.title,
        jurisdiction=body_create.jurisdiction,
        parties=merged_parties,
        purpose=body_create.purpose,
        payment_terms=body_create.payment_terms,
        duration=body_create.duration,
        due_date=body_create.due_date,
        effective_date=body_create.effective_date,
        payment_request=current.payment_request,
        payment_required=current.payment_required,
        pro_redline_v1=pro_redline if pro_redline else current.pro_redline_v1,
    )
    _save_draft_sync(next_draft.model_dump(), request)
    previous_hash = _corpus_fingerprint(current.purpose or "")
    updated_hash = _corpus_fingerprint(next_draft.purpose or "")
    log.info(
        "[owner-proposal-accepted] agreement_id=%s proposal_id=%s proposal_status=accepted previous_corpus_hash=%s updated_corpus_hash=%s accepted_corpus_hash=%s",
        agreement_id,
        pid_apply,
        previous_hash,
        updated_hash,
        updated_hash,
    )
    log.info(
        "[owner-corpus-updated] agreement_id=%s proposal_id=%s previous_corpus_hash=%s updated_corpus_hash=%s source=accept",
        agreement_id,
        pid_apply,
        previous_hash,
        updated_hash,
    )
    wm = _watermark_active_for_agreement(agreement_id)
    return {
        "id": agreement_id,
        "draft": next_draft.model_dump(),
        "rendered_html": _render_html(next_draft, watermark=wm),
        "canonical_json": canonicalize_agreement(next_draft.model_dump()),
    }


@router.post("/{agreement_id}/recipient-approve")
def recipient_approve_agreement(
    agreement_id: str,
    request: Request,
    body: RecipientApproveBody = RecipientApproveBody(),
) -> Dict[str, Any]:
    """Recipient acknowledges the current draft as acceptable (audit only; does not lock signing)."""
    assert_free_incomplete_draft_not_expired(agreement_id, surface="recipient_approve")
    assert_agreement_recipient_write_allowed(
        request,
        agreement_id,
        allowed_modes=("review",),
        bind_participant_id=body.participant_id,
    )
    draft = _load_or_404(agreement_id)
    imm = draft.audit_log or []
    if _signature_completed_participant_ids(imm) or _has_legacy_signature_without_participant(imm):
        raise HTTPException(status_code=400, detail="agreement_immutable_after_signature")
    lock = read_signing_lock(agreement_id)
    if lock and bool((lock or {}).get("locked_version_id")):
        raise HTTPException(status_code=400, detail="negotiation_locked")
    draft = _persist_party_id_backfill(draft)
    now = _utc_now_iso()
    msg = (body.message or "").strip()
    part_id = (body.participant_id or "").strip()
    part_name = (body.participant_display_name or "").strip()
    if part_id:
        found = False
        for p in draft.parties or []:
            if (p.id or "").strip() == part_id:
                found = True
                if not part_name:
                    part_name = p.name
                wr = _normalize_workflow_role(p.role)
                if wr == "owner":
                    raise HTTPException(status_code=403, detail="owner_uses_workspace_not_recipient_approve")
                if wr == "viewer":
                    raise HTTPException(status_code=403, detail="viewer_cannot_approve")
                break
        if not found:
            raise HTTPException(status_code=400, detail="participant_not_found")
    elif any((p.id or "").strip() for p in (draft.parties or [])):
        raise HTTPException(status_code=400, detail="participant_id_required")
    audit = [*(draft.audit_log or [])]
    approve_val: Dict[str, Any] = {"message": msg or "approved_current_draft"}
    if part_id:
        approve_val["participant_id"] = part_id
        approve_val["participant_display_name"] = part_name
        audit.append(
            AuditEvent(
                event_type="participant_approved",
                at=now,
                field="recipient",
                value=approve_val,
            )
        )
    audit.append(
        AuditEvent(
            event_type="recipient_approved",
            at=now,
            field="recipient",
            value=approve_val,
        )
    )
    next_draft = _merge_agreement_draft(draft, updated_at=now, audit_log=audit)
    _save_draft_sync(next_draft.model_dump(), request)
    try:
        from backend.services.email.review_delivery import (
            maybe_notify_counterparties_all_reviews_complete,
            maybe_notify_owner_after_reviewer_approval,
        )

        notify_audit = maybe_notify_owner_after_reviewer_approval(
            agreement_id=agreement_id,
            draft=next_draft.model_dump(),
            approver_participant_id=part_id or None,
            approver_display_name=part_name or None,
        )
        notify_events = [notify_audit] if notify_audit else []
        counterparty_audit = maybe_notify_counterparties_all_reviews_complete(
            agreement_id=agreement_id,
            draft=next_draft.model_dump(),
        )
        if counterparty_audit:
            notify_events.append(counterparty_audit)
        if notify_events:
            audit_with_notify = [*(next_draft.audit_log or []), *notify_events]
            next_draft = _merge_agreement_draft(
                next_draft,
                updated_at=_utc_now_iso(),
                audit_log=audit_with_notify,
            )
            _save_draft_sync(next_draft.model_dump(), request)
    except Exception:
        logging.getLogger(__name__).exception(
            "owner_review_notification_hook_failed agreement_id=%s",
            agreement_id,
        )
    return {"ok": True, "draft": next_draft.model_dump()}


# ---------------------------------------------------------------------------
# Agreement proof layer (deterministic receipt + Merkle batch queue)
# ---------------------------------------------------------------------------


def _batch_row_for_receipt(store: TimelineStore, rec: Dict[str, Any]) -> Optional[Dict[str, Any]]:
    bid = rec.get("batch_id")
    if not bid:
        return None
    try:
        return get_batch(store=store, batch_id=str(bid))
    except HTTPException:
        return None


class AgreementFinalizeReceiptRequest(BaseModel):
    finalized_version_id: str
    finalized_at: str
    content_sha256: str
    execution_packet_sha256: str
    parties_sha256: Optional[str] = None
    signer_count: Optional[int] = None
    anchor_network: str = "bitcoin-testnet"
    epoch_id: Optional[str] = None
    execution_packet: Optional[Dict[str, Any]] = Field(
        default=None,
        description="Optional full packet JSON; stored when unified artifact store is enabled.",
    )


def _persist_agreement_execution_packet_artifact(
    *,
    agreement_id: str,
    receipt_id: str,
    finalized_version_id: str,
    execution_packet_sha256: str,
    execution_packet: Optional[Dict[str, Any]],
) -> None:
    if execution_packet is None:
        return
    from backend.config.storage_runtime import unified_artifact_store_enabled
    from backend.proof.execution_packet_digest import (
        assert_execution_packet_matches_digest,
        execution_packet_canonical_json_bytes,
    )
    from backend.storage.artifact_repository import get_artifact_repository

    if not unified_artifact_store_enabled():
        return
    assert_execution_packet_matches_digest(
        execution_packet, declared_sha256_hex=execution_packet_sha256
    )
    payload = execution_packet_canonical_json_bytes(execution_packet)
    get_artifact_repository().put_artifact(
        artifact_type="agreement_execution_packet",
        logical_ref=receipt_id,
        data=payload,
        content_type="application/json",
        visibility="private",
        agreement_id=agreement_id,
        version_id=finalized_version_id,
        metadata={"execution_packet_sha256": execution_packet_sha256},
    )


def _agreement_anchor_proof_view(
    receipt: Dict[str, Any], batch: Optional[Dict[str, Any]]
) -> Dict[str, Any]:
    net = (receipt.get("network") or "").strip()
    b_status = ((batch or {}).get("anchor_status") or "").strip().lower()
    bt_raw = ((batch or {}).get("anchor_txid") or "").strip()
    b_err = ((batch or {}).get("anchor_error") or "").strip()
    b_attempts = batch.get("anchor_attempts") if batch else None

    txid: Optional[str] = None
    if bt_raw and bt_raw != "pending":
        txid = bt_raw
    elif (receipt.get("btc_txid") or "").strip() not in ("", "pending"):
        txid = str(receipt.get("btc_txid")).strip()

    if txid:
        anchor_status = "anchored"
    elif b_status == "anchoring":
        anchor_status = "anchoring"
    elif b_status == "failed":
        anchor_status = "failed"
    elif receipt.get("batch_id"):
        anchor_status = "batched"
    else:
        anchor_status = "queued"

    cadence = daily_equivalent_block_count_for_network(net) if net else None

    proof: Dict[str, Any] = {
        "receipt_id": receipt.get("receipt_id"),
        "receipt_hash_sha256": receipt.get("receipt_hash_sha256"),
        "anchor_network": net or None,
        "anchor_status": anchor_status,
    }
    if receipt.get("batch_id"):
        proof["batch_id"] = receipt.get("batch_id")
    if receipt.get("batch_merkle_root_sha256"):
        proof["batch_merkle_root_sha256"] = receipt.get("batch_merkle_root_sha256")
    if txid:
        proof["anchor_txid"] = txid
    if cadence is not None:
        proof["anchor_cadence_blocks"] = cadence
    if b_err and anchor_status == "failed":
        proof["anchor_error"] = b_err[:2000]
    if isinstance(b_attempts, int) and anchor_status == "failed":
        proof["anchor_attempts"] = b_attempts
    proof = {k: v for k, v in proof.items() if v is not None}
    from backend.anchoring.agreement_proof_enrichment import enrich_agreement_anchor_proof_view

    return enrich_agreement_anchor_proof_view(
        proof,
        receipt=receipt,
        timeline_batch=batch,
    )


@router.post("/{agreement_id}/finalized-receipt")
def post_agreement_finalized_receipt(agreement_id: str, body: AgreementFinalizeReceiptRequest):
    if not _agreements_write_allowed():
        raise HTTPException(status_code=403, detail="verifier_only")

    anchor_network = (body.anchor_network or "").strip()
    if anchor_network not in ALLOWED_AGREEMENT_ANCHOR_NETWORKS:
        raise HTTPException(
            status_code=400,
            detail=f"anchor_network must be one of {sorted(ALLOWED_AGREEMENT_ANCHOR_NETWORKS)}",
        )

    store = _agreements_timeline_store()
    receipt, artifact_body = create_agreement_receipt_response(
        agreement_id=agreement_id,
        finalized_version_id=body.finalized_version_id,
        finalized_at=body.finalized_at,
        content_sha256=body.content_sha256,
        execution_packet_sha256=body.execution_packet_sha256,
        parties_sha256=body.parties_sha256,
        signer_count=body.signer_count,
        anchor_network=anchor_network,
        epoch_id=body.epoch_id,
    )

    try:
        _persist_agreement_execution_packet_artifact(
            agreement_id=agreement_id,
            receipt_id=receipt["receipt_id"],
            finalized_version_id=body.finalized_version_id,
            execution_packet_sha256=body.execution_packet_sha256,
            execution_packet=body.execution_packet,
        )
    except ValueError:
        raise HTTPException(
            status_code=400,
            detail="The execution packet could not be validated. Check your data and try again.",
        ) from None

    try:
        existing = store.get_receipt(receipt["receipt_id"])
        b_existing = _batch_row_for_receipt(store, existing)
        return {
            "ok": True,
            "idempotent": True,
            "receipt": existing,
            "agreement_receipt_body": artifact_body,
            "proof": _agreement_anchor_proof_view(existing, b_existing),
        }
    except KeyError:
        pass

    store.create_receipt(
        receipt_id=receipt["receipt_id"],
        timeline_id=receipt["timeline_id"],
        protocol_version=receipt["protocol_version"],
        network=receipt["network"],
        epoch_id=receipt.get("epoch_id"),
        btc_txid=receipt["btc_txid"],
        commitment=receipt["commitment"],
        merkle_proof=receipt["merkle_proof"],
        zk_proof_refs=receipt.get("zk_proof_refs"),
        issued_at=receipt["issued_at"],
        receipt_hash_sha256=receipt.get("receipt_hash_sha256"),
    )
    out = store.get_receipt(receipt["receipt_id"])
    b_out = _batch_row_for_receipt(store, out)
    return {
        "ok": True,
        "idempotent": False,
        "receipt": out,
        "agreement_receipt_body": artifact_body,
        "proof": _agreement_anchor_proof_view(out, b_out),
    }


@router.get("/{agreement_id}/proof-status")
def get_agreement_proof_status(agreement_id: str):
    store = _agreements_timeline_store()
    timeline_id = f"agreement:{agreement_id}"
    rec = store.get_latest_receipt_for_timeline(timeline_id)
    if not rec:
        return {"proof": None, "cadence_defaults": anchor_cadence_summary()}

    batch = _batch_row_for_receipt(store, rec)

    return {
        "proof": _agreement_anchor_proof_view(rec, batch),
        "cadence_defaults": anchor_cadence_summary(),
    }
